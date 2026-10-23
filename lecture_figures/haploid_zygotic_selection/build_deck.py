"""Build 'BB512 - 14 - Haploid and zygotic selection' from BB512_Theme.pptx.
Run from the lecture folder:  python3 build_deck.py <path/to/BB512_Theme.pptx> <out.pptx>
Inline markup in text: *italic*, **bold**, ~subscript~, ^superscript^.
"""
import re, sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

THEME, OUT = sys.argv[1], sys.argv[2]
FOREST = RGBColor(0x1F, 0x4D, 0x3F); INK = RGBColor(0x1F, 0x26, 0x23)
PALE = RGBColor(0xEF, 0xEE, 0xE6); BG = RGBColor(0xFB, 0xFA, 0xF6)
GREY = RGBColor(0x6E, 0x6F, 0x68); CLAY = RGBColor(0xA4, 0x50, 0x3A); GOLD = RGBColor(0xC0, 0x8A, 0x2E)
FONT = "Avenir Next"

prs = Presentation(THEME)
# Remove the theme's example slides
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
LAYOUT = {l.name: l for l in prs.slide_layouts}

TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|~.+?~|\^.+?\^)")

def add_runs(par, text, size=None, color=None, bold=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        r = par.add_run()
        fmt = {}
        if part.startswith("**"):
            part = part[2:-2]; fmt["bold"] = True
        elif part.startswith("*"):
            part = part[1:-1]; fmt["italic"] = True
        elif part.startswith("~"):
            part = part[1:-1]; fmt["baseline"] = "-25000"
        elif part.startswith("^"):
            part = part[1:-1]; fmt["baseline"] = "30000"
        r.text = part
        if fmt.get("bold"): r.font.bold = True
        elif bold is not None: r.font.bold = bold
        if fmt.get("italic"): r.font.italic = True
        if "baseline" in fmt: r._r.get_or_add_rPr().set("baseline", fmt["baseline"])
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color

def set_text(shape, lines, size=None, color=None):
    tf = shape.text_frame
    if isinstance(lines, str):
        lines = [lines]
    # clear existing paragraphs but keep the first paragraph's properties
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, line in enumerate(lines):
        par = first if i == 0 else tf.add_paragraph()
        level = 0
        while line.startswith("  "):
            level += 1; line = line[2:]
        par.level = level
        add_runs(par, line, size=size, color=color)

def ph(slide, idx):
    for s in slide.placeholders:
        if s.placeholder_format.idx == idx:
            return s
    raise KeyError(f"placeholder idx {idx} not on layout {slide.slide_layout.name}")

def add_slide_number(slide):
    for lph in slide.slide_layout.placeholders:
        if lph.placeholder_format.type is not None and "SLIDE_NUMBER" in str(lph.placeholder_format.type):
            el = copy.deepcopy(lph._element)
            slide.shapes._spTree.append(el)
            # make sure it contains a slidenum field
            txBody = el.find(qn("p:txBody"))
            if txBody is not None and not txBody.findall(".//" + qn("a:fld")):
                pass
            return

def new(layout, fields=None, notes=None, pictures=None, number=True, hidden=False):
    slide = prs.slides.add_slide(LAYOUT[layout])
    if hidden:
        slide._element.set("show", "0")
    for idx, val in (fields or {}).items():
        if isinstance(val, tuple):
            set_text(ph(slide, idx), val[0], size=val[1])
        else:
            set_text(ph(slide, idx), val)
    for idx, path in (pictures or {}).items():
        ph(slide, idx).insert_picture(path)
    # remove unused empty placeholders (keeps the slide clean in Keynote)
    for s in list(slide.placeholders):
        if s.has_text_frame and not s.text_frame.text.strip() and s.placeholder_format.idx not in (fields or {}):
            s._element.getparent().remove(s._element)
    if number and layout not in ("Title", "Photo"):
        add_slide_number(slide)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes.strip()
    return slide

def table(slide, header, rows, col_w, top=2.15, row_h=0.52, size=15, first_col_bold=True, left=0.8,
          highlight_rows=()):
    nrows, ncols = len(rows) + 1, len(header)
    shape = slide.shapes.add_table(nrows, ncols, Inches(left), Inches(top), Inches(sum(col_w)), Inches(row_h * nrows))
    tbl = shape.table
    # strip the default table style banding
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    for i in range(nrows):
        tbl.rows[i].height = Inches(row_h)
        for j in range(ncols):
            cell = tbl.cell(i, j)
            text = header[j] if i == 0 else rows[i - 1][j]
            cell.text = ""
            par = cell.text_frame.paragraphs[0]
            par.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            if i == 0:
                add_runs(par, text, size=size, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
                cell.fill.solid(); cell.fill.fore_color.rgb = FOREST
            else:
                hl = (i - 1) in highlight_rows
                add_runs(par, text, size=size, color=CLAY if hl and j > 0 else INK,
                         bold=(j == 0 and first_col_bold) or hl)
                cell.fill.solid(); cell.fill.fore_color.rgb = PALE if i % 2 == 1 else BG
            for r in par.runs:
                r.font.name = FONT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
    return tbl

def textbox(slide, text, left, top, width, height, size=16, color=GREY, italic=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        add_runs(par, line, size=size, color=color)
        for r in par.runs:
            r.font.name = FONT
            if italic: r.font.italic = True
    return tb

F = "figures/"; I = "images/"

# ---------------------------------------------------------------- 1 Title
new("Title", {103: "Haploid and zygotic selection",
              104: "How fast can natural selection change allele frequencies?"},
    pictures={100: I + "title_birch_typica.jpg"},
    notes="""
(1 min) Open with the photo: there is a peppered moth (*Biston betularia*, typica form) resting on this birch bark. Ask the room to find it. We will come back to this moth: it is the best-known example of natural selection measured in the wild.
Guiding question for today: how fast can selection change allele frequencies, and what decides the speed?
Image: peppered moth on birch bark, carried over from the previous version of the lecture (file Biston_betularia_20110529_102239_8073M.JPG, Wikimedia Commons). Confirm author and licence before sharing the slides outside the course.
""")

# ---------------------------------------------------------------- 2 Objectives
new("Agenda", {101: "Haploid and zygotic selection",
               103: (["Calculate fitness and selection coefficients from data",
                      "Explain how the genetic system decides which selection model applies",
                      "Predict allele-frequency change under haploid selection",
                      "Explain how dominance changes the speed of zygotic selection",
                      "Estimate selection in the wild, and explain how heterozygote advantage maintains variation",
                      "Use selection models to evaluate the claims of eugenics"], 20)},
    notes="""
(1 min) Six objectives, one per section (sections 05 and 06 both serve objective 5). They are written as general skills; the peppered moth, sickle-cell haemoglobin and PKU are the evidence we use to practise them.
Reading: Neal (2004) Chapter 10 (the models) and Chapter 11 (applying them). Note that equation numbers in these slides follow Neal's 2nd edition (2004).
""")

# ---------------------------------------------------------------- 3 Hook
new("Big Fact", {100: "1% → 99%",
                 102: "Dark morph peppered moths around Manchester, 1848 to 1898: about 50 generations. How strong must selection have been?"},
    notes="""
(1 min) The hook. The slide says "dark morph" because students have not yet met the name *carbonaria*; it is introduced in section 05. In 1848 the black form was a collector's rarity; by the end of the century almost every peppered moth around Manchester was black. These are J.B.S. Haldane's (1924) round numbers, as reported by Neal (2004, section 11.2.1). One generation per year, so about 50 generations.
Ask: is that a lot of selection or a little? Hold the question; we answer it in section 05 (s ≈ 0.3), once we have the models.
""")

# ---------------------------------------------------------------- Section 01
new("Section", {102: "01", 103: "Fitness and selection"},
    notes="(10 s) First we need a way to measure how much better one genotype does than another.")

new("Key Concept", {102: "Fitness",
                    103: "Absolute fitness (λ) is how many times a genotype multiplies per generation. Relative fitness (W) is λ divided by the λ of the fittest genotype, so the fittest has W = 1.",
                    105: "Selection coefficient: s = 1 − W. The fittest genotype has s = 0; a lethal genotype has s = 1."},
    notes="""
(2 min) Absolute fitness, λ, is the multiplication rate defined in the population-growth lectures (Neal Chapter 4), applied to a genotype or allele rather than to a whole population.
We usually care about relative fitness, because natural selection is about differences between types. By convention the reference is the genotype with the highest λ (Neal section 10.1).
The selection coefficient s measures the disadvantage: s = 1 − W, and W = 1 − s. It runs from 0 (the reference) to 1 (lethal, λ = 0).
""")

new("Figure & Bullets", {100: "Relative fitness compares growth rates",
                         102: ["λ~A~ = 1.05, λ~B~ = 1.15",
                               "B is fittest: W~B~ = 1",
                               "W~A~ = 1.05 / 1.15 = 0.913",
                               "s~A~ = 1 − 0.913 = 0.087"]},
    pictures={101: F + "01_fitness_growth.png"},
    notes="""
(2 min) Kept from the previous version, redrawn. Both genotypes start from one individual and grow geometrically: N~t~ = λ^t^. After 20 generations B has about 16 descendants and A fewer than 3.
Both genotypes are increasing, yet A is losing ground: selection is about relative, not absolute, success.
Figure: analytical (geometric growth), script 01_fitness_growth.R.
""")

new("Activity", {101: "3 MIN · PAIRS",
                 102: "Which genotype is fittest, and how strong is selection against the others?",
                 103: ["Generation 0 → 1:   A 50 → 100,   B 60 → 90,   C 40 → 40",
                       "Calculate λ, then W, then s, for each genotype",
                       "What does s = 0.5 mean in words?"]},
    notes="""
(3 min) Neal Table 10.1.
Answers: λ = 2, 1.5, 1 for A, B, C. A is the reference: W = 1, 0.75, 0.5; s = 0, 0.25, 0.5.
s = 0.5: genotype C leaves half as many descendants per individual as genotype A.
Point out that C is not declining in number (λ = 1), but it is still being selected against.
""")

# ---------------------------------------------------------------- Section 02
new("Section", {102: "02", 103: "Genetic systems"},
    notes="(10 s) The same fitness differences produce different rates of change depending on how genes are packaged and passed on.")

new("Title & Bullets", {100: "Ploidy and sex decide which model we use",
                        101: ["**Ploidy:** chromosome sets in the adult: haploid (n) or diploid (2n)",
                              "**Reproduction:** asexual (offspring are copies) or sexual (alleles reshuffled every generation)",
                              "Together they decide which model of selection we use",
                              "Other systems exist, e.g. haplodiploid bees and ants"]},
    notes="""
(2 min) Diploid from Greek *diplous*, double; haploid from *haploos*, single.
Ploidy decides whether an allele can be hidden by another allele in the same individual (dominance). Sex decides whether genotypes are inherited intact (clones) or rebuilt each generation from the allele pool.
Diagrams of chromosome sets from the previous version can be shown here if wanted; the next two slides replace the four life-cycle slides.
Source on the old slide: learn.genetics.utah.edu/content/variation/reproduction/
""")

new("Title & Photo", {100: "Four genetic systems",
                      102: "Examples are typical rather than absolute: aphids, for example, are asexual in summer and sexual in autumn."},
    pictures={101: F + "02_genetic_systems.png"},
    notes="""
(2 min) Walk round the four boxes.
Correction from the previous version: dinoflagellates were listed as asexual diploids, but most dinoflagellates are haploid for most of their life cycle (only the zygote is diploid). They are replaced by aphids and water fleas (*Daphnia*), which are diploid in their asexual (parthenogenetic) phases. Avoid apomictic dandelions and many parthenogenetic whiptail lizards as examples: they are mostly triploid. Bdelloid rotifer ploidy is disputed.
The coloured line in each box says which model we will use: anything that reproduces clonally, haploid or diploid, follows the haploid selection model (Neal section 10.3). Only diploid sexual organisms need the zygotic model.
Schematic, script 02_genetic_systems.R.
""")

new("Title & Photo", {100: "Where does selection act in the life cycle?",
                      102: "Blue: haploid (n) stages. Green: diploid (2n) stages. Grey: clones of either ploidy."},
    pictures={101: F + "03_life_cycles.png"},
    notes="""
(2 min) Replaces the three hand-drawn life-cycle slides.
Asexual: offspring are copies, so genotypes are inherited intact.
Haploid-dominant sexual cycle (many fungi and algae): the diploid zygote undergoes meiosis straight away (zygotic meiosis), so the long-lived, selected stage is haploid.
Diploid-dominant cycle (most animals and flowering plants): meiosis happens in the gonads of adults (gametic meiosis); the long-lived stage is diploid. Selection between zygote and adult is **zygotic selection**.
Recombination (the shuffling of alleles during meiosis) is covered later in the course.
""")

new("Key Concept", {102: "Zygotic selection",
                    103: "Selection among diploid genotypes (A~1~A~1~, A~1~A~2~, A~2~A~2~), acting through differences in survival and reproduction between the zygote and the adult.",
                    105: "Haploid selection acts on haploid genotypes: bacteria, any clone, or the haploid phase of a life cycle. Neal calls selection on gametes gametic selection."},
    notes="""
(1 min) The title of the lecture uses this term, so define it explicitly (the previous version never did).
Neal (section 10.2): overall fitness combines components acting in the haploid stage (gametic selection) and in the diploid stage (zygotic selection). We model only zygotic selection on viability, and assume genotypes do not differ in fertility.
""")

# ---------------------------------------------------------------- Section 03
new("Section", {102: "03", 103: "Haploid selection"},
    notes="(10 s) The simplest model: each genotype copies itself.")

new("Title & Bullets", {100: "Haploid selection: the simplest case",
                        101: ["Each individual copies itself: A gives A, B gives B",
                              "Without selection, frequencies stay the same",
                              "With selection, the fitter type spreads until it is **fixed**",
                              "p = frequency of A, q = frequency of B; W~A~ = 1, W~B~ = 1 − s",
                              "Applies to bacteria and any asexual organism"]},
    notes="""
(2 min) Corrects a typo in the previous version ("extinction of A1 and the fixation of A1"): the disfavoured type is lost and the favoured type is fixed.
By Neal's convention, q is always the frequency of the type selected against.
Directional selection like this removes genetic variation; at the end of the lecture we meet a kind of selection that keeps it.
""")

new("Title & Photo", {100: "Selection is a filter",
                      102: "One generation with W~A~ = 1 and W~B~ = 0.75 (s = 0.25). The total after step 2 is the mean fitness."},
    pictures={101: F + "04_selection_filter.png"},
    notes="""
(2 min) Every selection model in this lecture uses the same three steps: (1) start with frequencies; (2) multiply each by its relative fitness; (3) divide by the total so the frequencies sum to 1 again.
The total after step 2 (0.875) is the mean relative fitness of the population.
Result: A rises from 0.5 to 0.571 in one generation.
Analytical, script 04_selection_filter.R.
""")

s = new("Title Only", {100: "Haploid selection: a worked example"},
    notes="""
(4 min) Kept from the previous version, with one correction: the 'contribution' row is frequency multiplied by **relative fitness** (1 − s), not by the selection coefficient.
Walk through row by row. A starts rare (1%) and grows 1.6× per generation; B grows 1.2×. W~B~ = 1.2/1.6 = 0.75, so s = 0.25.
After one generation A has risen from 0.010 to 0.013. Slow at first, because A is rare.
Arithmetic checked in scripts/00_check_numbers.R.
""")
table(s, ["", "Genotype A", "Genotype B", "Total"],
      [["Frequency, generation 0", "p~0~ = 0.01", "q~0~ = 0.99", "1"],
       ["Absolute fitness (λ)", "1.6", "1.2", ""],
       ["Relative fitness (W)", "1.6 / 1.6 = 1", "1.2 / 1.6 = 0.75", ""],
       ["Selection coefficient (s = 1 − W)", "0", "0.25", ""],
       ["Contribution (frequency × W)", "0.01 × 1 = 0.01", "0.99 × 0.75 = 0.7425", "0.7525"],
       ["Frequency, generation 1 (÷ total)", "0.01 / 0.7525 = 0.013", "0.7425 / 0.7525 = 0.987", "1"]],
      col_w=[3.9, 2.75, 3.03, 2.05], row_h=0.62, size=17, highlight_rows=(5,))

new("Key Concept", hidden=True, fields={102: "Δq = −spq / (1 − sq)",
                    103: "Change in frequency of the less fit type in one generation (Neal Eqn 10.1). It depends on s, and on how common both types are (pq).",
                    105: "Check: q = 0.99 and s = 0.25 give Δq = −0.0033, so q falls to 0.987, as in the table."},
    notes="""
(2 min) This is just the table written in algebra: q after selection = q(1 − s)/(1 − sq); subtract q to get Δq (derivation in Neal Box 10.1).
The pq term is the key: change is small when either type is rare, because pq is small. That is the origin of the slow–fast–slow pattern on the next slide.
For many generations, Neal gives the exact solution q~t~ = q~0~ / [q~0~ + (1 − q~0~)(1 − s)^−t^] (Eqn 10.2). Students do not need to memorise it.
""")

new("Title & Photo", {100: "Haploid selection: slow, fast, slow",
                      102: "Worked example continued (p~0~ = 0.01, s = 0.25). Change per generation: Δq = −spq / (1 − sq) (Neal Eqn 10.1)."},
    pictures={101: F + "05_haploid_dynamics.png"},
    notes="""
(2 min) Left: frequencies over 40 generations; A passes 50% at about generation 16 and is above 98% by generation 30. Right: the size of the change per generation against the frequency of A. Change is greatest at intermediate frequencies (here near p = 0.47; the curve is slightly skewed when s is large) and tiny when either type is rare.
The sigmoid shape is the same as logistic growth (Neal Chapter 5).
Biological consequence: a new advantageous mutation is rare for many generations, and while rare it can easily be lost by chance (genetic drift, earlier lecture).
Analytical iteration of Eqn 10.1, script 05_haploid_dynamics.R.
""")

new("Title, Bullets & Photo", {100: "Selection raises mean fitness",
                               103: "TAKEAWAY",
                               105: ("Mean fitness rises from 1.2 to 1.6 as A replaces B. Only relative fitness affects the change in frequency.", 20)},
    pictures={101: F + "06_mean_fitness.png"},
    notes="""
(1 min) Mean fitness each generation = p × 1.6 + q × 1.2 (Neal Fig. 10.2 uses λ = 2 and 1).
Neal's conclusions 3 and 4 (section 10.3.1): selection increases the overall fitness of the population, and only the ratio of growth rates matters.
Script 06_mean_fitness.R.
""")

new("Check Question (4 options)", {101: ("λ~A~ = 4 and λ~B~ = 2, or λ~A~ = 0.8 and λ~B~ = 0.4: in which population does A spread faster?", 30),
                                   105: "The first (growing)", 109: "The second (shrinking)", 113: "The same in both", 117: "Need population sizes"},
    notes="""
(2 min) Answer C. In both, W~B~ = 0.5 and s = 0.5, so the deterministic change in frequency is identical.
Discussion point: D is partly right in practice. A shrinking population becomes small, and then genetic drift matters; our models ignore drift.
""")

# ---------------------------------------------------------------- Section 04
new("Section", {102: "04", 103: "Zygotic selection"},
    notes="(10 s) Now diploid, sexually reproducing organisms: most animals and plants.")

new("Two Columns", {100: "What changes in diploids?",
                    101: "HAPLOID OR CLONAL", 103: ["Genotype = one allele", "Offspring copy their parent", "Every allele is exposed to selection"],
                    104: "DIPLOID AND SEXUAL", 106: ["Two alleles per individual",
                                                    "Random mating rebuilds genotypes: p², 2pq, q²",
                                                    "Dominance decides if selection 'sees' an allele in a heterozygote"]},
    notes="""
(2 min) Recap Hardy–Weinberg from the earlier lecture: with random mating and no evolutionary forces, genotype frequencies are p², 2pq and q². (Correction from the previous version: p² + 2pq + q² = 1 is always true; the equilibrium is that the genotype frequencies equal p², 2pq and q², and it also requires random mating.)
Genotypes are no longer independent: an A~1~A~2~ parent produces A~1~A~1~, A~1~A~2~ and A~2~A~2~ offspring. So we track allele frequency q and rebuild genotypes each generation.
Dominance here means dominance **with respect to fitness**, which can differ from dominance for the visible phenotype (Neal Example 11.3 calls this pleiotropy).
""")

new("Title & Photo", {100: "One generation of zygotic selection",
                      102: "No dominance, s = 0.3 (W = 1, 0.85, 0.70). Random mating restores Hardy–Weinberg proportions at the new q."},
    pictures={101: F + "07_zygotic_cycle.png"},
    notes="""
(2 min) The cycle that the circular arrow on the old table was trying to show.
Zygotes start in Hardy–Weinberg proportions at q = 0.4. Selection between zygote and adult changes the genotype frequencies, and the adults are no longer in Hardy–Weinberg proportions. Their gametes carry A~2~ at q' = 0.127 + 0.464/2 = 0.359. Random mating then produces the next zygotes in Hardy–Weinberg proportions at q' = 0.359.
(The old slide's 33% / 33% / 33% starting frequencies were not Hardy–Weinberg proportions for any q; this version uses consistent numbers.)
Script 07_zygotic_cycle.R.
""")

s = new("Title Only", {100: "Zygotic selection: a worked example"},
    notes="""
(3 min) Same three steps as the haploid filter, applied to genotypes. No dominance (Neal Case B): the heterozygote's fitness is halfway between the homozygotes.
Correction from the previous version: the last row is the genotype frequency in **adults after selection**, not in the next generation. To get the next generation, calculate the new allele frequency, then apply Hardy–Weinberg.
Check with Neal Eqn 10.4: Δq = −½spq/(1 − sq) = −0.5 × 0.3 × 0.6 × 0.4 / 0.88 = −0.041, so q' = 0.359.
""")
tz = table(s, ["", "A~1~A~1~", "A~1~A~2~", "A~2~A~2~", "Total"],
      [["Zygote frequency (q = 0.4)", "p² = 0.36", "2pq = 0.48", "q² = 0.16", "1"],
       ["Relative fitness (W)", "1", "1 − ½s = 0.85", "1 − s = 0.70", ""],
       ["Contribution (frequency × W)", "0.36", "0.408", "0.112", "1 − sq = 0.88"],
       ["Adult frequency (÷ 0.88)", "0.409", "0.464", "0.127", "1"],
       ["New allele frequency", "q' = 0.127 + ½ × 0.464 = 0.359", "", "", ""]],
      col_w=[4.05, 1.75, 2.25, 1.85, 1.83], row_h=0.66, size=17, highlight_rows=(4,))
tz.cell(5, 1).merge(tz.cell(5, 4))

s = new("Title Only", {100: "Four cases of dominance"},
    notes="""
(3 min) Neal Table 10.3 and 10.4, using Neal's case letters so students can find them in the book. Selection is always against A~2~, whose frequency is q.
Neal's example: insect body colour under bird predation. Case B: green, olive and brown insects on green vegetation. Case C: green dominant, brown insects eaten on green vegetation. Case D: the same insects once the vegetation turns brown. Case E: olive heterozygotes best camouflaged on olive vegetation.
Students should recognise the fitness pattern of each case from data; they do not need to memorise the Δq formulas.
Case E has an equilibrium where Δq = 0: q̂ = s~1~ / (s~1~ + s~2~) (Neal Eqn 10.15). We use it for sickle-cell haemoglobin.
The 'favoured allele is' column links the two vocabularies: from the next slide on, the cases are named by the favoured allele (C = favoured dominant, D = favoured recessive), and figures plot the favoured allele's frequency, p.
Why no Case A? Neal starts his lettering at B; the haploid model plays the role of the simplest case.
""")
table(s, ["Neal's case", "A~1~A~1~", "A~1~A~2~", "A~2~A~2~", "So the favoured allele is", "Δq in one generation"],
      [["B: no dominance", "1", "1 − ½s", "1 − s", "Neither dominant", "−½spq / (1 − sq)"],
       ["C: against the recessive", "1", "1", "1 − s", "**Dominant**", "−spq² / (1 − sq²)"],
       ["D: against the dominant", "1", "1 − s", "1 − s", "**Recessive**", "−sp²q / (1 − 2sq + sq²)"],
       ["E: heterozygote superior", "1 − s~1~", "1", "1 − s~2~", "Neither: heterozygote best", "pq(s~1~p − s~2~q) / (1 − s~1~p² − s~2~q²)"]],
      col_w=[2.95, 1.05, 1.05, 1.05, 2.3, 3.33], row_h=0.66, size=15)
textbox(s, "Relative fitness of each genotype. A~2~ (frequency q) is the allele selected against; in Case E both homozygotes are.",
        0.8, 5.45, 11.7, 0.5, size=14)

new("Title, Bullets & Photo", {100: "Dominance shifts the fastest change",
                               103: "TAKEAWAY",
                               105: ("A rare favoured recessive allele hides in heterozygotes, so it barely changes while p is small (red).", 20)},
    pictures={101: F + "08_dq_by_case.png"},
    notes="""
(2 min) Neal Fig. 10.3a, redrawn with s = 0.1 and flipped to show the frequency of the **favoured** allele (p = 1 − q), so the axis matches the next slide. The haploid model (dashed) is added for comparison.
Peaks: no dominance at p ≈ 0.5; favoured dominant (Case C) at p ≈ 0.33; favoured recessive (Case D) at p ≈ 0.67 (Neal section 10.4.1, point 3, gives the same peaks in terms of q; checked numerically in 00_check_numbers.R).
Read it from both ends: a favoured recessive is slow while rare (left of the red curve); a favoured dominant is slow when nearly fixed (right of the green curve), because the last disfavoured recessive copies hide in heterozygotes.
With no dominance, Δq is exactly half the haploid value (compare Eqns 10.1 and 10.4), because the homozygote's disadvantage s is shared out at s/2 per allele.
Script 08_dq_by_case.R.
""")

new("Title & Photo", {100: "How fast does a favoured allele spread?",
                      102: "s = 0.1, p~0~ = 0.01 (compare Neal Fig. 10.4). Deterministic: real rare alleles are often lost by drift first."},
    pictures={101: F + "09_spread_by_case.png"},
    notes="""
(3 min) Replaces the textbook figure (ho04f06) with the same comparison drawn from our own models.
Haploid (dashed) is fastest: >99% in about 90 generations. No dominance: about 175 generations.
A favoured dominant allele rises fastest of the diploid cases while rare (it is expressed in heterozygotes) but then crawls towards fixation: it is still below 99% after 1,000 generations, because the disfavoured allele survives hidden in heterozygotes.
A favoured recessive allele is almost invisible to selection while rare: it takes about 900 generations even to reach 50%.
Correction from the previous version: 'dominant alleles reach fixation relatively rapidly' is not right. They **increase** rapidly but **fix** slowly (Neal section 10.4.1, point 7).
Script 09_spread_by_case.R.
""")

new("Big Fact", {100: "100,000",
                 102: "generations for a favoured recessive allele to go from 0.1% to 50% (s = 0.01); nine-tenths of that time is spent below 1%. A dominant allele gets to 50% in about 800."},
    notes="""
(1 min) Same starting point (0.1%), same s (0.01); only dominance differs. Deterministic iteration of Neal's models (00_check_numbers.R):
Favoured allele recessive (Case D): 0.1% to 1% in 89,332 generations; 0.1% to 50% in 99,493.
No dominance (Case B): 459 and 1,372.
Favoured allele dominant (Case C): 231 and 787.
Neal Table 10.5 (after Crow 1986) gives matching values for the 0.1% to 1% step (90,150, 459 and 231).
Why: at p = 0.001, almost every copy of the allele is in a heterozygote (only a fraction p of copies are in homozygotes). A recessive allele is not expressed there, so selection cannot see it; a dominant allele is expressed in every carrier. Once the recessive allele reaches 1%, it gets to 50% in about another 10,000 generations.
Caveat: the models ignore drift. In a real population a recessive allele at 0.1% would very probably be lost by chance long before selection could act.
A frequency of 0.001 is one new mutation in a population of 500 diploid individuals (2N = 1,000 gene copies).
Extension (mirror image): a favoured dominant allele is slow at the other end. Going from 99% to 99.9% takes it about 90,000 generations, because the last disfavoured recessive copies hide in heterozygotes (see the previous slide's figure).
""")

new("Check Question (4 options)", {101: ("Which new advantageous mutation takes longest to start spreading?", 30),
                                   105: "In a haploid bacterium", 109: "Diploid, no dominance",
                                   113: "Diploid, dominant", 117: "Diploid, recessive"},
    notes="""
(2 min) Answer D. While rare, a recessive allele occurs almost only in heterozygotes, where it is not expressed, so selection cannot see it. It is also at the mercy of drift for longer.
Follow-up: which case takes longest to reach fixation once common? C (favoured dominant), because the last copies of the disfavoured recessive hide in heterozygotes.
""")

new("Figure & Bullets", {100: "Why it matters: microbes evolve fast",
                         102: ["Haploid: new mutations are never hidden",
                               "Generations of hours, huge populations",
                               "Drug resistance in TB can evolve within one patient",
                               "Hosts keep up through sex and immunity"]},
    pictures={101: I + "mtb_sem.jpg"},
    notes="""
(2 min) Most microbes are haploid; most animals and plants are diploid. Time in our models is measured in generations, and bacterial generations are short.
Correction from the previous version: 'haploid asexual pathogens will tend to win coevolutionary struggles with diploid sexual hosts' overstates the case. Ploidy is one advantage among several (generation time and population size matter more), and hosts persist; the Red Queen idea comes later in the course.
TB treatment uses several drugs together for months because resistance to a single drug evolves readily.
Image: scanning electron micrograph of *Mycobacterium tuberculosis* (CDC Public Health Image Library; public domain), carried over from the previous version.
""")

new("Title & Bullets", {100: "What the models tell us",
                        101: (["Change is fastest at intermediate frequencies, where variation in fitness is greatest",
                              "For the same s, haploid change is twice as fast as diploid with no dominance",
                              "Rare recessives hide in heterozygotes: slow to spread if favoured, slow to remove if harmful",
                              "A favoured dominant spreads fast but takes very long to fix",
                              "The more variation in fitness, the faster mean fitness rises (Fisher)"], 20)},
    notes="""
(2 min) Synthesis of Neal section 10.4.1.
Correction from the previous version: Fisher's (1930) 'fundamental theorem of natural selection' (not 'of evolution') concerns the rate of increase in **mean fitness**, and it is the **additive** genetic variance in fitness, not total genetic variance, that sets it. That is why dominance shifts the peak of Δq away from q = 0.5.
On the slide this is simplified to 'more variation in fitness, faster change', because additive genetic variance is not introduced until the quantitative genetics lecture. Precise statement: the rate of increase in mean fitness due to natural selection equals the additive genetic variance in relative fitness (divided by mean fitness in absolute terms).
The 'twice as fast' bullet depends on Neal's definition of s (the diploid homozygote has fitness 1 − s, the heterozygote 1 − s/2).
Assumptions behind all of these models (Neal section 10.5): one locus of large effect; constant selection; no mutation, migration, drift or linkage to other selected genes.
""")

# ---------------------------------------------------------------- Section 05
new("Section", {102: "05", 103: "Industrial melanism"},
    notes="(10 s) Back to the peppered moth: now we can answer the opening question.")

new("Title, Bullets & Photo", {100: "Did you spot the moth?",
                               103: "THE OPENING PHOTO",
                               105: ("A pale *typica* peppered moth on lichen-covered birch bark. Its camouflage is what birds select on.", 20)},
    pictures={101: I + "spot_the_moth_reveal.jpg"},
    notes="""
(1 min) The reveal for the title-slide photo. Ask who found it. Now imagine the same bark blackened by soot: the pale moth would stand out, and a dark one would vanish. That is the hypothesis this section tests.
Image: as on the title slide (file Biston_betularia_20110529_102239_8073M.JPG, Wikimedia Commons; licence to confirm), with the moth circled.
""")

new("Title, Bullets & Photo", {100: "The peppered moth, *Biston betularia*",
                               103: "TWO FORMS",
                               105: ("*typica* (left) is pale; *carbonaria* (right) is black and dominant. It is caused by a transposable element in the *cortex* gene that arose in the early 1800s.", 20)},
    pictures={101: I + "morphs_pair.jpg"},
    notes="""
(1 min) Other dark forms (*insularia*) exist; following Neal, we treat the system as two alleles with *carbonaria* (C) dominant to *typica* (c).
New since the textbook: van't Hof et al. (2016) Nature 534: 102–105 identified the mutation as a large transposable element inserted in the *cortex* gene, and dated its origin to about 1819. That is a statistical estimate with an uncertainty of a few decades, which is why the slide says 'early 1800s': the first known *carbonaria* specimen was caught before 1811 (next slide). (Correction from the previous version: 'a single mutation' is right, but its nature is now known.)
Images: Olaf Leillinger, Wikimedia Commons (files Biston.betularia.7200.jpg and Biston.betularia.f.carbonaria.7209.jpg), CC BY-SA 2.5 as far as I know: please confirm on Commons before redistributing.
""")

new("Figure & Bullets", {100: "The rise of *carbonaria*",
                         102: ["First specimens: before 1811, then Manchester 1848",
                               "By the 1890s: over 90% near Manchester",
                               "Spread by migration; rose where favoured",
                               "Soot darkened bark; pollution killed lichens"]},
    pictures={101: I + "uk_morph_map.jpg"},
    notes="""
(2 min) Neal section 11.2.1. Corrections from the previous version: the first *carbonaria* was caught 'prior to 1811' (exact date unknown; the specimen is in Oxford). 'It spread rapidly through migration' is incomplete: migration carried the allele between areas, but selection explains why it became common in industrial areas and stayed rare in rural ones.
Map: frequencies of *typica* (white), *carbonaria* (black) and *insularia* in Britain, from Lees (1981) as reproduced in Neal (2004) Fig. 11.3 (carried over from the previous version; the copyright is Elsevier's, so keep it for teaching use only).
Note Norfolk and Lincolnshire: rural but downwind of the industrial Midlands.
""")

s = new("Title Only", {100: "Haldane's question: how strong?"},
    notes="""
(3 min) Set-up, following Haldane (1924) and Neal (2004) section 11.2.1.
*typica* is the recessive homozygote cc, so its frequency is q². In 1848 typica = 0.99, so q~0~ = √0.99 = 0.995. In 1898 typica = 0.01, so q~t~ = √0.01 = 0.1.
Selection is against the recessive: Neal's Case C. 50 years = 50 generations (one generation per year).
Correction from the previous version: the reference is Neal (2004) Eqn 10.11 in section 11.2.1, not 'Box 15.1' (that was the 1st edition).
""")
table(s, ["Year", "*carbonaria* (CC + Cc)", "*typica* (cc) = q²", "Allele frequency of c (q)"],
      [["1848", "0.01", "0.99", "q~0~ = √0.99 = 0.995"],
       ["1898", "0.99", "0.01", "q~t~ = √0.01 = 0.10"]],
      col_w=[1.7, 3.2, 3.0, 3.83], row_h=0.62, size=17)
textbox(s, ["Selection is against the recessive (*typica*): **Case C**, over t = 50 generations.",
            "Solve for s (Neal Eqn 10.11):   s × t = (q~0~ − q~t~) / (q~0~q~t~) + ln[q~0~(1 − q~t~) / (q~t~(1 − q~0~))]"],
        0.8, 4.4, 11.7, 1.4, size=20, color=INK)

new("Title, Bullets & Photo", {100: "Selection against *typica* was strong",
                               103: "TAKEAWAY",
                               105: ("1% to 99% in 50 generations needs s ≈ 0.3: *typica* had about 70% of the fitness of *carbonaria*.", 20)},
    pictures={101: F + "10_haldane_carbonaria.png"},
    notes="""
(2 min) Neal's approximate equation (Eqn 10.11) gives s = 0.33, so W~typica~ = 0.67; equivalently *carbonaria* had about a 50% advantage (1/0.67 = 1.49), which is how Haldane expressed it.
Iterating the exact model gives s = 0.30. Neal warns that the approximations are accurate only when s is small (about 0.01); with s this large the approximation overestimates by about 10%. Either way, selection was far stronger than anyone then thought possible.
Haldane did not know what was causing it. That came 30 years later.
Model trajectory: deterministic Case C, script 10_haldane_carbonaria.R; the two points are Haldane's round figures, not a time series.
""")

new("Activity", {101: "4 MIN · PAIRS",
                 102: ("With W(*typica*) = 0.67, how long for *carbonaria* to rise from 10% to 90%?", 32),
                 103: ["Convert each phenotype frequency to q (remember *typica* = q²)",
                       "Decide which case of dominance applies",
                       "Use s × t = (q~0~ − q~t~) / (q~0~q~t~) + ln[q~0~(1 − q~t~) / (q~t~(1 − q~0~))]"]},
    notes="""
(4 min) Neal problem 11.4, 2(a).
q~0~ = √0.90 = 0.949; q~t~ = √0.10 = 0.316. Case C, s = 0.33.
Eqn 10.11: st = (q~0~ − q~t~)/(q~0~q~t~) + ln[q~0~(1 − q~t~)/(q~t~(1 − q~0~))] = 2.11 + 3.70 = 5.81, so t ≈ 17.6 generations. Exact iteration: 16 generations.
So most of the change in Manchester could have happened in under 20 years.
""")

new("Title, Bullets & Photo", {100: "Camouflaged moths are recaptured more",
                               103: "TAKEAWAY",
                               105: ("Kettlewell released and recaptured marked moths. Each form was recaptured about twice as often where it was camouflaged.", 20)},
    pictures={101: F + "11_kettlewell.png"},
    notes="""
(3 min) H.B.D. Kettlewell (1955, 1956) released marked moths of both forms in a polluted wood near Birmingham and an unpolluted wood in Dorset, then recaptured them at light traps over the following days (data: Neal Table 11.1).
Relative fitness: typica in Birmingham W = 0.47 (naively) or 0.45 after Neal's log correction; carbonaria in Dorset W = 0.51, or 0.75 after correction. The correction (ratio of log survival rates, Neal Example 11.2) is needed because survival is compounded over several days.
These data show differential survival, not its cause; recapture rates also depend on dispersal and on how moths respond to traps. The evidence that birds were responsible comes next: Kettlewell, with Niko Tinbergen, filmed birds taking resting moths, which convinced sceptics who doubted that birds ate many moths.
(The previous version attributed the dead-moth experiments to Kettlewell; those were mainly by others, e.g. Bishop 1972: see next slide.)
Script 11_kettlewell.R.
""")

new("Figure & Bullets", {100: "Bird predation: criticised, confirmed",
                         102: (["Kettlewell and Tinbergen filmed birds taking moths",
                               "Critics: unnatural densities and resting sites",
                               "Dead moths on trees gave the same pattern",
                               "Majerus: 4,864 moths, natural densities, six years",
                               "Birds now take more *carbonaria*: s ≈ 0.09 **per day**, not per generation"], 20)},
    pictures={101: I + "morphs_on_bark.jpg"},
    notes="""
(2 min) Kettlewell's methods were criticised from the 1970s onwards, and the peppered moth became a target for creationist attacks in the 1990s and 2000s.
Majerus designed the definitive test: moths released singly at night into a 1 ha rural garden, resting positions recorded, survival checked next morning, and bird predation watched with binoculars (nine bird species). He died in 2009; the results were published by Cook, Grant, Saccheri and Mallet (2012) Biology Letters 8: 609–612. Maximum-likelihood selection against *carbonaria*: 0.091 per day (95% CI 0.028–0.157). Stress the units: this is per day of adult life, so it cannot be compared directly with the per-generation values of s elsewhere in the lecture (e.g. 0.3 from Haldane).
Also: Bishop's (1972) Sefton Park data give W(typica) = 0.795 relative to carbonaria (Neal Example 11.2).
Image: the two forms on bark, carried over from the previous version; source and licence to confirm.
""")

new("Title, Bullets & Photo", {100: "Clean air: the decline of *carbonaria*",
                               103: "TAKEAWAY",
                               105: ("Clean air reversed selection: now the dominant *carbonaria* allele is selected against (so q is now *carbonaria*). At West Kirby: 93% to 33%, s ≈ 0.15.", 18)},
    pictures={101: F + "12_west_kirby.png"},
    notes="""
(2 min) Data: West Kirby, near Liverpool (Neal section 11.2.1): 93% (1959), 84% (1976), 33% (1990).
Point out the switch: in Haldane's calculation q was the *typica* allele; here q is the *carbonaria* allele, because Neal's q is always the allele selected against. Now *carbonaria* is the dominant allele being selected against: Neal's Case D. After the Clean Air Act (1956).
The constant-s model misses the 1976 point (75% predicted, 84% observed), consistent with selection strengthening over time. typica = p², so p = √(typica frequency).
Exact fit from 1959 to 1990: s = 0.15 (Neal's approximation gives 0.16). Fitting each interval separately (Neal) gives 0.11 for 1959–1976 and 0.23 for 1976–1990: selection strengthened as lichens returned.
Correction from the previous version: the main UK legislation was the Clean Air Act 1956 (extended 1968), not the 1960s–1970s.
*carbonaria* is now rare across Britain. Selection reversed, and the moth is close to its pre-industrial state.
Script 12_west_kirby.R.
""")

# ---------------------------------------------------------------- Section 06
new("Section", {102: "06", 103: "Heterozygote advantage"},
    notes="(10 s) Case E: when the heterozygote is fittest, selection keeps both alleles.")

new("Figure & Bullets", {100: "Sickle-cell haemoglobin",
                         102: (["HbS: a point mutation in β-globin",
                               "SS: sickle-cell disease; often fatal in childhood without treatment",
                               "AS: usually healthy, and protected against severe malaria",
                               "AA: no anaemia, but vulnerable to malaria"], 22)},
    pictures={101: I + "sickle_cells.jpg"},
    notes="""
(2 min) Corrections from the previous version: the malaria protection belongs to heterozygous carriers (AS), not to people with sickle-cell disease (SS). 'SS = death' was true historically in malarial Africa without treatment, but with modern care many people with sickle-cell disease live into adulthood, so phrase this carefully. Protection in AS individuals is strong against severe malaria but not complete.
Image: normal red blood cells and a sickle cell, coloured scanning electron micrograph (file 1911_Sickle_Cells.jpg, OpenStax Anatomy and Physiology, CC BY 4.0, as far as I know: confirm), carried over from the previous version.
""")

s = new("Title Only", {100: "Fitness in Tanzania (Allison 1956)"},
    notes="""
(3 min) Allison (1956) genotyped 287 infants and 654 adults in Tanzania (Neal Table 11.3). Neal calls the ratio of adult to infant frequency 'absolute fitness'; the slide calls it a survival ratio, because that is what it is. It compares two different groups sampled at the same time (not the same children followed to adulthood), and it measures survival only. We assume the genotypes are equally fertile, which is probably not true.
The relative fitnesses fit Case E: s~1~ = 0.24 against AA (malaria), s~2~ = 0.80 against SS (sickle-cell disease).
Note on numbers: the previous slide's values (λ = 1.228 and 0.244; W~SS~ = 0.199) are correct. Neal's printed table gives 0.2420 and 0.1971 for SS, which do not follow from his own counts.
Expected HWE in infants from q = 0.186: 0.662, 0.303, 0.035 (observed 0.659, 0.310, 0.031).
""")
table(s, ["", "AA", "AS", "SS", "Total", "q (HbS)"],
      [["Infants (number)", "189", "89", "9", "287", ""],
       ["Infants (frequency)", "0.659", "0.310", "0.031", "1", "0.186"],
       ["Adults (number)", "400", "249", "5", "654", ""],
       ["Adults (frequency)", "0.612", "0.381", "0.008", "1", "0.198"],
       ["Survival ratio (adult ÷ infant frequency)", "0.929", "1.228", "0.244", "", ""],
       ["Relative fitness (W)", "0.757", "1", "0.199", "", ""],
       ["Selection coefficient", "s~1~ = 0.24", "0", "s~2~ = 0.80", "", ""]],
      col_w=[4.35, 1.5, 1.3, 1.5, 1.3, 1.78], row_h=0.55, size=16, highlight_rows=(5, 6))

new("Title, Bullets & Photo", {100: "Heterozygote advantage keeps both",
                               103: "TAKEAWAY",
                               105: ("Equilibrium q = s~1~ / (s~1~ + s~2~) = 0.23, close to the 0.19 to 0.20 observed: a balanced polymorphism.", 20)},
    pictures={101: F + "13_sickle_equilibrium.png"},
    notes="""
(2 min) Neal Eqn 10.15. At equilibrium Δq = 0, which happens when s~1~p = s~2~q.
Trajectories use the fitnesses estimated on the previous slide (deterministic, four starting frequencies). The model predicts 0.23; the observed 0.19–0.20 is close, given that only survival (not fertility) was measured.
This links back to the 'heterozygote advantage' slide in the mutation and genetic variation lecture, which used stylised fitnesses; these are the estimated ones.
Script 13_sickle_equilibrium.R.
""")

new("Title & Photo", {100: "Where malaria is common, so is HbS",
                      102: "(a) HbS allele frequency; (b) historical range of malaria. Without malaria, HbS declines, but slowly. Source to confirm."},
    pictures={101: I + "hbs_malaria_map_wide.jpg"},
    notes="""
(2 min) Say: HbS is common only where malaria is, or was. Without malaria it declines, but slowly, because most copies are in healthy carriers. Heterozygote advantage is one way selection maintains genetic variation.
The map is now shown as two large panels; the source image is low resolution, so a better copy would help.
Correction from the previous version: the heterozygote is not 'maladaptive' outside malarial regions; it simply loses its advantage, while SS remains strongly disadvantageous, so the S allele declines.
I have removed the '~2% per generation decline in the US Black population'. Neal (section 11.2.2) gives a different figure (about 7% per generation, part of it due to gene flow), and neither has a primary source I could check. If you want a number, cite a primary source.
Map: (a) HbS allele frequency and (b) historical distribution of malaria. Carried over from the previous version; source not recorded (a textbook figure), so confirm and add the credit.
""")

# ---------------------------------------------------------------- Section 07
new("Section", {102: "07", 103: "Selection against recessive alleles"},
    notes="(10 s) Case C with s = 1, and why it matters for a dark chapter of history.")

new("Figure & Bullets", {100: "Eugenics",
                         102: ["Aim: stop people with 'undesirable' traits from reproducing",
                               "Forced sterilisation in many countries, including Denmark (1929 to 1967)",
                               "Ethically indefensible, and genetically ineffective"]},
    pictures={101: I + "eugenics_poster.jpg"},
    notes="""
(2 min) The main objection to eugenics is ethical: it violates individual rights and was used to justify forced sterilisation and, under the Nazis, murder. The Nordic countries all had sterilisation laws; Denmark's ran from 1929 to 1967 (nordics.info, Aarhus University). The previous version called it 'a foolish endeavour'; the ethical point should come first, then the science.
Image: 1930s eugenics poster, 'Release the strangle-hold of hereditary disease and unfitness', carried over from the previous version; source and licence to confirm.
""")

new("Big Fact", {100: "99%",
                 102: "of the alleles behind a recessive condition seen in 1 in 10,000 births are carried by healthy heterozygotes. Selection, natural or imposed, cannot see them."},
    notes="""
(1 min) Incidence q² = 1/10,000, so q = 0.01. Carriers: 2pq = 0.0198, about 1 person in 50. Each carrier has one copy and each affected person two, so the fraction of copies in carriers is pq/(pq + q²) = p = 0.99.
Correction from the previous version: 'Carriers must be homozygous to have the disease' should read 'People must be homozygous to have the disease; carriers are heterozygous and unaffected.'
""")

new("Title, Bullets & Photo", {100: "Removing a rare recessive is very slow",
                               103: "TAKEAWAY",
                               105: ("Even with s = 1, halving an incidence of 1 in 10,000 takes 1/q~t~ − 1/q~0~ = 41 generations: about 1,000 years.", 20)},
    pictures={101: F + "14_eugenics.png"},
    notes="""
(3 min) Phenylketonuria (PKU) is recessive and occurs in about 1 in 10,000 births in Europe (roughly 1 in 12,000 in Denmark). It is now detected by newborn screening and treated with diet, so people with PKU live normal lives: an example where treatment, not selection, is the answer.
Calculation (Neal Eqn 10.13, Case C with s = 1): q~0~ = 0.01, q~t~ = √0.00005 = 0.00707, so t = 141.4 − 100 = 41 generations, about 1,000 years at 25 years per generation. Exact iteration: 42. (The equation was missing from the previous slide.)
Neal Example 11.5: a dictator wanting a recessive trait reduced from 1 in 10,000 to 1 in 1,000,000 would wait 1,000 − 100 = 900 generations, about 22,500 years. This ignores new mutations, which would slow it further.
Script 14_eugenics.R.
""")

new("Title & Bullets", hidden=True, fields={100: "Mutation–selection balance",
                        101: ["New mutations replace the alleles selection removes",
                              "Achondroplasia in Denmark (Mørch 1941): 10 in 94,075 births; s = 0.80",
                              "Mutation rate, estimated two ways: 4.3 × 10⁻⁵ per generation",
                              "When selection is already strong, more selection changes little: even s = 1 gives 8.5 rather than 10.6 per 100,000"]},
    notes="""
(2 min) HIDDEN by default to keep the lecture within time; unhide it if there is time. Neal Example 11.6 and Box 11.1; a Danish study.
Framing: present this as the population genetics of new mutations, not as a policy question. Coming straight after the eugenics slides, avoid any wording that treats people with achondroplasia as a problem to remove.
Achondroplasia is dominant, so Neal Eqn 11.3 applies: equilibrium heterozygote frequency H = 2u/s. H = 10/94,075 = 0.000106 and s = 1 − (27/108)/(582/457) = 0.80, giving u = 4.3 × 10⁻⁵.
Directly: 8 of the 10 affected babies had unaffected parents, so they carried new mutations: u = 8/(2 × 94,075) = 4.3 × 10⁻⁵. The close match is partly luck; both estimates are imprecise.
With s = 1, H = 2u = 0.000085. When selection is already strong, imposing more achieves almost nothing.
For a recessive, the equilibrium is q̂ = √(u/s) (Neal Eqn 11.2).
""")

# ---------------------------------------------------------------- Summary
new("Summary", {101: "Key points",
                102: (["Fitness is relative: s = 1 − W",
                      "Genetic system decides the model: haploid or zygotic",
                      "Slow, fast, slow; dominance sets the speed",
                      "Selection in the wild can be strong (s ≈ 0.3)",
                      "Heterozygote advantage maintains variation",
                      "Eugenics: unethical and ineffective"], 20),
                106: (["**Read:** Neal (2004), Chapters 10 and 11",
                      "**Practise:** Neal problems 11.4"], 16)},
    notes="""
(2 min) One summary point per section. Practical: if this lecture is followed by an R practical on selection models, link it here (add the URL to these notes). The models in R/models.R in the lecture-figures folder could be adapted for it.
""")

new("Activity", {101: "AFTER THE LECTURE · R",
                 102: ("Can you reproduce the spread of *carbonaria* in R?", 34),
                 103: ["Write a function for one generation of zygotic selection: × W, ÷ total, new q",
                       "Start the *carbonaria* allele at 0.005, s = 0.30 against *typica* (Case C); iterate 50 generations",
                       "Repeat with *carbonaria* recessive (Case D). How long now?"]},
    notes="""
(1 min) Points students to independent practice. If there is an R practical for this lecture, add its URL here.
Model answer: R/models.R in the lecture-figures folder (zygotic_step and iterate). With carbonaria recessive and s = 0.3, the rise from 1% is far slower at first: the same idea as the 100,000-generations slide.
""")

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")

"""Build "BB512 - 06 - Age and stage-structured models" from BB512_Theme.pptx.
Run from the lecture folder:  python3 build_deck.py <path/to/BB512_Theme.pptx> <out.pptx>
Inline markup in text: *italic*, **bold**, ~subscript~, ^superscript^, `code` (Menlo).
"""
import re, sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

THEME, OUT = sys.argv[1], sys.argv[2]
FOREST = RGBColor(0x1F, 0x4D, 0x3F); INK = RGBColor(0x1F, 0x26, 0x23)
PALE = RGBColor(0xEF, 0xEE, 0xE6); BG = RGBColor(0xFB, 0xFA, 0xF6)
GREY = RGBColor(0x6E, 0x6F, 0x68); CLAY = RGBColor(0xA4, 0x50, 0x3A); GOLD = RGBColor(0xC0, 0x8A, 0x2E)
FONT = "Avenir Next"

prs = Presentation(THEME)
# Remove the theme's example slides
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
LAYOUT = {l.name: l for l in prs.slide_layouts}

TOKEN = re.compile(r"(`.+?`|\*\*.+?\*\*|\*.+?\*|~.+?~|\^.+?\^)")

def add_runs(par, text, size=None, color=None, bold=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        r = par.add_run()
        fmt = {}
        if part.startswith("`"):
            part = part[1:-1]; fmt["code"] = True
        elif part.startswith("**"):
            part = part[2:-2]; fmt["bold"] = True
        elif part.startswith("*"):
            part = part[1:-1]; fmt["italic"] = True
        elif part.startswith("~"):
            part = part[1:-1]; fmt["baseline"] = "-25000"
        elif part.startswith("^"):
            part = part[1:-1]; fmt["baseline"] = "30000"
        r.text = part
        if fmt.get("bold"): r.font.bold = True
        elif bold is not None: r.font.bold = bold
        if fmt.get("italic"): r.font.italic = True
        if fmt.get("code"): r.font.name = "Menlo"
        if "baseline" in fmt: r._r.get_or_add_rPr().set("baseline", fmt["baseline"])
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color

def set_text(shape, lines, size=None, color=None):
    tf = shape.text_frame
    if isinstance(lines, str):
        lines = [lines]
    # clear existing paragraphs but keep the first paragraph's properties
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, line in enumerate(lines):
        par = first if i == 0 else tf.add_paragraph()
        level = 0
        while line.startswith("  "):
            level += 1; line = line[2:]
        par.level = level
        add_runs(par, line, size=size, color=color)

def ph(slide, idx):
    for s in slide.placeholders:
        if s.placeholder_format.idx == idx:
            return s
    raise KeyError(f"placeholder idx {idx} not on layout {slide.slide_layout.name}")

def add_slide_number(slide):
    for lph in slide.slide_layout.placeholders:
        if lph.placeholder_format.type is not None and "SLIDE_NUMBER" in str(lph.placeholder_format.type):
            el = copy.deepcopy(lph._element)
            slide.shapes._spTree.append(el)
            # make sure it contains a slidenum field
            txBody = el.find(qn("p:txBody"))
            if txBody is not None and not txBody.findall(".//" + qn("a:fld")):
                pass
            return

def new(layout, fields=None, notes=None, pictures=None, number=False, hidden=False, keep=()):
    slide = prs.slides.add_slide(LAYOUT[layout])
    if hidden:
        slide._element.set("show", "0")
    for idx, val in (fields or {}).items():
        if isinstance(val, tuple):
            set_text(ph(slide, idx), val[0], size=val[1])
        else:
            set_text(ph(slide, idx), val)
    for idx, path in (pictures or {}).items():
        ph(slide, idx).insert_picture(path)
    # remove unused empty placeholders (keeps the slide clean in Keynote)
    for s in list(slide.placeholders):
        if s.has_text_frame and not s.text_frame.text.strip() and s.placeholder_format.idx not in (fields or {}) \
                and s.placeholder_format.idx not in keep:
            s._element.getparent().remove(s._element)
    # No slide numbers: BB512 decks are built without them (the layouts' number placeholders are not copied)
    if number:
        pass
    if notes:
        slide.notes_slide.notes_text_frame.text = notes.strip()
    return slide

def table(slide, header, rows, col_w, top=2.15, row_h=0.52, size=15, first_col_bold=True, left=0.8,
          highlight_rows=()):
    nrows, ncols = len(rows) + 1, len(header)
    shape = slide.shapes.add_table(nrows, ncols, Inches(left), Inches(top), Inches(sum(col_w)), Inches(row_h * nrows))
    tbl = shape.table
    # strip the default table style banding
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    for i in range(nrows):
        tbl.rows[i].height = Inches(row_h)
        for j in range(ncols):
            cell = tbl.cell(i, j)
            text = header[j] if i == 0 else rows[i - 1][j]
            cell.text = ""
            par = cell.text_frame.paragraphs[0]
            par.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            if i == 0:
                add_runs(par, text, size=size, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
                cell.fill.solid(); cell.fill.fore_color.rgb = FOREST
            else:
                hl = (i - 1) in highlight_rows
                add_runs(par, text, size=size, color=CLAY if hl and j > 0 else INK,
                         bold=(j == 0 and first_col_bold) or hl)
                cell.fill.solid(); cell.fill.fore_color.rgb = PALE if i % 2 == 1 else BG
            for r in par.runs:
                r.font.name = FONT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
    return tbl

def textbox(slide, text, left, top, width, height, size=16, color=GREY, italic=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        add_runs(par, line, size=size, color=color)
        for r in par.runs:
            if r.font.name != "Menlo":
                r.font.name = FONT
            if italic: r.font.italic = True
    return tb




from pptx.enum.shapes import MSO_CONNECTOR
BLUE = RGBColor(0x3E, 0x6E, 0x7E); SAGE = RGBColor(0x9A, 0xAE, 0x9F)

def line(slide, x1, y1, x2, y2, color=INK, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    return c

def matrix(slide, cells, left, top, cw=1.0, rh=0.6, size=22, colors=None, bold=None,
           col_labels=None, row_labels=None, label_size=14, name=None, name_size=26):
    """Editable matrix: one text box per cell plus square brackets drawn as lines.
    cells: list of rows of strings; colors: same shape (RGBColor or None for grey zeros)."""
    nr, nc = len(cells), len(cells[0])
    W, H = nc * cw, nr * rh
    if name:
        textbox(slide, name, left - 1.25, top + H / 2 - 0.3, 1.1, 0.6, size=name_size, color=INK,
                align=PP_ALIGN.RIGHT)
    for i, row in enumerate(cells):
        for j, val in enumerate(row):
            col = colors[i][j] if colors else INK
            if col is None:
                col = SAGE
            tb = textbox(slide, val, left + j * cw, top + i * rh, cw, rh, size=size, color=col,
                         align=PP_ALIGN.CENTER)
            tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            tb.text_frame.word_wrap = False
            tb.text_frame.margin_left = tb.text_frame.margin_right = 0
            if bold and bold[i][j]:
                for r in tb.text_frame.paragraphs[0].runs: r.font.bold = True
    t = 0.14
    for x, s in ((left - 0.05, 1), (left + W + 0.05, -1)):
        line(slide, x, top, x, top + H)
        line(slide, x, top, x + s * t, top); line(slide, x, top + H, x + s * t, top + H)
    if col_labels:
        for j, lab in enumerate(col_labels):
            textbox(slide, lab, left + j * cw, top - 0.45, cw, 0.4, size=label_size, color=GREY,
                    align=PP_ALIGN.CENTER)
    if row_labels:
        for i, lab in enumerate(row_labels):
            tb = textbox(slide, lab, left + W + 0.15, top + i * rh, 1.6, rh, size=label_size, color=GREY)
            tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    return left + W

def picture(slide, path, left, top, width=None, height=None):
    kw = {}
    if width: kw["width"] = Inches(width)
    if height: kw["height"] = Inches(height)
    return slide.shapes.add_picture(path, Inches(left), Inches(top), **kw)


F = "figures/"; I = "images/"
BOOK = "https://jonesor.github.io/BB512_Book/"
URL_MPM = BOOK + "matrix-population-modelling.html"
URL_CENSUS = BOOK + "pre--and-post-reproduction-census.html"
URL_LTRE = BOOK + "life-table-response-experiments.html"
URL_STOCH = BOOK + "stochastic-population-growth.html"

# ================================================================ 1 Title
new("Title", {103: ("Age- and stage-structured models", 38),
              104: "Why do individuals of different ages and stages contribute differently to population growth?"},
    pictures={100: I + "title_hatchling.jpg"},
    notes="""
(30 s) Guiding question for today. In the exponential and logistic models every individual was identical. Real populations are made of newborns, juveniles, breeders and old individuals that survive and reproduce at very different rates.
Photo: leatherback turtle (*Dermochelys coriacea*) hatchling, carried over from the previous version of the lecture (source and licence to confirm: see images/CREDITS.md).
""")

# ================================================================ 2 Objectives
new("Agenda", {101: "Structured populations",
               103: (["Decide whether an age- or a stage-based model suits a life cycle",
                      "Turn a life table or life cycle diagram into a projection matrix, including fertility that allows for survival",
                      "Project a population and interpret λ, the stable stage distribution and transient dynamics",
                      "Use sensitivity, elasticity and *in silico* experiments to compare management options",
                      "State the assumptions of the basic model, and how stochastic models relax one of them"], 18)},
    notes="""
(1 min) One objective per section. The sea turtles, the resident killer whales and the loggerhead conservation case are the evidence we use to practise them.
Reading: Neal (2004) Chapter 15, section 15.3 "Matrix models" (the course schedule lists "Neal 7", which is the chapter number in the current edition).
Builds on the previous lecture (life tables: l~x~, m~x~, p~x~; same notation as Neal and lecture 05) and on geometric growth (n~t+1~ = λn~t~).
After the lecture: practical "Matrix population modelling" (""" + URL_MPM + """), with "Pre- and post-reproduction census" and "Life Table Response Experiments" as follow-ups.
""")

# ================================================================ Section 01
new("Section", {102: "01", 103: "Age or stage?"},
    notes="(10 s) First: what kind of structure does a population have?")

new("Title, Bullets & Photo", {100: "Age-based: when age is what matters",
                               103: "AGE",
                               105: (["Survival and reproduction change with age", "We can tell (or know) how old individuals are",
                                      "Examples: humans, many mammals and birds"], 18)},
    pictures={101: I + "ages.jpg"},
    notes="""
(1.5 min) Age-based models work well when we know individuals' ages and when age itself is meaningful for the organism: survival and fecundity change predictably as individuals get older. This is the world of the life tables from the previous lecture.
Illustration: Freepik ("person-different-ages-style"), carried over from the previous version; Freepik's free licence requires attribution: check before sharing the slides outside the course.
""")

new("Title, Bullets & Photo", {100: "Stage-based: when size or stage is what matters",
                               103: "STAGE",
                               105: (["Fate depends on size or developmental stage, not age",
                                      "Age is often unknown, and individuals can stay in a stage for years",
                                      "Examples: sea turtles, trees, fish, insects"], 18)},
    pictures={101: I + "turtle_stages.jpg"},
    notes="""
(1.5 min) Leatherback turtle (*Dermochelys coriacea*): hatchling and adult. We cannot tell how old a turtle is by looking at it, and what matters for its survival is how big it is.
Neal (2004, section 15.3) gives more examples: insects pass through stages (egg, larva, pupa, adult) whose duration depends on food, temperature and moisture; in many plants and fish reproduction depends more on size than on age, so a small plant may be a fast-growing juvenile or a stunted "adult"; some organisms reproduce both sexually and clonally, producing offspring in different stages.
Photos carried over from the previous version of the lecture (sources and licences to confirm: images/CREDITS.md).
""")

# ================================================================ Section 02
new("Section", {102: "02", 103: "Age-based models"},
    notes="(10 s) We start with age, because it links directly to last lecture's life tables.")

s = new("Title Only", {100: "Start with a life table"},
        notes="""
(1.5 min) The same cohort and the same notation as last lecture's life table (Neal's symbols plus p(x)). x is age at the start of the interval; S(x) the number of females alive at age x; l(x) = S(x)/S(0) survivorship; m(x) female offspring per female of age x; p(x) = l(x+1)/l(x), the probability of surviving from age x to x+1.
Point to the two columns we need: p(x) and m(x). S(x) and l(x) are how we got p(x).
Other books differ: Gotelli calls m(x) b(x) and p(x) g(x) (the previous version of this lecture used Gotelli's letters).
""")
table(s, ["Age, *x*", "S~x~", "l~x~", "m~x~", "p~x~ = l~x+1~/l~x~"],
      [["0", "500", "1.0", "0", "0.80"], ["1", "400", "0.8", "2", "0.50"], ["2", "200", "0.4", "3", "0.25"],
       ["3", "50", "0.1", "1", "0.00"], ["4", "0", "0.0", "0", "–"]],
      col_w=[1.2, 1.0, 1.0, 1.0, 2.3], top=2.3, row_h=0.58, size=18)
textbox(s, ["**m**~x~: female offspring per female aged *x*", "",
            "**p**~x~: chance of surviving from age *x* to *x* + 1", "",
            "These two columns are all the matrix needs"], 8.2, 2.5, 4.3, 3.5, size=19, color=INK)

new("Title & Photo", {100: "Survival moves individuals one age class along",
                      102: "Post-breeding census: individuals are counted just after breeding, so newborns are age 0."},
    pictures={101: F + "01a_leslie_survival.png"},
    notes="""
(1.5 min) Each circle is an age class at the census; each arrow is what happens over one projection interval (one year). The arrow values are p(x) from the life table.
Census timing matters for the next step. We assume a birth-pulse population (breeding once a year) counted just after breeding, as in Neal section 15.3.1. Nobody survives beyond age 3 here (p(3) = 0).
Figure: scripts/01_leslie_diagram.R (analytical).
""")

new("Key Concept", {102: "Fertility = survival × births",
                    103: "An individual counted at age *x* adds offspring to the next census only if it survives the year, p~x~, and then breeds, m~x+1~.",
                    105: "F~x~ = p~x~ × m~x+1~.  Age 0: 0.80 × 2 = 1.6.  Age 1: 0.50 × 3 = 1.5.  Age 2: 0.25 × 1 = 0.25.  Common mistake: putting m~x~ straight into the matrix."},
    notes="""
(2 min) This is Neal's Eqn 15.8, F~i~ = m~i~P~i~, written with ages. Neal numbers the age classes i = 1, 2, 3 (class i runs from age i − 1 to i), so his P~1~ is our p~0~ and his F~1~ = m~1~P~1~ is our F~0~ = p~0~m~1~: the same numbers, shifted by one in the label.
Walk through age 0: a newborn counted this census must survive its first year (0.8) to breed at age 1 (2 offspring), so on average it contributes 0.8 × 2 = 1.6 newborns to next year's census.
The previous version of this lecture put m(x) (0, 2, 3, 1) straight into the top row. That delays every birth by a year and ignores parental survival: λ = 1.55 instead of the correct 2.17 (scripts/00_check_numbers.R). The practical "Pre- and post-reproduction census" (""" + URL_CENSUS + """) shows a related slip: using the wrong survival term in the fertility for the census type.
With a pre-breeding census the bookkeeping changes (fertility is multiplied by the offspring's first-year survival instead), but done correctly λ is the same.
""")

new("Title & Photo", {100: "Add reproduction: arrows back to age 0",
                      102: "Clay: fertility, F~x~ = p~x~ m~x+1~. Green: survival, p~x~. Age-3 individuals have bred and do not survive."},
    pictures={101: F + "01b_leslie_full.png"},
    notes="""
(1 min) Every age class that can produce newborns by the next census gets an arrow back to age 0, including age 0 itself (a loop): newborns this year can be parents of next year's newborns.
Age 3 has no outgoing arrows: these individuals have just bred (their newborns came in through the 0.25 arrow from age 2, as F(2) = p(2) m(3)) and will not survive another year.
Figure: scripts/01_leslie_diagram.R (analytical).
""")

s = new("Title Only", {100: "The Leslie matrix"},
        notes="""
(2 min) Named after P. H. Leslie (1945); Bernardelli and Lewis developed the same idea independently in the 1940s (Neal 15.3).
Rule for every matrix model (also in the practical): the arrow's starting stage gives the COLUMN, the stage it points to gives the ROW. So the column for age 0 reads: 1.6 newborns and 0.8 survivors to age 1.
In an age-based (Leslie) matrix, fertilities sit in the top row and survival on the sub-diagonal (Neal's P~i~: the p(x) values from the life table, as promised in lecture 05); everything else is zero, because nobody can stay the same age or get younger.
The last column is all zeros (age 3 contributes nothing). We could drop it without changing λ; it is kept here to match the diagram.
Check: 6 arrows in the diagram, 6 non-zero cells. λ = 2.17 (very fast growth: an illustrative life table). This is the same cohort as last lecture, and the same answer: λ = e^r^ with r = 0.776 from the Euler–Lotka equation. The raw-m(x) matrix of the previous version (λ = 1.55) disagreed with last lecture's Euler–Lotka result.
""")
C, Fo, Z = CLAY, FOREST, None
matrix(s, [["1.6", "1.5", "0.25", "0"], ["0.8", "0", "0", "0"], ["0", "0.5", "0", "0"], ["0", "0", "0.25", "0"]],
       left=2.3, top=2.9, cw=1.25, rh=0.75, size=26,
       colors=[[C, C, C, Z], [Fo, Z, Z, Z], [Z, Fo, Z, Z], [Z, Z, Fo, Z]],
       col_labels=["Age 0", "Age 1", "Age 2", "Age 3"], row_labels=["Age 0", "Age 1", "Age 2", "Age 3"],
       name="**A** =", name_size=30)
textbox(s, "FROM (this year) →", 2.3, 2.05, 5.0, 0.4, size=14, color=GREY, align=PP_ALIGN.CENTER)
textbox(s, ["**Columns:** age this year (from)", "**Rows:** age next year (to)", "",
            "Top row: fertilities, F~x~", "Sub-diagonal: survival, p~x~", "(Neal: F~i~ and P~i~)", "",
            "One arrow = one non-zero cell"], 8.6, 2.5, 4.0, 3.8, size=19, color=INK)

s = new("Title Only", {100: "Generation time: three answers"},
        notes="""
(1 min) Last lecture promised other definitions of generation time. Same cohort as before (R~0~ = Σ l(x)m(x) = 2.9; λ = 2.17 from the Leslie matrix).
1. Cohort generation time (lecture 05): T~c~ = Σ x l(x)m(x) / R~0~ = 4.3 / 2.9 = 1.48 years, the mean age at which a cohort's females have their offspring.
2. The time it takes the population to grow by a factor R~0~ at its long-run rate: T = ln R~0~ / ln λ = 1.065 / 0.776 = 1.37 years. (This is the matrix-model definition; R~0~ can also be calculated from the matrix itself.)
3. The mean age of the mothers of this year's newborns, once the population has its stable age distribution: Ā = Σ x λ^−x^ l(x)m(x) = 1.27 years.
Why they differ: in a growing population each cohort is larger than the one before, so young mothers are over-represented among this year's parents. In a stationary population (λ = 1) the three coincide. This population is growing fast, so the differences are large; for most real populations they are small.
Values: scripts/00_check_numbers.R.
""")
table(s, ["Definition", "Formula", "Value"],
      [["Cohort generation time (lecture 05)", "Σ *x* l~x~m~x~ / R~0~", "1.48 years"],
       ["Time to grow by R~0~ at rate λ", "ln R~0~ / ln λ", "1.37 years"],
       ["Mean age of mothers, stable population", "Σ *x* λ^−x^ l~x~m~x~", "1.27 years"]],
      col_w=[5.2, 3.6, 1.9], top=2.45, row_h=0.62, size=18, first_col_bold=False)
textbox(s, ["In a growing population young mothers are over-represented, so the answers differ.",
            "In a stationary population (λ = 1) they are the same."], 0.8, 5.0, 11.5, 1.2, size=19, color=INK)

# ================================================================ Section 03
new("Section", {102: "03", 103: "Stage-based models"},
    notes="(10 s) The same machinery, but the classes are stages, and individuals can stay put.")

new("Title & Photo", {100: "Individuals grow from stage to stage",
                      102: "Each arrow: probability of moving to the next stage in one year. Illustrative values."},
    pictures={101: F + "02a_stage_progression.png"},
    notes="""
(1 min) A simple three-stage life cycle. 26% of juveniles survive and become adults each year; 12% of adults survive and become post-reproductive (they no longer breed).
Figure: scripts/02_stage_diagram.R (illustrative values from the previous version of the lecture).
""")

new("Title & Photo", {100: "Adults produce juveniles",
                      102: "Fertility (the practical says fecundity): 4.55 new juveniles per adult per year."},
    pictures={101: F + "02b_stage_reproduction.png"},
    notes="(30 s) Only adults reproduce in this model. Vocabulary: demographers call this matrix entry fertility (Neal: fertility coefficient); the practical uses the everyday word fecundity for the same thing. As in the age-based model, the fertility value is per individual per projection interval, and includes whatever survival is needed for offspring to be counted at the next census.")

new("Title & Photo", {100: "Individuals can also stay in a stage",
                      102: "Staying (stasis): a juvenile can remain a juvenile for several years. Age-based models cannot do this."},
    pictures={101: F + "02c_stage_full.png"},
    notes="""
(1 min) The loops are survival without moving on: 14% of juveniles survive and stay juvenile, and so on.
Survival out of a stage is the sum of its non-fertility arrows, and cannot exceed 1: juveniles 0.14 + 0.26 = 0.40, adults 0.32 + 0.12 = 0.44, post-reproductive 0.18. Fertility does not count towards this total, because newborns are added rather than being a fate of the parent.
Other possible transitions (not here): skipping a stage, or moving backwards (shrinking, e.g. plants after damage, or "regression").
""")

new("Title, Bullets & Photo", {100: "Write the matrix for this life cycle",
                               103: "YOUR TURN · 2 MIN",
                               105: (["Pairs: fill in a 3 × 3 matrix", "Columns = from, rows = to",
                                      "One arrow = one non-zero cell", "Leave the rest as 0"], 18)},
    pictures={101: F + "02d_stage_full_takeaway.png"},
    notes="""
(2 min) Give pairs about 90 seconds. Ask: how many non-zero cells should you have? (7, one per arrow.)
Common errors to look for: transposing the matrix (rows = from), and putting 4.55 in the adult row instead of the juvenile row.
""")

s = new("Title Only", {100: "A stage-classified (Lefkovitch) matrix"},
        notes="""
(1.5 min) Stage-classified matrices are often called Lefkovitch matrices (after L. P. Lefkovitch, 1965), to distinguish them from age-based Leslie matrices. The algebra is the same; only the pattern of non-zero cells differs: here there are values on the diagonal (stasis).
Neal's Fig. 15.1 / Table 15.3 show the same structure for a plant (dormant seeds, immature, mature) with both sexual and clonal reproduction.
Interesting aside for discussion: the post-reproductive stage never contributes to future population growth (no arrows leave it except its own loop), so in this model it does not affect λ at all. We will meet the same thing in killer whales. (The practical's example organism has a 'senescent' stage that does still reproduce: different life histories, same rules.)
""")
matrix(s, [["0.14", "4.55", "0"], ["0.26", "0.32", "0"], ["0", "0.12", "0.18"]],
       left=2.6, top=2.9, cw=1.4, rh=0.85, size=28,
       colors=[[Fo, C, Z], [Fo, Fo, Z], [Z, Fo, Fo]],
       col_labels=["Juvenile", "Adult", "Post-repro."], row_labels=["Juvenile", "Adult", "Post-repro."],
       name="**A** =", name_size=30)
textbox(s, "FROM (this year) →", 2.6, 2.05, 4.2, 0.4, size=14, color=GREY, align=PP_ALIGN.CENTER)
textbox(s, ["7 arrows, 7 non-zero cells", "",
            "Survival out of a stage ≤ 1:", "juveniles 0.14 + 0.26 = 0.40", "",
            "Fertility can be greater than 1", "",
            "Same rules as the Leslie matrix"], 8.6, 2.4, 4.0, 4.0, size=19, color=INK)

new("Title, Bullets & Photo", {100: "Matrix models fit almost any life cycle",
                               103: "USES",
                               105: (["Project population size and structure", "Find which transitions matter most (elasticity)",
                                      "Run *in silico* experiments"], 18)},
    pictures={101: I + "mpm_examples.jpg"},
    notes="""
(1.5 min) A matrix population model (MPM) describes a life cycle mathematically. It works for almost any life cycle: animals and plants, age-based, stage-based or a mixture (e.g. stages for juveniles, ages for adults).
To be complete it needs a projection interval (the time step, usually one year), and a census time relative to breeding.
Thousands of published MPMs are collected in the COMPADRE (plants) and COMADRE (animals) databases.
Photos (sea turtle, tree ferns, Atlantic cod) carried over from the previous version of the lecture (sources to confirm: images/CREDITS.md).
""")

# ================================================================ Section 04
new("Section", {102: "04", 103: "Projecting the population"},
    notes="(10 s) What will the population size and structure be next year, and in the long run?")

new("Key Concept", {102: "n~t+1~ = **A** n~t~",
                    103: "Multiply the projection matrix (also called a transition matrix) by this year's population vector to get next year's. It is geometric growth, n~t+1~ = λn~t~, with the single number λ replaced by a matrix.",
                    105: "Order matters: **A** × n~t~, matrix first. Each row of **A** says where next year's individuals in that stage come from."},
    notes="""
(1.5 min) n~t~ is a column vector: the number of individuals in each stage at time t.
Link back to the geometric growth lecture: there, a single λ multiplied a single number. Here a matrix multiplies a vector. We will see that in the long run the two become the same thing.
Matrix multiplication is not commutative: n~t~ × A is not even defined for a column vector.
""")

def projection_slide(title, results, worked, notes):
    s = new("Title Only", {100: title}, notes=notes)
    matrix(s, [["0.14", "4.55", "0"], ["0.26", "0.32", "0"], ["0", "0.12", "0.18"]],
           left=1.4, top=2.6, cw=1.05, rh=0.7, size=22, colors=[[C, C, Z], [Fo, Fo, Z], [Z, Fo, Fo]],
           col_labels=["Juv.", "Adult", "Post-r."])
    textbox(s, "×", 4.75, 3.0, 0.5, 0.8, size=30, color=INK, align=PP_ALIGN.CENTER)
    matrix(s, [["100"], ["50"], ["20"]], left=5.4, top=2.6, cw=0.95, rh=0.7, size=22,
           colors=[[INK], [INK], [INK]], col_labels=["n~t~"], label_size=16)
    textbox(s, "=", 6.55, 3.0, 0.5, 0.8, size=30, color=INK, align=PP_ALIGN.CENTER)
    matrix(s, [[r] for r in results], left=7.2, top=2.6, cw=1.1, rh=0.7, size=22,
           colors=[[GOLD if r == "?" else INK] for r in results], bold=[[True]] * 3,
           col_labels=["n~t+1~"], row_labels=["Juveniles", "Adults", "Post-repro."], label_size=16)
    tb = textbox(s, worked, 1.0, 5.05, 11.3, 1.5, size=18, color=INK)
    return s

projection_slide("Projecting one year: your turn", ["241.5", "?", "?"],
                 ["**Juveniles:** 0.14 × 100 + 4.55 × 50 + 0 × 20 = 14 + 227.5 = 241.5",
                  "**Your turn:** use the adult row and the post-reproductive row in the same way."],
                 """
(1.5 min) Work through the juvenile row together: go along the row of A and down the vector, multiply pairs and add. Most next-year juveniles (227.5 of 241.5) are newborns from adults.
Give students a minute for the other two rows before the next slide.
""")
s = projection_slide("Projecting one year: the answer", ["241.5", "42", "9.6"],
                 ["**Juveniles:** 0.14 × 100 + 4.55 × 50 + 0 × 20 = 241.5",
                  "**Adults:** 0.26 × 100 + 0.32 × 50 + 0 × 20 = 26 + 16 = 42",
                  "**Post-reproductive:** 0 × 100 + 0.12 × 50 + 0.18 × 20 = 6 + 3.6 = 9.6"],
                 """
(1 min) Each row tells us how many individuals enter that stage next year, and from where. Total: 170 now, 293.1 next year. Repeat the multiplication with the new vector to project further: R does this with `A %*% n` (the practical uses popdemo::project()).
Fractions of individuals are fine: the model tracks expected numbers.
""")
textbox(s, ["In R:", "`A %*% n`"], 10.1, 2.6, 2.6, 1.0, size=18, color=GREY)

new("Figure & Bullets", {100: "Example: resident killer whales",
                         102: (["*Orcinus orca*, coastal British Columbia and Washington State",
                                "Every whale recognised from photographs, followed since 1973",
                                "Brault & Caswell (1993): females only, four stages",
                                "Females can live for decades after they stop breeding"], 18)},
    keep=(101,),
    notes="""
(1 min) Context before the model. Resident (fish-eating) killer whales live in stable family groups (pods). Bigg and colleagues began a photo-identification census in 1973: every individual is recognised from its dorsal fin and saddle patch, so births and deaths are known almost exactly. Brault & Caswell (1993, *Ecology* 74: 1444–1454) used the 1973–1987 data for 18 pods of the northern and southern resident populations.
The model follows females only, in four stages: yearlings, juveniles, mature (reproductive) females and post-reproductive females. It is the worked example in Caswell's textbook (Caswell 2001, Example 5.1).
PHOTO PLACEHOLDER: drop in a photograph of resident killer whales (e.g. from Wikimedia Commons; downloads were blocked when the deck was built) and add the credit here.
""")

new("Title & Photo", {100: "The killer whale life cycle",
                      102: "Brault & Caswell (1993). Survival per year: yearling 0.98, juvenile 0.98, mature 0.999, post-reproductive 0.98."},
    pictures={101: F + "03_whale_diagram.png"},
    notes="""
(1 min) Yearlings all become juveniles (0.98 survive). Juveniles stay juvenile (0.91) or mature (0.07); mature females stay mature (0.95) or become post-reproductive (0.05); post-reproductive females stay post-reproductive (0.98). Annual survival is very high in every stage.
Fertility: 0.11 female yearlings per mature female per year, roughly one daughter every nine years. The tiny juvenile fertility (0.004) is a bookkeeping detail: a few juveniles mature and give birth within the same year.
Post-reproductive females have no arrow back to yearlings. Ask: what will that mean for their effect on λ? (Answer coming in the hidden elasticity slide, or ask it now.)
Figure: scripts/03_whale_diagram.R (empirical matrix).
""")

s = new("Title Only", {100: "From diagram to matrix"}, hidden=True,
        notes="""
(1.5 min) HIDDEN by default to keep to 45 minutes. Left: the matrix written with Neal's symbols (Fig. 15.1, Table 15.3): P = probability of surviving and staying in a stage, G = probability of surviving and growing to the next stage, F = fertility. So G~Y~ is yearlings surviving to become juveniles (column = from, row = to), P~J~ is juveniles staying juvenile, and F~M~ is the fertility of mature females. Crouse et al. use the same letters for the loggerhead. Right: the published values (Brault & Caswell 1993, Eqn 2).
Ask: why is the last column all zeros except the diagonal? (Post-reproductive females only survive; they produce nobody and move nowhere.)
""")
sym = [["0", "F~J~", "F~M~", "0"], ["G~Y~", "P~J~", "0", "0"], ["0", "G~J~", "P~M~", "0"], ["0", "0", "G~M~", "P~P~"]]
num = [["0", "0.0043", "0.1132", "0"], ["0.9775", "0.9111", "0", "0"], ["0", "0.0736", "0.9534", "0"], ["0", "0", "0.0452", "0.9804"]]
colp = [[Z, C, C, Z], [Fo, Fo, Z, Z], [Z, Fo, Fo, Z], [Z, Z, Fo, Fo]]
matrix(s, sym, left=1.7, top=2.9, cw=1.0, rh=0.72, size=22, colors=colp,
       col_labels=["Y", "J", "M", "P"], name="**A** =", name_size=28)
matrix(s, num, left=7.0, top=2.9, cw=1.2, rh=0.72, size=20, colors=colp,
       col_labels=["Y", "J", "M", "P"], row_labels=["Y", "J", "M", "P"], name="=", name_size=28)

new("Title & Photo", {100: "Different starts, same long-run growth",
                      102: "Log scale. Straight, parallel lines: every stage grows at the same rate, λ. From 100 yearlings this takes about 40 years."},
    pictures={101: F + "04_whale_projection.png"},
    notes="""
(2 min) Left: starting with 100 yearlings only; right: 25 in every stage. Early on the two look completely different (transient dynamics): from yearlings, there are no mature females for several years, so births collapse and the yearling line crashes. Eventually the lines become straight and parallel on the log scale: every stage grows at the same constant rate, so the proportions in each stage stop changing. (This is practical question 1: straight on a log scale means exponential growth; parallel means the same rate in every stage.)
How long that takes depends on the life history. Here it is slow: within 5% of the stable structure after about 11 years from the even start, but about 40 years from yearlings only (scripts/00_check_numbers.R). Long-lived species with high survival take decades to settle, so transients matter for real conservation problems, such as recovery after a die-off or a reintroduction.
Figure: scripts/04_whale_projection.R (analytical, no randomness).
""")

new("Big Fact", {100: "λ = 1.025",
                 102: "The dominant eigenvalue of **A**: in the long run the population grows by 2.5% per year, whatever the starting population."},
    notes="""
(1 min) Once the stage structure has settled, **A**n~t~ = λn~t~: the matrix model behaves exactly like geometric growth from lecture 3. λ is the largest (dominant) eigenvalue of **A**; R calculates it (eigen(), or popdemo::eigs() in the practical).
Value: 1.0254, as published by Brault & Caswell (1993), who also report a bootstrap 90% confidence interval of 1.018 to 1.032 (scripts/00_check_numbers.R reproduces 1.0254). λ > 1: growth; λ = 1: stable; λ < 1: decline.
In R (as in the practical): popdemo::eigs(A, "lambda").
A projection is not a forecast: λ describes the 1973–1987 vital rates. Since then the Southern Resident population has declined; it was listed as endangered in Canada in 2001 and under the US Endangered Species Act in 2005 (74 whales in July 2025; US Marine Mammal Commission).
""")

new("Title, Bullets & Photo", {100: "The stable stage distribution",
                               103: "TAKEAWAY",
                               105: (["Once transients die out, the share in each stage stays fixed", "It is the dominant (right) eigenvector of **A**",
                                      "Only 4% yearlings; a third post-reproductive",
                                      "In R: `eigs(A, \"ss\")`"], 17)},
    pictures={101: F + "05_whale_ssd.png"},
    notes="""
(1.5 min) Stable stage distribution (SSD): 4% yearlings, 32% juveniles, 32% mature females, 32% post-reproductive females. This is what both projections converged to. It is the eigenvector that goes with λ, scaled to sum to 1. Like λ, it depends only on the matrix, not on the starting population.
Why so few yearlings? Females have about one daughter every nine years, and the yearling stage lasts only one year. Why so many post-reproductive females? They survive well (0.98 per year) and stay in that stage for decades.
Practical use: if an observed population is far from its SSD, expect transient dynamics; the stable distribution is also a check on whether the model resembles the real population.
(The same calculation for the age-based model gives the stable age distribution; Neal 15.3.2.)
Figure: scripts/05_whale_ssd.R.
""")

s = new("Title Only", {100: "Stochastic: a new matrix each year"},
        notes="""
(2 min) So far the matrix was the same every year. Real environments vary: good and bad years change survival and breeding. In a stochastic projection we draw a matrix at random each year (e.g. from matrices estimated in good, average and bad years), and multiply as before.
What it gives: the stochastic growth rate λ~s~ (the long-run average growth rate on a log scale), which is usually lower than the λ of the average matrix, because bad years hurt more than good years help; and the probability of falling below a threshold within a given time (extinction risk), from many replicate runs.
Links to the practicals "Stochastic population growth" (""" + URL_STOCH + """) and "Matrix Population Models: projection and simulation".
Correction from the previous version: the slide showed n~t~ × A; it is **A**~t~ × n~t~.
Icons carried over from the previous version.
""")
textbox(s, "n~t+1~ = **A**~t~ n~t~", 0.8, 2.25, 6.0, 1.0, size=40, color=FOREST)
for k, (icon, lab) in enumerate([("icon_sun.png", "Good year"), ("icon_cloud_sun.png", "Average year"),
                                 ("icon_rain.png", "Bad year")]):
    x = 0.9 + k * 2.0
    picture(s, I + "source/" + icon, x + 0.35, 3.75, width=1.0)
    textbox(s, ["**A**~" + lab.split()[0].lower() + "~", lab], x, 4.85, 1.7, 0.9, size=16, color=INK,
            align=PP_ALIGN.CENTER)
textbox(s, "One drawn at random each year", 0.9, 5.95, 5.7, 0.4, size=16, color=GREY, align=PP_ALIGN.CENTER)
textbox(s, ["Gives the **stochastic growth rate**, λ~s~", "usually lower than λ of the average matrix", "",
            "and the **risk of extinction** (or of falling below a threshold), from many runs"],
        7.6, 2.6, 4.9, 3.5, size=19, color=INK)

# ================================================================ Section 05
new("Section", {102: "05", 103: "Which transitions matter?"},
    notes="(10 s) Sensitivity, elasticity and in silico experiments.")

new("Title, Bullets & Photo", {100: "Which parts of the life cycle matter most?",
                               103: "QUESTIONS",
                               105: (["Where should conservation or management effort go?",
                                      "Which transitions are likely to be under the strongest selection?"], 18)},
    pictures={101: I + "loggerhead_takeaway.jpg"},
    notes="""
(1 min) λ depends on every element of the matrix, but not equally. If we could change one transition a little, which would change λ most?
The same question has an evolutionary reading: λ is a measure of fitness (when growth is density-independent), so transitions to which λ is most sensitive are where selection is expected to act most strongly (next lecture: life-history evolution).
Photo: loggerhead turtle (*Caretta caretta*), Wikimedia Commons "Loggerhead sea turtle.jpg" (author and licence to confirm: images/CREDITS.md).
""")

new("Two Columns", {100: "Sensitivity and elasticity",
                    101: "SENSITIVITY",
                    103: (["Change in λ after a small **absolute** change in one element (e.g. +0.01)",
                           "Juvenile survival 0.20 → 0.21: λ + 0.005", "Adult survival 0.80 → 0.81: λ + 0.010",
                           "Hard to compare survival (≤ 1) with fertility (can be ≫ 1)"], 17),
                    104: "ELASTICITY",
                    106: (["**Proportional** change in λ after a small **proportional** change (e.g. +5%)",
                           "Juvenile survival +5%: λ +0.5% (elasticity 0.10)", "Adult survival +5%: λ +4% (elasticity 0.80)",
                           "Comparable across all elements; they sum to 1",
                           "In R: `popdemo::elas(A)`"], 17)},
    notes="""
(2 min) A hypothetical example, with λ = 1.00 now.
Sensitivity: give each element the same absolute nudge (+0.01). Juvenile survival 0.20 → 0.21 raises λ by 0.005 (sensitivity 0.5); adult survival 0.80 → 0.81 raises it by 0.010 (sensitivity 1.0).
But +0.01 is a 5% increase for juveniles and only 1.25% for adults. Elasticity instead gives each element the same percentage nudge. +5%: juveniles 0.20 → 0.21 gives λ +0.5%; adults 0.80 → 0.84 gives λ +4%. Elasticity = % change in λ / % change in the element: 0.5/5 = 0.10 and 4/5 = 0.80. (Equivalently, elasticity = sensitivity × element / λ: 0.5 × 0.20 / 1 = 0.10; 1.0 × 0.80 / 1 = 0.80.)
Because elasticities of all matrix elements sum to 1, each can be read as that transition's share of the influence on λ. This summing to 1 is a property of λ elasticities; it is not how they are calculated.
A high elasticity says λ responds strongly to that transition; it does not say the transition is easy or cheap to change.
""")

new("Title, Bullets & Photo", {100: "Killer whale: adult survival matters most",
                               103: "TAKEAWAY",
                               105: (["Mature females surviving: 0.54; juveniles surviving: 0.34",
                                      "Fertility: 0.04 in total",
                                      "Post-reproductive females: 0"], 18)},
    pictures={101: F + "06_whale_elasticity.png"}, hidden=True,
    notes="""
(1.5 min) HIDDEN by default to keep to 45 minutes; show it if time allows, as a gentle 4 × 4 introduction before the 7 × 7 loggerhead matrix. Elasticities of λ to each element of the killer whale matrix, laid out like the matrix (columns from, rows to). Blank cells are zero in the matrix.
Survival dominates: mature females staying alive and mature (0.54) and juveniles staying alive (0.34). Fertility contributes only 0.04. This is what Brault & Caswell (1993) concluded, and is typical of long-lived species: because survival is close to 1, even small proportional declines in survival have large effects on λ.
Post-reproductive females have zero elasticity: in this model they contribute nothing to future growth (zero reproductive value). Ask: does that mean they do not matter? Brault & Caswell themselves point out their likely social roles (helping to raise others' young, remembering where to find food); a model that tracks only births and deaths cannot see this.
Note that yearling → juvenile, juvenile → mature and mature fertility all come out at about 0.04: they lie on the same loop through the life cycle (yearling → juvenile → mature → back to yearling), and every element of a loop gets the same share of elasticity (van Groenendael et al. 1994). Yearling → juvenile is slightly higher (0.043) because it also lies on the small loop through juvenile fertility.
Figure: scripts/06_whale_elasticity.R (analytical; elasticities sum to 1).
""")

new("Figure & Bullets", {100: "Loggerhead turtles: a declining population",
                         102: (["*Caretta caretta*, nesting on beaches in the south-east USA",
                                "Crouse, Crowder & Caswell (1987): a 7-stage model from a life table",
                                "λ = 0.945: declining by about 5% per year",
                                "Conservation focused on protecting eggs on nesting beaches"], 18)},
    pictures={101: I + "loggerhead_half.jpg"},
    notes="""
(1.5 min) Context before the data. Crouse, D. T., Crowder, L. B. & Caswell, H. (1987) A stage-based population model for loggerhead sea turtles and implications for conservation. *Ecology* 68: 1412–1423. They built a stage-classified matrix from Frazer's (1983) life table for loggerheads on Little Cumberland Island, Georgia.
Seven stages: eggs and hatchlings, small juveniles, large juveniles, subadults, novice breeders, first-year remigrants, mature breeders. Females lay many eggs (fertility 127 for novice breeders, 80 for mature breeders) but mature only after about 20 years or more.
λ = 0.945 (r = ln λ = −0.057). Management at the time concentrated on the most visible mortality: eggs and hatchlings on the beaches.
""")

new("Check Question (2 options)", {101: ("You have limited money to save the loggerheads. Which would you choose?", 30),
                                   105: "Protect nests and hatchlings on the beaches",
                                   109: "Reduce deaths of large juveniles in fishing nets"},
    notes="""
(1.5 min) Show of hands. Many choose A: eggs and hatchlings suffer the highest mortality, and beaches are where people see turtles.
Then: what would elasticity analysis tell us? Next slide.
""")

# Loggerhead projection matrix and its elasticity matrix, side by side (values from
# scripts/00_check_numbers.R: Crouse et al. 1987 Table 4, elasticities recalculated)
CR_CODES = ["E", "SJ", "LJ", "SA", "NB", "R", "MB"]
CR_A = [["0", "0", "0", "0", "127", "4", "80"],
        ["0.67", "0.74", "0", "0", "0", "0", "0"],
        ["0", "0.05", "0.66", "0", "0", "0", "0"],
        ["0", "0", "0.01", "0.69", "0", "0", "0"],
        ["0", "0", "0", "0.05", "0", "0", "0"],
        ["0", "0", "0", "0", "0.81", "0", "0"],
        ["0", "0", "0", "0", "0", "0.81", "0.81"]]
CR_E = [["0", "0", "0", "0", "0.01", "0", "0.04"],
        ["0.05", "0.18", "0", "0", "0", "0", "0"],
        ["0", "0.05", "0.12", "0", "0", "0", "0"],
        ["0", "0", "0.05", "0.14", "0", "0", "0"],
        ["0", "0", "0", "0.05", "0", "0", "0"],
        ["0", "0", "0", "0", "0.04", "0", "0"],
        ["0", "0", "0", "0", "0", "0.04", "0.23"]]

def crouse_matrices_slide(title, reveal, caption, notes):
    s = new("Title Only", {100: title}, notes=notes)
    colA = [[None if v == "0" else (GOLD if reveal and v in ("127", "80") else INK) for v in row] for row in CR_A]
    colE = [[None if v == "0" else (CLAY if reveal and float(v) >= 0.10 else INK) for v in row] for row in CR_E]
    boldA = [[reveal and v in ("127", "80") for v in row] for row in CR_A]
    boldE = [[reveal and v != "0" and float(v) >= 0.10 for v in row] for row in CR_E]
    matrix(s, CR_A, left=1.35, top=2.85, cw=0.64, rh=0.46, size=16, colors=colA, bold=boldA,
           col_labels=CR_CODES, name="**A** =", name_size=24, label_size=14)
    matrix(s, CR_E, left=7.45, top=2.85, cw=0.64, rh=0.46, size=16, colors=colE, bold=boldE,
           col_labels=CR_CODES, row_labels=CR_CODES, name="**E** =", name_size=24, label_size=14)
    textbox(s, "Projection matrix (from → columns, to → rows)", 0.8, 1.95, 5.6, 0.4, size=14, color=GREY,
            align=PP_ALIGN.CENTER)
    textbox(s, "Elasticity of λ to each element (sums to 1)", 6.9, 1.95, 5.6, 0.4, size=14, color=GREY,
            align=PP_ALIGN.CENTER)
    textbox(s, caption, 0.8, 6.2, 11.7, 0.5, size=15, color=INK if reveal else GREY, align=PP_ALIGN.CENTER)
    return s

crouse_matrices_slide("Where are the big numbers?", False,
    "E eggs & hatchlings · SJ, LJ small, large juveniles · SA subadults · NB novice breeders · R remigrants · MB mature",
    """
(1.5 min) Left: the loggerhead projection matrix A (Crouse et al. 1987, rounded to 2 decimals). Right: the elasticity matrix E, one elasticity per matrix element, in the same positions. Zeros in A always have zero elasticity.
Ask students, in pairs, for 30 seconds: which three elements of E are biggest? And where are the biggest numbers in A? Do they match?
Stage codes: E eggs and hatchlings, SJ small juveniles, LJ large juveniles, SA subadults, NB novice breeders, R first-year remigrants, MB mature breeders.
""")

crouse_matrices_slide("Big numbers in A are not big in E", True,
    "Fecundity is huge in **A** (127, 80) but tiny in **E** (0.01, 0.04). Survival of juveniles and adults dominates **E**.",
    """
(1 min) Reveal. Largest elasticities: mature breeders staying alive (0.23), small juveniles staying (0.18), subadults staying (0.14) and large juveniles staying (0.12); together 0.67 of the total of 1. The growth transitions on the sub-diagonal are 0.05 each (eggs to small juveniles included).
The biggest numbers in A are the fecundities (127 and 80 eggs), yet their elasticities are only 0.01 and 0.04: a 1% increase in egg production gives almost no change in λ. Elasticity is about proportional change, so a large number in A does not mean a large influence on λ.
This is the same message as the bar chart on the next slide, which sums each column of E by stage and type of transition.
Values: scripts/00_check_numbers.R.
""")

new("Title, Bullets & Photo", {100: "Elasticity is low for eggs",
                               103: "TAKEAWAY",
                               105: (["Eggs and hatchlings: 0.05", "Adult survival is the largest single element (0.23)",
                                      "Juvenile and subadult survival together: 0.59"], 18)},
    pictures={101: F + "07_crouse_elasticity.png"},
    notes="""
(1 min) The elasticity matrix summarised by stage. Elasticities from the published Crouse et al. (1987) matrix (recalculated: scripts/07_crouse_elasticity.R; λ = 0.945 matches the paper). Each bar is one stage, split into staying, growing and reproducing.
Eggs and hatchlings: 0.05. Small juveniles 0.23, large juveniles 0.17, subadults 0.19 (0.59 together); mature breeders 0.27, the most elastic single stage. The largest single element is mature breeders staying alive (0.23), then small juveniles staying (0.18). All three breeding stages together: 0.36. Fecundity elasticities are tiny (0.01 to 0.04).
So "what matters most" depends on how you group: element by element and stage by stage it is adult survival; summed over the long juvenile period (about 20 years in three stages) it is juvenile and subadult survival.
Their abstract: "Current management practices appear to be focused on the least responsive life stage, eggs on nesting beaches. Alternative protection efforts for juvenile loggerheads, such as using turtle excluder devices (TEDs), may be far more effective."
This replaces the screenshots of the paper in the previous version of the lecture.
""")

new("Title, Bullets & Photo", {100: "*In silico*: halve the mortality of one stage",
                               103: "IN SILICO",
                               105: (["Even 100% egg survival: λ = 0.965, still declining",
                                      "Halving large-juvenile mortality: λ = 1.04, growing",
                                      "Basis for turtle excluder devices (TEDs)",
                                      "In R: edit a copy, `A_mod`; compare λ"], 17)},
    pictures={101: F + "08_crouse_in_silico.png"},
    notes="""
(2 min) An in silico experiment: change the matrix to represent a management action and ask what happens to λ (or other properties). Here, as in Crouse et al. (1987): halve the annual mortality of one stage at a time (e.g. large-juvenile survival 0.676 → 0.838), and recalculate that stage's "stay" and "grow" entries from its new survival and its duration in years.
Results (scripts/08_crouse_in_silico.R): eggs and hatchlings 0.956; small juveniles 0.994; large juveniles 1.044; subadults 0.998; novice breeders and remigrants 0.949; mature breeders 0.979. Even making egg survival perfect (gold circle) only gives 0.965: the population still declines.
Halving large-juvenile mortality is the only single change that turns decline into growth, and large juveniles are the stage drowned in shrimp trawls. TEDs became required in US shrimp trawls from 1989 onwards; Crowder et al. (1994, *Ecological Applications*) extended the analysis to predict their effect.
Why does this ranking differ from the elasticities, where adult survival was the largest single element? Elasticity compares small, equal proportional changes. Halving mortality is a large change, and not an equal one: large juveniles die at 32% a year and mature breeders at 19%, so halving mortality raises large-juvenile survival by about 24% (0.676 → 0.838) but mature survival by only 12% (0.809 → 0.905).
Why recalculate rather than just scale the matrix entries? Large juveniles spend about 8 years in their stage. Higher survival raises the chance of staying alive long enough to grow out of it, so the "grow" entry rises proportionally more than the "stay" entry. (Rebuilding the matrix from Crouse et al.'s stage survivals and durations reproduces their published matrix to 4 decimal places.)
In the practical you will run your own "what if" experiments: predict, change the matrix, compare λ, explain.
""")

# ---------------------------------------------------------------- Hidden: a fast life history (red fox)
FOX_A = [["0.55", "0.77", "0.71", "0.74"], ["0.48", "0", "0", "0"], ["0", "0.54", "0", "0"], ["0", "0", "0.53", "0.51"]]
FOX_E = [["0.25", "0.15", "0.06", "0.06"], ["0.27", "0", "0", "0"], ["0", "0.12", "0", "0"], ["0", "0", "0.06", "0.05"]]

new("Title & Photo", {100: "A fast life history: the red fox",
                      102: "*Vulpes vulpes*, Bristol, UK; females, age classes (Devenish-Nelson et al. 2013). λ = 1.14."},
    pictures={101: F + "09_fox_diagram.png"}, hidden=True,
    notes="""
(1 min) HIDDEN: an optional contrast with the two long-lived species. Red foxes breed in their first year, have litters of several cubs, and rarely live more than a few years: annual survival is only about 0.5 at every age.
The matrix is age-based (a Leslie matrix with a final 3+ class that can stay put), from the review by Devenish-Nelson, Harris, Soulsbury, Richards & Stephens (2013) *Oikos* 122: 705–716, "Demography of a carnivore, the red fox: what have we learnt from 70 years of published studies?". Only 8 of the 56 populations they reviewed had enough data for a matrix model; this is the mean matrix for urban foxes in Bristol (as stored in COMADRE).
Post-breeding census: the age-0 loop (0.55) is cubs counted this year that survive and breed next year, as in the Leslie example earlier.
λ = 1.14. This population converges to its stable age distribution within about 3 years (damping ratio 5.3), compared with decades for the killer whale.
Figure: scripts/09_fox_diagram.R.
""")

s = new("Title Only", {100: "Fox: the big numbers move"}, hidden=True,
        notes="""
(1.5 min) HIDDEN. Same layout as the loggerhead slides: projection matrix A on the left, elasticity matrix E on the right.
Ask students to compare with the loggerhead: where are the big elasticities now? Answer: in the top row (reproduction, 0.25 + 0.15 + 0.06 + 0.06 = 0.51 in total) and in first-year survival (0.27). Survival of older foxes (0.12, 0.06, 0.05) matters much less. This matches the paper's conclusion that survival and fecundity of the youngest age class contribute most to population growth.
Values: scripts/00_check_numbers.R.
""")
colA = [[INK if v != "0" else None for v in row] for row in FOX_A]
colE = [[None if v == "0" else (CLAY if float(v) >= 0.10 else INK) for v in row] for row in FOX_E]
boldE = [[v != "0" and float(v) >= 0.10 for v in row] for row in FOX_E]
matrix(s, FOX_A, left=1.9, top=2.85, cw=0.95, rh=0.62, size=20, colors=colA,
       col_labels=["0", "1", "2", "3+"], name="**A** =", name_size=24, label_size=14)
matrix(s, FOX_E, left=7.9, top=2.85, cw=0.95, rh=0.62, size=20, colors=colE, bold=boldE,
       col_labels=["0", "1", "2", "3+"], row_labels=["0", "1", "2", "3+"], name="**E** =", name_size=24, label_size=14)
textbox(s, "Projection matrix (age this year → columns)", 0.8, 2.0, 5.6, 0.4, size=14, color=GREY, align=PP_ALIGN.CENTER)
textbox(s, "Elasticity of λ to each element (sums to 1)", 6.9, 2.0, 5.6, 0.4, size=14, color=GREY, align=PP_ALIGN.CENTER)
textbox(s, "Reproduction (top row) and first-year survival dominate **E**: the opposite of the loggerhead.",
        0.8, 5.75, 11.7, 0.5, size=16, color=INK, align=PP_ALIGN.CENTER)

new("Title, Bullets & Photo", {100: "Fast and slow life histories",
                               103: "TAKEAWAY",
                               105: (["Fox: half the elasticity is in reproduction",
                                      "Turtle and whale: survival dominates; reproduction ≤ 0.05",
                                      "Next lecture: why life histories differ"], 18)},
    pictures={101: F + "10_life_history_contrast.png"}, hidden=True,
    notes="""
(1.5 min) HIDDEN. Total elasticity of λ split three ways for each species: reproduction (fertility elements), survival before breeding age, and survival of breeders (including post-reproductive whales).
Fox (fast): reproduction 0.51, pre-breeding survival 0.27, breeders 0.22. Loggerhead (slow, late maturity): reproduction 0.05, pre-breeding survival 0.64, breeders 0.31. Killer whale (slow): reproduction 0.04, pre-breeding survival 0.42, breeders 0.54.
Generation time (T = ln R~0~ / ln λ, as on the generation-time slide): fox about 2 years, loggerhead about 18 years, killer whale about 28 years. General pattern (e.g. Heppell, Caswell & Crowder 2000; Sæther & Bakke 2000): in fast life histories λ is most elastic to reproduction and early survival; in slow ones to adult (or late juvenile) survival. This is why management advice differs between species, and it leads into next lecture on the evolution of life histories.
Figure: scripts/10_life_history_contrast.R (analytical, from the three published matrices).
""")

new("Title & Bullets", {100: "Assumptions of the basic model",
                        101: (["Vital rates are constant through time (no environmental variation)",
                               "No density dependence: vital rates do not depend on population size",
                               "Closed population: no immigration or emigration",
                               "All individuals in a stage are equivalent",
                               "Females only, with no shortage of mates",
                               "A fixed projection interval and census time"], 20)},
    notes="""
(1.5 min) These are close to the assumptions of the geometric growth model, plus the structure-specific ones. Under them, the population eventually grows (or declines) exponentially at rate λ with a stable stage distribution: that is a consequence of the assumptions, not an assumption (the previous version listed it as one).
None is a fundamental limitation of matrix models: stochastic models relax the first (section 04), density-dependent matrices the second, multi-site models the third, and finer stage classes or integral projection models the fourth.
Practical question 3 asks for these assumptions.
""")

new("Summary (Key Points Only)", {101: "Key points",
                102: (["Individuals of different ages and stages contribute differently to growth",
                       "A life cycle diagram converts directly to a matrix: columns from, rows to",
                       "Fertility includes the survival needed to reach the next census",
                       "n~t+1~ = **A**n~t~ projects size and structure; λ and the stable stage distribution are its dominant eigenvalue and eigenvector",
                       "Elasticity and *in silico* experiments rank management options"], 20)},
    notes="""
(1 min) One point per section. "MPMs are exponential growth models on steroids."
Reading: Neal (2004) Chapter 15, section 15.3 (schedule: "Neal 7" in the current edition).
Practicals: """ + URL_MPM + """ ; """ + URL_CENSUS + """ ; """ + URL_LTRE + """
""")

new("Activity", {101: "PRACTICAL · PAPER, THEN R",
                 102: ("Can you model the life cycle of an organism of your choice?", 30),
                 103: (["Draw a life cycle (2–3 stages, up to 7 arrows), add values, check them, write the matrix",
                        "Predict, then project: does the starting population matter?",
                        "Find λ, the stable stage distribution and the elasticities",
                        "*In silico*: predict, change a copy of the matrix, compare λ, explain",
                        "Evolution: is a mutant with one more offspring but lower juvenile survival fitter?"], 17)},
    notes="""
(30 s) Previews the practical "Matrix population modelling": """ + URL_MPM + """
The bullets follow the practical's sections in order: life cycle diagram and sanity checks; projection ("Does the starting population matter?"); growth rate and stable stage distribution; elasticity; in silico experiments; the evolutionary experiment, which treats λ as fitness (valid here because growth is density-independent).
The practical uses the popdemo package (project(), eigs(), elas()). Any organism, real or fantasy, plant or animal. Sanity checks: survival out of a stage ≤ 1; one arrow per non-zero cell.
LTRE (practical: """ + URL_LTRE + """): compares two matrices (e.g. two sites or treatments) and splits the difference in λ into contributions from each transition, roughly (difference in the element) × (sensitivity).
""")

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")

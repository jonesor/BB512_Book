"""Build "BB512 - 05 - Life tables and age-structured population growth" from BB512_Theme.pptx.
Run from the lecture folder:  python3 build_deck.py <path/to/BB512_Theme.pptx> <out.pptx>
Inline markup in text: *italic*, **bold**, ~subscript~, ^superscript^, `code` (Menlo).
"""
import re, sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

THEME, OUT = sys.argv[1], sys.argv[2]
FOREST = RGBColor(0x1F, 0x4D, 0x3F); INK = RGBColor(0x1F, 0x26, 0x23)
PALE = RGBColor(0xEF, 0xEE, 0xE6); BG = RGBColor(0xFB, 0xFA, 0xF6)
GREY = RGBColor(0x6E, 0x6F, 0x68); CLAY = RGBColor(0xA4, 0x50, 0x3A); GOLD = RGBColor(0xC0, 0x8A, 0x2E)
FONT = "Avenir Next"

prs = Presentation(THEME)
# Remove the theme's example slides
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
LAYOUT = {l.name: l for l in prs.slide_layouts}

TOKEN = re.compile(r"(`.+?`|\*\*.+?\*\*|\*.+?\*|~.+?~|\^.+?\^)")

def add_runs(par, text, size=None, color=None, bold=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        r = par.add_run()
        fmt = {}
        if part.startswith("`"):
            part = part[1:-1]; fmt["code"] = True
        elif part.startswith("**"):
            part = part[2:-2]; fmt["bold"] = True
        elif part.startswith("*"):
            part = part[1:-1]; fmt["italic"] = True
        elif part.startswith("~"):
            part = part[1:-1]; fmt["baseline"] = "-25000"
        elif part.startswith("^"):
            part = part[1:-1]; fmt["baseline"] = "30000"
        r.text = part
        if fmt.get("bold"): r.font.bold = True
        elif bold is not None: r.font.bold = bold
        if fmt.get("italic"): r.font.italic = True
        if fmt.get("code"): r.font.name = "Menlo"
        if "baseline" in fmt: r._r.get_or_add_rPr().set("baseline", fmt["baseline"])
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color

def set_text(shape, lines, size=None, color=None):
    tf = shape.text_frame
    if isinstance(lines, str):
        lines = [lines]
    # clear existing paragraphs but keep the first paragraph's properties
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, line in enumerate(lines):
        par = first if i == 0 else tf.add_paragraph()
        level = 0
        while line.startswith("  "):
            level += 1; line = line[2:]
        par.level = level
        add_runs(par, line, size=size, color=color)

def ph(slide, idx):
    for s in slide.placeholders:
        if s.placeholder_format.idx == idx:
            return s
    raise KeyError(f"placeholder idx {idx} not on layout {slide.slide_layout.name}")

def add_slide_number(slide):
    for lph in slide.slide_layout.placeholders:
        if lph.placeholder_format.type is not None and "SLIDE_NUMBER" in str(lph.placeholder_format.type):
            el = copy.deepcopy(lph._element)
            slide.shapes._spTree.append(el)
            # make sure it contains a slidenum field
            txBody = el.find(qn("p:txBody"))
            if txBody is not None and not txBody.findall(".//" + qn("a:fld")):
                pass
            return

def new(layout, fields=None, notes=None, pictures=None, number=False, hidden=False):
    slide = prs.slides.add_slide(LAYOUT[layout])
    if hidden:
        slide._element.set("show", "0")
    for idx, val in (fields or {}).items():
        if isinstance(val, tuple):
            set_text(ph(slide, idx), val[0], size=val[1])
        else:
            set_text(ph(slide, idx), val)
    for idx, path in (pictures or {}).items():
        ph(slide, idx).insert_picture(path)
    # remove unused empty placeholders (keeps the slide clean in Keynote)
    for s in list(slide.placeholders):
        if s.has_text_frame and not s.text_frame.text.strip() and s.placeholder_format.idx not in (fields or {}):
            s._element.getparent().remove(s._element)
    # No slide numbers: BB512 decks are built without them (the layouts' number placeholders are not copied)
    if number:
        pass
    if notes:
        slide.notes_slide.notes_text_frame.text = notes.strip()
    return slide

def table(slide, header, rows, col_w, top=2.15, row_h=0.52, size=15, first_col_bold=True, left=0.8,
          highlight_rows=(), left_cols=(0,)):
    nrows, ncols = len(rows) + 1, len(header)
    shape = slide.shapes.add_table(nrows, ncols, Inches(left), Inches(top), Inches(sum(col_w)), Inches(row_h * nrows))
    tbl = shape.table
    # strip the default table style banding
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    for i in range(nrows):
        tbl.rows[i].height = Inches(row_h)
        for j in range(ncols):
            cell = tbl.cell(i, j)
            text = header[j] if i == 0 else rows[i - 1][j]
            cell.text = ""
            par = cell.text_frame.paragraphs[0]
            par.alignment = PP_ALIGN.LEFT if j in left_cols else PP_ALIGN.CENTER
            if i == 0:
                add_runs(par, text, size=size, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
                cell.fill.solid(); cell.fill.fore_color.rgb = FOREST
            else:
                hl = (i - 1) in highlight_rows
                add_runs(par, text, size=size, color=CLAY if hl and j > 0 else INK,
                         bold=(j == 0 and first_col_bold) or hl)
                cell.fill.solid(); cell.fill.fore_color.rgb = PALE if i % 2 == 1 else BG
            for r in par.runs:
                r.font.name = FONT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
    return tbl

def textbox(slide, text, left, top, width, height, size=16, color=GREY, italic=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        add_runs(par, line, size=size, color=color)
        for r in par.runs:
            r.font.name = FONT
            if italic: r.font.italic = True
    return tb



def picture(slide, path, left=0.8, top=2.0, width=11.73, height=None):
    if height is None:
        return slide.shapes.add_picture(path, Inches(left), Inches(top), width=Inches(width))
    return slide.shapes.add_picture(path, Inches(left), Inches(top), Inches(width), Inches(height))


def lt_table(slide, cols, new=(), top=2.25, left=0.8, size=18, col_w=0.92, row_h=0.56, rows=None):
    """Life table with selected columns; columns in `new` get a gold header (the one to notice)."""
    data = rows or LT
    header = [HEAD[c] for c in cols]
    body = [[r[c] for c in cols] for r in data]
    nrows, ncols = len(body) + 1, len(header)
    shape = slide.shapes.add_table(nrows, ncols, Inches(left), Inches(top),
                                   Inches(col_w * ncols), Inches(row_h * nrows))
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    for j in range(ncols):
        tbl.columns[j].width = Inches(col_w)
    for i in range(nrows):
        tbl.rows[i].height = Inches(row_h)
        for j in range(ncols):
            cell = tbl.cell(i, j)
            text = header[j] if i == 0 else body[i - 1][j]
            cell.text = ""
            par = cell.text_frame.paragraphs[0]
            par.alignment = PP_ALIGN.CENTER
            is_new = cols[j] in new
            if i == 0:
                add_runs(par, text, size=size, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
                cell.fill.solid(); cell.fill.fore_color.rgb = GOLD if is_new else FOREST
            else:
                add_runs(par, text, size=size, color=CLAY if is_new else INK, bold=is_new)
                cell.fill.solid(); cell.fill.fore_color.rgb = PALE if i % 2 == 1 else BG
            for r in par.runs:
                r.font.name = FONT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.05)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
    return tbl


F = "figures/"; I = "images/"
URL = "https://jonesor.github.io/BB512_Book/life-tables-and-survivorship-types.html"

# The lecture's hypothetical cohort (the same numbers as Gotelli's Table 3.1 and the practical's
# "Worked example" sheet). Checked by scripts/00_check_numbers.R.
HEAD = {"x": "*x*", "S": "S~x~", "D": "D~x~", "l": "l~x~", "d": "d~x~", "q": "q~x~", "p": "p~x~",
        "e": "e~x~", "m": "m~x~", "lm": "l~x~m~x~", "xlm": "x l~x~m~x~"}
LT = [
    {"x": "0", "S": "500", "D": "100", "l": "1.0", "d": "0.2", "q": "0.2",  "p": "0.8",  "e": "1.8",   "m": "0", "lm": "0",   "xlm": "0"},
    {"x": "1", "S": "400", "D": "200", "l": "0.8", "d": "0.4", "q": "0.5",  "p": "0.5",  "e": "1.125", "m": "2", "lm": "1.6", "xlm": "1.6"},
    {"x": "2", "S": "200", "D": "150", "l": "0.4", "d": "0.3", "q": "0.75", "p": "0.25", "e": "0.75",  "m": "3", "lm": "1.2", "xlm": "2.4"},
    {"x": "3", "S": "50",  "D": "50",  "l": "0.1", "d": "0.1", "q": "1.0",  "p": "0",    "e": "0.5",   "m": "1", "lm": "0.1", "xlm": "0.3"},
    {"x": "4", "S": "0",   "D": "–",   "l": "0",   "d": "–",   "q": "–",    "p": "–",    "e": "–",     "m": "0", "lm": "0",   "xlm": "0"},
]
SUM_ROW = {"x": "Σ", "S": "", "D": "", "l": "", "d": "", "q": "", "p": "", "e": "", "m": "", "lm": "**2.9**", "xlm": "**4.3**"}

# ================================================================ 1 Title
new("Title", {103: ("Life tables and age-structured population growth", 36),
              104: "How do survival and reproduction at each age determine how fast a population grows?"},
    pictures={100: I + "title_elephants.jpg"},
    notes="""
(30 s) Guiding question for today. Elephant calves: an elephant does not breed until its teens and then has one calf every four years or so. Its survival and reproduction depend strongly on age, and that is what today is about.
Photo: African elephant calves (Loxodonta africana), cropped from the photo used on the "R0 is not lambda" slide of the previous version (file ELEPHANTS-K6C5EP_401765-1.jpg; looks like an Alamy stock image). Confirm the licence before sharing outside the course.
""")

# ================================================================ 2 Objectives
new("Agenda", {101: "Life tables",
               103: (["Explain why survival and reproduction must be tracked by age",
                      "Build a cohort life table and interpret l~x~, q~x~, p~x~ and e~x~",
                      "Distinguish the probability of dying (q~x~) from the force of mortality (μ~x~), and link the shape of mortality to survivorship curves",
                      "Calculate R~0~, T and r from a life table, and explain why generation time matters"], 19)},
    notes="""
(1 min) One objective per section of the lecture. The numbers we use (a hypothetical cohort of 500 females, Dall sheep skulls) are the evidence; the skills are general.
Reading: Neal, "Life tables and age-specific death rates", and the first section of "Age-specific reproduction and population growth rates" (calculating growth rates from age-specific birth and death rates). The matrix-model section of that chapter belongs with the next lecture.
Practical (second 45 minutes): "Life tables and survivorship types", in Excel. URL: """ + URL + """
""")

# ================================================================ 3 Why age
new("Title & Bullets", {100: "So far, every individual has been identical",
                        101: (["Exponential and logistic models use one rate (r or λ) for the whole population",
                               "That rate lumps together two very different processes: births and deaths",
                               "But a newborn, a 30-year-old and an 80-year-old differ in their chance of dying and of reproducing",
                               "Today: models with **age**. Next lecture: **stage** (matrix population models)"], 22)},
    notes="""
(1 min) Combines the two introduction slides of the previous version. "Not clones!": the earlier models treat every individual as interchangeable.
Ask: why might it matter for population growth if most individuals are juveniles, or most are old? (Juveniles do not breed yet; old individuals die soon.)
Age = time since birth. Stage = a category such as seed, juvenile, adult, or a size class; for many plants and invertebrates stage predicts fate better than age (next lecture).
""")

# ================================================================ Section 01
new("Section", {102: "01", 103: "Why age matters"},
    notes="(10 s) First: survival and reproduction both change with age.")

new("Title & Photo", {100: "A life is a sequence of stages",
                      102: "Stages differ in length: a year as a baby, decades as an adult."},
    pictures={101: F + "01a_stages.png"},
    notes="""
(45 s) Build 1 of 3. The age axis runs left to right; the stages are drawn the same width, although they are not the same length.
Photos carried over from the previous version of the lecture (the "teen" image is a record sleeve, "Die Grosse Teenager-Party"). Sources and licences to confirm.
""")

new("Title & Photo", {100: "The risk of dying changes with age",
                      102: "Schematic, humans, not to scale: high in infancy, lowest in childhood, then rising with age (senescence)."},
    pictures={101: F + "01b_stages_mortality.png"},
    notes="""
(1 min) Build 2 of 3. Ask before revealing: when in a human life is the risk of dying highest? Lowest?
This U shape (a "bathtub curve") is common in mammals and birds: Neal shows it for the cactus finch, Dall sheep, Himalayan thar and Canadian women. We will see real Dall sheep data in section 02.
The rise in mortality with age in adults is senescence (ageing), now named on the slide. Not to scale: in modern populations infant mortality is far lower than old-age mortality (in Neal's Canadian table, 1000q is 9.5 for infants but 152 at 85+); the shape is closer to historical populations.
Schematic: script 01_age_stage.R; hand-chosen values, no data.
""")

new("Title & Photo", {100: "So does reproduction",
                      102: "Schematic. Births and deaths come from different ages, so the age structure of a population affects its growth."},
    pictures={101: F + "01c_stages_mortality_fertility.png"},
    notes="""
(1 min) Build 3 of 3. Fertility is zero before maturity, peaks in early adulthood and declines. Replaces the old text overlay ("Different stages (and ages) have different mortality/survival and reproductive output"), which covered the figure.
Point: at first, a population full of young adults will grow faster than one of the same size full of children or old people, even with identical age-specific rates. In the long run the growth rate depends only on the age-specific rates, once the population settles to its stable age distribution (next lecture). So we need age-specific survival AND age-specific reproduction.
""")

s = new("Title Only", {100: "Ageing is diverse: match the pairs"},
        notes="""
(2.5 min) Activity, pairs. Red = mortality, blue = fertility, each from maturity to the age at which only 5% of adults are still alive (so the x-axes differ: centuries for one species, 4 years for another). Ask pairs to match the six species to the six panels, then take votes before the next slide.
Hints if needed: which organism might not age at all? Which is short-lived? Which has mortality that falls with age?
Figure: Jones et al. (2014) Diversity of ageing across the tree of life. Nature 505: 169–173, carried over from the previous version (screenshot, rearranged as a quiz). Teaching use only.
""")
picture(s, I + "jones2014_match.jpg", top=2.0, height=4.35)

s = new("Title Only", {100: "Senescence is not universal"},
        notes="""
(1.5 min) Answers (checked against the arrows on the previous version): 1 hydra, 2 netleaf oak, 3 great tit, 4 desert tortoise, 5 lion, 6 Scots pine.
Hydra (Hydra magnipapillata): flat mortality and fertility over centuries, no detectable ageing. Netleaf oak (Quercus rugosa) and desert tortoise (Gopherus agassizii): mortality FALLS with age. Lion (Panthera leo): strong senescence. Humans (not shown) have one of the steepest increases of all.
Senescence = mortality rising (and fertility falling) with age after maturity. Lesson: senescence, and the U shape of mammals, is not universal, which is why we need a tool that records mortality age by age: the life table.
The previous version also had two slides with all 46 species of Jones et al. (2014); they are unreadable at slide size and have been dropped. Students can look at Fig. 1 of the paper.
""")
picture(s, I + "jones2014_match.jpg", left=1.4, top=1.95, height=3.9)
textbox(s, "**1** hydra  ·  **2** netleaf oak  ·  **3** great tit  ·  **4** desert tortoise  ·  **5** lion  ·  **6** Scots pine",
        0.8, 5.95, 11.73, 0.45, size=19, color=CLAY, align=PP_ALIGN.CENTER)
textbox(s, "Jones et al. (2014) *Nature* 505: 169–173", 0.8, 6.4, 8.0, 0.35, size=14)

# ================================================================ Section 02
new("Section", {102: "02", 103: "Building a life table"},
    notes="(10 s) How do we turn counts of survivors into age-specific rates?")

new("Two Columns", {100: "Two ways to collect life-table data",
                    101: "COHORT LIFE TABLE",
                    103: (["Follow one group born together until all have died",
                           "Record S~x~, the number alive at each age",
                           "Best for short-lived, sedentary species",
                           "e.g. cactus finches, Galápagos"], 19),
                    104: "STATIC LIFE TABLE",
                    106: (["One snapshot: ages of the living, or ages at death",
                           "Assumes a stationary population and constant environment",
                           "Used for long-lived or mobile species",
                           "e.g. Dall sheep skulls, Alaska"], 19)},
    notes="""
(1.5 min) Neal calls these cohort (horizontal, generation) and static (stationary, time-specific, "vertical") life tables.
Cohort tables need no assumptions about population growth, but take a lifetime to collect and are hard for animals that disperse (dispersal looks like death).
Static tables are quick, but assume (1) the population is stationary (births = deaths, so the age structure equals survivorship), (2) the environment, and so age-specific mortality, has been constant, and (3) individuals can be aged accurately, and (for skulls) remains of all ages are equally likely to be found. Self-study slide at the end has the details.
History (corrected from the previous version): John Graunt (1662) analysed London's weekly "Bills of Mortality" in his "Natural and Political Observations"; Edmond Halley (1693) built the first true life table, from births and deaths in Breslau.
Grants' finches: Geospiza conirostris on Genovesa, 1244 nestlings banded 1978–83 (Neal's first example).
""")

new("Figure & Bullets", {100: "Following a cohort",
                         102: ["Each line is one individual, from birth (left) to death (●)",
                               "Count how many are still alive at the start of each year: S~x~",
                               "Age *x* means the **start** of an interval: *x* = 0 is age 0 to 1"]},
    pictures={101: I + "cohort_lollipop.jpg"},
    notes="""
(1 min) Carried over from the previous version ("Focus on the number entering and leaving intervals"). Count the lines that cross each dashed line: that is S(x) for that age.
The Lexis diagram of the previous version has been dropped: it was not explained on the slide and is not needed for what follows.
""")

def build_slide(title, cols, new_cols, text, notes, size=18):
    s = new("Title Only", {100: title}, notes=notes)
    lt_table(s, cols, new=new_cols, col_w=0.95 if len(cols) <= 6 else 0.88)
    left = 0.8 + (0.95 if len(cols) <= 6 else 0.88) * len(cols) + 0.5
    textbox(s, text, left, 2.25, 12.53 - left, 4.2, size=size, color=INK)
    return s

s = new("Title Only", {100: "From counts to survivorship"},
        notes="""
(1.5 min) Hypothetical cohort of 500 newborn females, followed until all have died: the same numbers as Gotelli's Table 3.1 and the "Worked example" sheet in the practical spreadsheet. Replaces five build slides of the previous version.
S(x) = number alive at the start of age x; D(x) = S(x) − S(x+1) = number dying during the year (100 in the first year, 200 in the second...). Check: D(x) sums to S(0) = 500 (Neal suggests this check).
l(x) = S(x)/S(0): the probability of surviving from birth to age x. l(1) = 400/500 = 0.8, so 80% of newborns reach their first birthday. It always starts at 1, so cohorts of different sizes, populations and species can be compared; some books use 1000 l(x) (Neal).
The plot is the same column as a survivorship curve (linear scale here; log scale later). Script 07_cohort_lx.R.
Notation reference (with d(x)) is a self-study slide at the end; students have it in the PDF.
""")
lt_table(s, ["x", "S", "D", "l"], new=["l"], col_w=0.95)
picture(s, F + "07_cohort_lx.png", left=6.4, top=2.05, width=5.57)
textbox(s, ["l~x~ = S~x~ / S~0~", "Probability of surviving from birth to age *x*"], 0.8, 5.75, 5.3, 0.9, size=18, color=INK)

s = build_slide("Probability of dying and of surviving", ["x", "S", "D", "l", "q", "p"], ["q", "p"],
                ["q~x~ = D~x~ / S~x~: chance of dying during the year, if alive at its start",
                 "q~1~ = 200 / 400 = 0.5",
                 "", "p~x~ = 1 − q~x~",
                 "Not the same as l~x~: p~2~ = 0.25 (age 2 to 3); l~2~ = 0.4 (birth to 2)",
                 "", "Life expectancy, e~x~: further years expected at age *x*; e~0~ = 1.8 years (worked in self-study)"],
                notes="""
(2 min) q(x) is the age-specific mortality rate in Neal's terminology: strictly a probability over the interval (section 03). q(1) = 0.5: half of those that reach age 1 die before age 2. q(3) = 1: all remaining individuals die during their fourth year.
p(x) = 1 − q(x) = l(x+1)/l(x); p(x) and q(x) always sum to 1. l(x) is cumulative, p(x) is for one step: l(3) = l(2) × p(2) = 0.4 × 0.25 = 0.1. This is the most common confusion in the practical (question 4). p(x) is g(x) in Gotelli; in next lecture's matrix models it becomes the survival entry P(i) of a Leslie matrix.
Life expectancy: the idea only. e(x) = (sum of l from age x onwards)/l(x) − 0.5; the arithmetic is a self-study slide, and students calculate e(x) in the practical. e(x) can rise with age when early mortality is high (Neal's finch: e(0) = 0.61, e(1) = 2.06 years).
The previous version's d(x) column (fraction of the original cohort dying in each interval) is in the self-study table.
""", size=17)

new("Title & Photo", {100: "Dall sheep: mortality is U-shaped with age",
                      102: "*Ovis dalli*, Denali, Alaska: 608 skulls aged from horn growth (Murie 1944; life table by Deevey 1947; data from Neal)."},
    pictures={101: F + "02_dall_sheep.png"},
    notes="""
(2 min) Real data. Olaus Murie collected 608 skulls of Dall sheep in what is now Denali National Park and aged them from their horns; Edward Deevey (1947) turned these ages at death into a static life table: all 608 were alive at birth, 121 died in the first year leaving 487, and so on.
Left: survivorship on a log scale. On a log scale the steepness of the curve is proportional to the mortality rate (the slope of ln l(x) is −μ; on these base-10 axes it is −μ/2.3): shallow from 1 to 7 years (prime adults rarely die), steep after 8.
Right: q(x) is U-shaped: lambs 0.2, prime adults about 0.015, rising to over 0.6 by age 10 (the main predator is the wolf, which takes lambs and old sheep). Ages 12 and 13 rest on only four and two sheep, so they are greyed out.
Assumptions (Neal): stationary population; skulls of all ages equally likely to be found (lamb skulls are fragile, so lamb mortality is probably underestimated).
Script 02_dall_sheep.R; empirical. Values checked in 00_check_numbers.R.
""")

# ================================================================ Section 03
new("Section", {102: "03", 103: "Mortality as a rate"},
    notes="(10 s) q(x) is a probability over a whole interval. What about mortality at an instant?")

new("Key Concept", {102: "Force of mortality, μ~x~",
                    103: "The instantaneous rate of dying at exact age *x*, per individual still alive. Also called the hazard.",
                    105: "q~x~ is a probability over a whole interval (0 to 1). μ~x~ is a rate at an instant (0 to ∞). q~x~ cannot tell whether deaths came early or late in the interval."},
    notes="""
(1.5 min) Replaces "A complication: probability is bounded" (its memento-mori AI image dropped).
Analogy: speed versus distance. μ is like your speed at one instant; q is like the fraction of a journey completed in an hour, which cannot exceed 100% however fast you drive.
Units: μ is per individual per unit time (e.g. per year). It can be greater than 1: μ = 2 per year means that, if it stayed constant, an individual would on average last half a year.
""")

new("Title & Photo", {100: "One mortality pattern, two descriptions",
                      102: "Illustrative: mortality rising exponentially with age (Gompertz). Dashed lines at 1."},
    pictures={101: F + "03a_same_mortality_two_ways.png"},
    notes="""
(1.5 min) Kept from the previous version (the staircase and the smooth curve), redrawn in one figure. Left: q(x), one step per year, levelling off at 1. Right: μ(x), continuous, passing 1 at about age 6.5 and reaching 8 by age 10.
We can convert an interval probability into an equivalent average hazard, but we cannot reconstruct how the hazard varied within the interval (from the previous notes).
Analytical: μ(x) = 0.02 e^(0.6x); q(x) = 1 − exp(−integral of μ over the year). Script 03_rate_vs_probability.R.
""")

new("Title, Bullets & Photo", {100: "Probability is bounded; rate is not",
                               103: "TAKEAWAY",
                               105: (["q = 1 − e^−μ^ (for a constant rate over the year)",
                                      "Low mortality: q ≈ μ",
                                      "High mortality: μ keeps rising, but q only approaches 1"], 18)},
    pictures={101: F + "03b_q_vs_mu.png"},
    notes="""
(2 min) CORRECTION from the previous version, which said: "Imagine 365 individuals. A rate of death of 1 per day would kill the whole population over a year (q = 1)." That mixes up a count of deaths per day with a per-capita rate. μ is a rate per individual still alive: as individuals die there are fewer left to die, so with any constant μ the expected fraction still alive falls exponentially and never reaches zero (a real, finite population can of course die out). q approaches 1 but never equals 1.
Numbers: μ = 0.1 per year gives q = 0.095; μ = 1 gives 0.63; μ = 2 gives 0.86; μ = 5 gives 0.993. Going the other way: μ = −ln(1 − q) = −ln(p); q = 0.5 corresponds to μ = 0.69; q = 0.99 to μ = 4.6.
At low mortality the two are nearly equal (dashed 1:1 line): that is why we can often treat them as the same for adults of long-lived species.
Analytical, script 03_rate_vs_probability.R; values in 00_check_numbers.R.
""")

s = new("Title Only", {100: "Mortality shape sets survivorship type"}, notes="""
(2 min) Why think of mortality as a rate that changes with age? Because how μ changes with age determines the shape of l(x) (kept from the previous notes). The slope of ln l(x) at age x is −μ(x) (on base-10 axes, −μ/2.3), so on any log scale steeper means higher mortality.
Type I: mortality rises with age, so l(x) stays high then falls steeply (humans in developed countries; zoo and domestic animals). The caption no longer says "large mammals": wild large mammals such as Dall sheep are usually a mixture (next slide).
Type II: constant mortality, a straight line on a log scale (Belding's ground squirrel in Neal; many songbirds roughly).
Type III: mortality falls with age: heavy early losses, then survivors live long (snapping turtle in Neal; fish, many trees with prolific seeds).
Real species are usually mixtures: Dall sheep and the cactus finch have Type III-like juvenile mortality and Type I-like adult mortality.
The practical's "Survivorship curves" sheet asks students to diagnose the three types from ln(l(x)).
Analytical, illustrative parameters: Gompertz, constant and declining hazards. Script 04_hazard_types.R. Redraws the figure of the previous version (the silhouettes are named on the slide instead).
""")
picture(s, F + "04_hazard_types.png", top=1.95, height=4.35)
textbox(s, "Type I: humans in developed countries  ·  Type II: many small birds  ·  Type III: fish, many trees. Illustrative hazards.",
        0.8, 6.3, 10.9, 0.35, size=14)

new("Check Question (4 options)", {101: ("Dall sheep: q~x~ = 0.2 for lambs, about 0.015 for prime adults, rising to over 0.6 by age 10. Which survivorship type, or mixture?", 26),
                                   105: "Type I", 109: "Type II", 113: "Type III", 117: "Type III early in life, then Type I"},
    notes="""
(1.5 min) Answer: D. Lamb mortality is high and falling (Type III-like), then mortality rises with age (Type I-like). Most mammals and birds look like this: a U-shaped hazard. The pure types are idealisations.
If students say "Type I": the log-scale survivorship curve is broadly Type I in shape (Neal compares it to humans), but the hazard tells us more than the curve alone.
""")

# ================================================================ Section 04
new("Section", {102: "04", 103: "From life table to population growth"},
    notes="(10 s) So far only deaths. Add births, and the life table tells us how fast the population grows.")

build_slide("Adding reproduction, m~x~", ["x", "l", "m", "lm"], ["m", "lm"],
            ["m~x~ = female offspring per female aged *x*",
             "", "Count daughters only: with a 50 : 50 sex ratio, m~x~ = births ÷ 2",
             "", "l~x~m~x~ = daughters expected at age *x* by a newborn female (survive to *x*, then reproduce)"],
            notes="""
(1.5 min) Why females only? The female segment of the population determines births; with a constant sex ratio, males grow at the same rate (Neal).
Example from Neal: 30 two-year-old females produce 120 young, so the birth rate is 4.0 and m = 2.0 daughters per female.
l(x)m(x) combines survival and reproduction: a female has 3 daughters at age 2 if she is alive, but only 40% of newborns are, so the expected contribution from age 2 is 1.2.
""")

build_slide("Net reproductive rate, R~0~", ["x", "l", "m", "lm"], ["lm"],
            ["R~0~ = Σ l~x~m~x~",
             "", "Daughters per female over her whole life: growth **per generation**",
             "", "Here: R~0~ = 0 + 1.6 + 1.2 + 0.1 = **2.9**",
             "", "R~0~ > 1 growing;  R~0~ = 1 stationary;  R~0~ < 1 shrinking"],
            notes="""
(1.5 min) Correction from the previous version: R0 is the NET REPRODUCTIVE RATE (Neal, Gotelli). "Basic reproduction number" is the epidemiological term for infections (the number of new cases caused by one case), which uses the same symbol.
Also corrected: R0 counts female offspring per female, not all offspring.
Each female replaces herself 2.9 times over: in one generation the population multiplies by 2.9. Merges the two R0 slides of the previous version (the second was a duplicate).
""", size=19)

s = build_slide("Generation time, T", ["x", "l", "m", "lm", "xlm"], ["xlm"],
                ["T = Σ *x* l~x~m~x~ / Σ l~x~m~x~",
                 "", "Mean age of mothers when their daughters are born",
                 "", "Here: T = 4.3 / 2.9 = **1.48 years**"],
                notes="""
(1.5 min) Added the formula, which the previous version described but did not show.
T weights each age by the daughters produced at that age. Neal's check: if l(x)m(x) were 2, 4 and 2 at ages 1, 2 and 3, T = (2 + 8 + 6)/8 = 2 years.
Sum of x l(x)m(x) = 0 + 1.6 + 2.4 + 0.3 = 4.3.
This is the cohort generation time; other definitions exist (next lecture), and they differ when the population is growing.
""", size=19)

new("Title, Bullets & Photo", {100: "R~0~ is per generation, not per year",
                               103: "TAKEAWAY",
                               105: (["λ and r: growth per unit of **time**",
                                      "R~0~: growth per **generation**",
                                      "A harvest mouse generation takes months; an elephant's about 25 years"], 18)},
    pictures={101: I + "mice_elephants.jpg"},
    notes="""
(1.5 min) Kept from the previous version. It is tempting to write r = ln(R0), by analogy with r = ln(λ). But λ is the multiplication per time step; R0 is the multiplication per generation. To get a rate per year we must divide by the length of a generation.
Numbers in context: suppose both had R0 = 2. Mouse with T = 0.5 years: r = ln 2 / 0.5 = 1.39 per year (λ = 4). Elephant with T = 25 years: r = 0.028 per year (λ = 1.03). Same R0, a 50-fold difference in r.
Reminder (from earlier lectures and the book's "Lambda and r" appendix): λ = e^r, r = ln λ; λ is a ratio of population sizes, r a per-capita rate.
Photos carried over from the previous version: harvest mice (Micromys minutus, as far as can be judged from the photo; file d41586-021-00216-7, from a Nature news article) and African elephant calves. Licences to confirm; teaching use only.
""")

new("Key Concept", {102: "r ≈ ln(R~0~) / T",
                    103: "Growth per year ≈ log of growth per generation ÷ generation length. r > 0 growing; r = 0 stationary; r < 0 shrinking.",
                    105: "Here r ≈ ln(2.9) / 1.48 = 0.72 per year (λ = e^r^ ≈ 2.05). An approximation: the exact r is 0.78 (λ = 2.17), from the Euler–Lotka equation (self-study) or a matrix model (next lecture)."},
    notes="""
(2 min) Derivation (Neal): N(T) = N(0) e^(rT) and N(T)/N(0) = R0, so ln(R0) = rT and r = ln(R0)/T. Neal calls this r(c), the capacity for increase.
Corrections from the previous version: (1) the equation images on the slide did not display; (2) it said "r = 0 when the population is stable": the right word is STATIONARY; in demography a "stable" population is one with a stable age distribution, which can be growing; (3) r = ln(R0)/T is an approximation, not an identity.
Why approximate? T is calculated for a cohort, but in a growing population mothers are on average younger, because there are more young females. The approximation is exact when generations do not overlap (annual plants, many insects) and good when r is small (Neal's ground squirrels: 0.048 vs 0.049).
Numbers checked in 00_check_numbers.R.
""")

new("Title, Bullets & Photo", {100: "Shorter generations, faster growth",
                               103: "THE NUMBERS",
                               105: (["R~0~ = 2 in both",
                                      "T = 1 year: r = 0.69, λ = 2",
                                      "T = 2 years: r = 0.35, λ = 1.41",
                                      "After 6 years: 6,400 vs 800"], 18)},
    pictures={101: F + "05_generation_time.png"},
    notes="""
(1.5 min) Makes "generation time is an important factor" (old summary) concrete. Both populations double every generation; the one with shorter generations doubles twice as often.
Consequence (link to life-history evolution later in the course): breeding earlier can raise fitness even without producing more offspring, and changes early in life matter more than changes late in life. The practical's "what if" experiment shows this.
Analytical: N(t) = 100 × 2^(t/T). Script 05_generation_time.R.
""")


new("Title, Bullets & Photo", {100: "Real data: fewer squirrels, faster growth",
                               103: "THE NUMBERS",
                               105: (["Uinta ground squirrel, Utah",
                                      "Density halved: R~0~ 1.07 → 1.59",
                                      "r ≈ 0.05 → 0.26 per year (λ 1.05 → 1.29)",
                                      "Better survival, not more births"], 17)},
    pictures={101: F + "08_uinta.png"},
    notes="""
(1.5 min) Real data for section 04 (Neal, "Calculating population growth rates from age-specific birth and death rates"). Uinta ground squirrel (Spermophilus armatus in Neal; now Urocitellus armatus), lawn habitat in Utah (Slade & Balph 1974). Every individual was marked and watched from observation towers. Undisturbed 1964–68; then density was roughly halved (1968–71).
Before: R0 = 1.069, T = 1.39 years, r(c) = 0.048 (λ = 1.05): almost stationary (Slade and Balph thought surplus animals emigrated). After: R0 = 1.589, T = 1.81 years, r(c) = 0.255 (λ = 1.29). Neal's exact r(m): 0.049 and 0.282.
Survivorship improved at every age and animals lived longer (the green line); m(x) hardly changed (1.96–2.73 before, 1.75–2.76 after). Link back to density dependence: lower density, higher survival, faster growth.
Note T rose (1.39 to 1.81 years) because more females survived to breed at older ages, yet r still rose five-fold.
Empirical, script 08_uinta.R; numbers checked in 00_check_numbers.R.
""")

new("Check Question (2 options)", {101: ("Population A: R~0~ = 3, T = 4 years. Population B: R~0~ = 2, T = 1 year. Which grows faster per year?", 28),
                                   105: "A: it has the higher R~0~", 109: "B: it has the shorter T"},
    notes="""
(1.5 min) Answer: B. r(A) ≈ ln 3 / 4 = 0.27 per year (λ ≈ 1.32); r(B) ≈ ln 2 / 1 = 0.69 per year (λ = 2).
Follow-up: what would A need to match B? R0 = e^(0.69 × 4) = 16 daughters per lifetime.
""")

new("Activity", {101: "PRACTICAL · EXCEL",
                 102: ("Build a life table, then find R~0~, T and r", 32),
                 103: ["*Life table* sheet: calculate l~x~, p~x~, q~x~, e~x~ and l~x~m~x~, then R~0~, T and r ≈ ln(R~0~)/T",
                       "What if? Add one daughter at one age at a time: which ages change R~0~ most?",
                       "*Survivorship curves* sheet: diagnose Types I, II and III from ln(l~x~)"]},
    notes="""
(30 s) Previews the practical "Life tables and survivorship types": """ + URL + """
The spreadsheet's "Worked example" sheet is today's lecture table, with the formulas. Model answers are in the book's Solutions chapter.
Expected results: R0 = 4.06, T = 3.78 years, r = 0.37 per year (exact 0.41); adding a daughter at age 1 raises R0 by 17%, at age 9 by 2%.
""")

new("Summary", {101: "Key points",
                102: (["Survival and reproduction change with age (senescence is not universal)",
                       "A life table turns counts (S~x~) into l~x~, q~x~, p~x~ and e~x~",
                       "q~x~ is a probability (at most 1); μ~x~ is a rate (no upper limit)",
                       "How mortality changes with age sets the survivorship type (I, II, III)",
                       "R~0~ = Σ l~x~m~x~ per generation;  T = Σ *x* l~x~m~x~ / R~0~;  r ≈ ln(R~0~) / T per year"], 17),
                106: (["**Read:** Neal, *Life tables and age-specific death rates*; start of *Age-specific reproduction and population growth rates*",
                       "**Practical:** Life tables and survivorship types",
                       "**Next:** matrix models"], 14)},
    notes="""
(1 min) One point per section, plus the growth equations.
Practical URL: """ + URL + """
""")

# ================================================================ Self-study
new("Section", {102: "S", 103: "Self-study: worked examples"},
    notes="Not presented in class; students have these slides in the PDF.")

s = new("Title Only", {100: "Self-study: life-table notation"},
        notes="""
Reference slide for the PDF; the lecture builds the main columns on the "From counts to survivorship" and "Probability of dying and of surviving" slides.
Notation varies between books: survival probability p(x) is g(x) in Gotelli; m(x) is b(x) in Gotelli; generation time T is G in Gotelli. The lecture, the practical spreadsheet and the book all now use Neal's symbols plus p(x).
Correction from the previous version: l(x) is not a number of individuals; S(x) (or N(x)) is the number, l(x) the proportion surviving from birth.
""")
table(s, ["Symbol", "Meaning", "Calculation"],
      [["S~x~", "Number alive at the start of age *x*", "counted"],
       ["D~x~", "Number dying between *x* and *x* + 1", "S~x~ − S~x+1~"],
       ["l~x~", "Survivorship: probability of surviving from birth to age *x*", "S~x~ / S~0~"],
       ["d~x~", "Fraction of the original cohort dying between *x* and *x* + 1", "l~x~ − l~x+1~"],
       ["q~x~", "Probability of dying between *x* and *x* + 1, if alive at *x*", "d~x~ / l~x~"],
       ["p~x~", "Probability of surviving from *x* to *x* + 1", "1 − q~x~"],
       ["e~x~", "Life expectancy: further years expected at age *x*", "self-study e~x~ slide"]],
      col_w=[1.4, 7.6, 2.7], top=2.0, row_h=0.52, size=17, left_cols=(0, 1))
textbox(s, "Notation varies between books, e.g. p~x~ is g~x~ in Gotelli.", 0.8, 6.25, 9.0, 0.35, size=14)

s = new("Title Only", {100: "Self-study: the complete life table"},
        notes="""
Every column for the lecture's cohort. Numbers checked in 00_check_numbers.R. The practical spreadsheet's "Worked example" sheet has the same table with the formulas.
""")
lt_table(s, ["x", "S", "D", "l", "d", "q", "p", "e", "m", "lm", "xlm"], new=(),
         rows=LT + [SUM_ROW], col_w=1.06, size=17, row_h=0.52, top=2.05)
textbox(s, "R~0~ = 2.9;  T = 4.3 / 2.9 = 1.48 years;  r ≈ ln(2.9) / 1.48 = 0.72 per year;  λ ≈ e^0.72^ = 2.05",
        0.8, 5.9, 11.73, 0.45, size=18, color=INK)

build_slide("Self-study: life expectancy, e~x~", ["x", "S", "D", "l", "d", "q", "p", "e"], ["e"],
            ["e~x~ = (sum of l from age *x* onwards) / l~x~ − 0.5",
             "", "e~0~ = (1 + 0.8 + 0.4 + 0.1) / 1 − 0.5 = 1.8 years",
             "e~1~ = (0.8 + 0.4 + 0.1) / 0.8 − 0.5 = 1.125 years",
             "", "− 0.5: deaths assumed spread evenly through each year"], size=17,
            notes="""
Correction from the previous version: the slide said "sum of l(x) from point x, minus 0.5 times the interval" and left out dividing by l(x) (the worked examples did divide).
Intuition: the sum of l from x onwards is the expected number of whole years lived after x by those alive at x, if all deaths happened at the end of a year; dividing by l(x) gives it per individual; subtracting half a year corrects for deaths happening on average halfway through the year. The −0.5 assumes one-year intervals (half the interval length in general).
e(2) = 0.5/0.4 − 0.5 = 0.75 and e(3) = 0.1/0.1 − 0.5 = 0.5 years (self-study slide).
Life expectancy can increase with age when early mortality is high (Neal's finch: e(0) = 0.61, e(1) = 2.06 years).
""")


new("Title & Bullets", {100: "Self-study: worked answers",
                        101: (["e~0~ = (1 + 0.8 + 0.4 + 0.1) / 1 − 0.5 = 1.8 years",
                               "e~1~ = (0.8 + 0.4 + 0.1) / 0.8 − 0.5 = 1.125 years",
                               "e~2~ = (0.4 + 0.1) / 0.4 − 0.5 = 0.75 years",
                               "e~3~ = 0.1 / 0.1 − 0.5 = 0.5 years",
                               "R~0~ = 1.6 + 1.2 + 0.1 = 2.9 daughters per female",
                               "T = (1 × 1.6 + 2 × 1.2 + 3 × 0.1) / 2.9 = 4.3 / 2.9 = 1.48 years",
                               "r ≈ ln(2.9) / 1.48 = 1.065 / 1.48 = 0.72 per year;  λ ≈ e^0.72^ = 2.05"], 19)},
    notes="Worked versions of the calculations shown in the lecture. Checked in 00_check_numbers.R.")

new("Title, Bullets & Photo", {100: "Self-study: finding r exactly",
                               103: "EULER–LOTKA",
                               105: (["Find r so that Σ e^−r*x*^ l~x~m~x~ = 1",
                                      "Try values of r until the sum is 1: r = 0.776 (λ = 2.17)",
                                      "ln(R~0~)/T = 0.718 gives 1.08: a slight underestimate",
                                      "Next lecture's matrix model gives the same λ"], 16)},
    pictures={101: F + "06_euler_lotka.png"},
    notes="""
The Euler–Lotka equation (Euler, 18th century; Lotka, early 20th century; Neal, "The intrinsic rate of natural increase and the Euler equation"). Logic: in a population growing at rate r with a stable age distribution, the daughters born now come from mothers born x years ago, when births were smaller by a factor e^(−rx); of those, a fraction l(x) survive and each has m(x) daughters. Adding over all ages must give back today's births, hence the sum equals 1.
At r = 0 the sum is just R0 = 2.9. Neal calls the exact value r(m), the intrinsic rate of natural increase; the population only grows at this rate once it has reached a stable age distribution.
The Leslie matrix built from the same table (post-breeding census, next lecture) has dominant eigenvalue 2.173 = e^0.776 (checked in 00_check_numbers.R).
Analytical, script 06_euler_lotka.R.
""")

new("Two Columns", {100: "Self-study: when static life tables mislead",
                    101: "ASSUMPTIONS",
                    103: (["Ages are estimated accurately",
                           "Remains of all ages are equally likely to be found",
                           "The population is stationary: births = deaths",
                           "The environment has been constant"], 18),
                    104: "IF THEY FAIL",
                    106: (["Fragile lamb skulls missed: juvenile mortality underestimated",
                           "Growing population: young over-represented",
                           "Declining population: old over-represented",
                           "Cohort tables avoid these problems but take a lifetime"], 18)},
    notes="""
From Neal, "Static (stationary or time-specific) life tables". Corrections for growth are possible (Caughley 1977) but assume exponential growth with a stable age distribution, which for a long-lived species like Dall sheep would mean constant growth for decades.
""")

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")

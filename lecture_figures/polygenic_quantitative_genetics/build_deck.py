"""Build "BB512 - 15 - Polygenic inheritance and quantitative genetics" from BB512_Theme.pptx.
Run from the lecture folder:  python3 build_deck.py <path/to/BB512_Theme.pptx> <out.pptx>
Inline markup in text: *italic*, **bold**, ~subscript~, ^superscript^.
"""
import re, sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

THEME, OUT = sys.argv[1], sys.argv[2]
FOREST = RGBColor(0x1F, 0x4D, 0x3F); INK = RGBColor(0x1F, 0x26, 0x23)
PALE = RGBColor(0xEF, 0xEE, 0xE6); BG = RGBColor(0xFB, 0xFA, 0xF6)
GREY = RGBColor(0x6E, 0x6F, 0x68); CLAY = RGBColor(0xA4, 0x50, 0x3A); GOLD = RGBColor(0xC0, 0x8A, 0x2E)
FONT = "Avenir Next"

prs = Presentation(THEME)
# Remove the theme's example slides
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
LAYOUT = {l.name: l for l in prs.slide_layouts}

TOKEN = re.compile(r"(`.+?`|\*\*.+?\*\*|\*.+?\*|~.+?~|\^.+?\^)")

def add_runs(par, text, size=None, color=None, bold=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        r = par.add_run()
        fmt = {}
        if part.startswith("`"):
            part = part[1:-1]; fmt["code"] = True
        elif part.startswith("**"):
            part = part[2:-2]; fmt["bold"] = True
        elif part.startswith("*"):
            part = part[1:-1]; fmt["italic"] = True
        elif part.startswith("~"):
            part = part[1:-1]; fmt["baseline"] = "-25000"
        elif part.startswith("^"):
            part = part[1:-1]; fmt["baseline"] = "30000"
        r.text = part
        if fmt.get("bold"): r.font.bold = True
        elif bold is not None: r.font.bold = bold
        if fmt.get("italic"): r.font.italic = True
        if fmt.get("code"): r.font.name = "Menlo"
        if "baseline" in fmt: r._r.get_or_add_rPr().set("baseline", fmt["baseline"])
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color

def set_text(shape, lines, size=None, color=None):
    tf = shape.text_frame
    if isinstance(lines, str):
        lines = [lines]
    # clear existing paragraphs but keep the first paragraph's properties
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, line in enumerate(lines):
        par = first if i == 0 else tf.add_paragraph()
        level = 0
        while line.startswith("  "):
            level += 1; line = line[2:]
        par.level = level
        add_runs(par, line, size=size, color=color)

def ph(slide, idx):
    for s in slide.placeholders:
        if s.placeholder_format.idx == idx:
            return s
    raise KeyError(f"placeholder idx {idx} not on layout {slide.slide_layout.name}")

def add_slide_number(slide):
    for lph in slide.slide_layout.placeholders:
        if lph.placeholder_format.type is not None and "SLIDE_NUMBER" in str(lph.placeholder_format.type):
            el = copy.deepcopy(lph._element)
            slide.shapes._spTree.append(el)
            # make sure it contains a slidenum field
            txBody = el.find(qn("p:txBody"))
            if txBody is not None and not txBody.findall(".//" + qn("a:fld")):
                pass
            return

def new(layout, fields=None, notes=None, pictures=None, number=False, hidden=False):
    slide = prs.slides.add_slide(LAYOUT[layout])
    if hidden:
        slide._element.set("show", "0")
    for idx, val in (fields or {}).items():
        if isinstance(val, tuple):
            set_text(ph(slide, idx), val[0], size=val[1])
        else:
            set_text(ph(slide, idx), val)
    for idx, path in (pictures or {}).items():
        ph(slide, idx).insert_picture(path)
    # remove unused empty placeholders (keeps the slide clean in Keynote)
    for s in list(slide.placeholders):
        if s.has_text_frame and not s.text_frame.text.strip() and s.placeholder_format.idx not in (fields or {}):
            s._element.getparent().remove(s._element)
    # No slide numbers: BB512 decks are built without them (the layouts' number placeholders are not copied)
    if number:
        pass
    if notes:
        slide.notes_slide.notes_text_frame.text = notes.strip()
    return slide

def table(slide, header, rows, col_w, top=2.15, row_h=0.52, size=15, first_col_bold=True, left=0.8,
          highlight_rows=()):
    nrows, ncols = len(rows) + 1, len(header)
    shape = slide.shapes.add_table(nrows, ncols, Inches(left), Inches(top), Inches(sum(col_w)), Inches(row_h * nrows))
    tbl = shape.table
    # strip the default table style banding
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    for i in range(nrows):
        tbl.rows[i].height = Inches(row_h)
        for j in range(ncols):
            cell = tbl.cell(i, j)
            text = header[j] if i == 0 else rows[i - 1][j]
            cell.text = ""
            par = cell.text_frame.paragraphs[0]
            par.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            if i == 0:
                add_runs(par, text, size=size, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
                cell.fill.solid(); cell.fill.fore_color.rgb = FOREST
            else:
                hl = (i - 1) in highlight_rows
                add_runs(par, text, size=size, color=CLAY if hl and j > 0 else INK,
                         bold=(j == 0 and first_col_bold) or hl)
                cell.fill.solid(); cell.fill.fore_color.rgb = PALE if i % 2 == 1 else BG
            for r in par.runs:
                r.font.name = FONT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
    return tbl

def textbox(slide, text, left, top, width, height, size=16, color=GREY, italic=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        add_runs(par, line, size=size, color=color)
        for r in par.runs:
            r.font.name = FONT
            if italic: r.font.italic = True
    return tb


F = "figures/"; I = "images/"
URL = "https://jonesor.github.io/BB512_Book/heritability-from-a-linear-regression.html"

# ---------------------------------------------------------------- 1 Title
new("Title", {103: ("Polygenic inheritance and quantitative genetics", 36),
              104: "How do many genes of small effect make a continuous trait, and how fast can it evolve?"},
    pictures={100: I + "title_nicotiana.jpg"},
    notes="""
(30 s) Guiding question for today. Mendel's peas had discrete traits (round or wrinkled); most traits we care about in ecology, agriculture and medicine (size, yield, timing, survival) vary continuously.
Photo: long-flowered tobacco, *Nicotiana longiflora*, the plant in Edward East's (1916) classic experiment, which we analyse in section 02. Photo: jacilluch, Flickr, "Flor de la Tarde - Nicotiana longiflora", CC BY-SA 2.0 (via Wikimedia Commons; carried over from the previous version of the lecture).
""")

# ---------------------------------------------------------------- 2 Objectives
new("Agenda", {101: "Quantitative genetics",
               103: (["Explain how many loci of small, additive effect produce continuous variation",
                      "Partition phenotypic variance into genetic and environmental components",
                      "Estimate narrow-sense heritability from a parent–offspring regression, and say what it does not mean",
                      "Predict the response to selection with the breeder's equation",
                      "Explain why heritability depends on the population and the environment",
                      "Interpret evidence of selection on quantitative traits in wild populations"], 18)},
    notes="""
(1 min) One objective per section. They are general skills; Nilsson-Ehle's wheat, East's tobacco, Darwin's finches, great tits and bighorn sheep are the evidence we use to practise them.
Reading: Neal (2004), "Polygenic inheritance, quantitative genetics and heritability", all sections except "Intelligence, race and societal class". That section is not covered in the lecture beyond the check question in section 05, but it is worth reading for its clear explanation of why within-group heritability says nothing about differences between groups.
After the lecture: practical "Heritability from a linear regression".
""")

# ---------------------------------------------------------------- 3 Hook
new("Title, Bullets & Photo", {100: "Many traits vary continuously",
                               103: "PUZZLE",
                               105: ("Mendel's alleles come in discrete types: A or a. How can they produce a smooth range of heights?", 18)},
    pictures={101: I + "living_histogram.jpg"},
    notes="""
(1 min) A "living histogram": college students arranged by height, 1975 (photo credited to Brian Joiner; reproduced in Joiner 1975, *The American Statistician*). Each column is one height class; the outline is a bell-shaped (approximately normal) distribution.
Ask the room: if height were controlled by one gene with two alleles and no dominance, how many height classes would we see? (Three.) So what is going on?
Answer, coming up: many genes (polygenes), each with a small effect, plus the environment.
Photo carried over from the previous version of the lecture; confirm the licence before sharing the slides outside the course.
""")

# ---------------------------------------------------------------- Section 01
new("Section", {102: "01", 103: "Polygenic traits"},
    notes="(10 s) First: how discrete genes give continuous variation.")

new("Figure & Bullets", {100: "Nilsson-Ehle's wheat: one trait, three loci",
                         102: (["Kernel colour in wheat (*Triticum aestivum*), 1909",
                               "Three loci, each with a red and a white allele",
                               "No dominance: each red allele adds the same amount of red",
                               "Pure white × pure dark red → F~1~ all intermediate",
                               "F~1~ × F~1~ → F~2~: predict!"], 18)},
    pictures={101: I + "wheat.jpg"},
    notes="""
(1.5 min) Context before content. Herman Nilsson-Ehle (Swedish geneticist, Lund) published the first genetic analysis of a quantitative trait in 1909 (Neal, "Polygenic inheritance").
The alleles are additive: the phenotype depends only on how many red alleles a plant carries, not which loci they are at.
The F~1~ is heterozygous at all three loci, so it is intermediate in colour. Ask students to predict the F~2~ before showing the next slide.
Photo: white (left) and red (right) wheat grain, carried over from the previous version of the lecture (file "White-wheat-left-and-red-wheat-right-used-in-this-study.png", from a research article). Source and licence to confirm.
""")

new("Title & Photo", {100: "Three loci give seven colour classes",
                      102: "Each dot is one of the 64 equally likely combinations of F~1~ gametes (8 × 8). After Neal (2004), wheat kernel colour figure."},
    pictures={101: F + "01_wheat_f2.png"},
    notes="""
(1.5 min) The F~1~ (heterozygous at three loci) makes 2^3^ = 8 kinds of gamete, so F~1~ × F~1~ gives 8 × 8 = 64 equally likely combinations.
Be precise about the words: there are 27 different **genotypes** (3 genotypes at each of 3 loci, 3^3^ = 27) but only 7 **phenotype classes**, because colour depends only on the number of red alleles (0 to 6).
The proportions 1 : 6 : 15 : 20 : 15 : 6 : 1 are binomial (the number of red alleles out of 6, each red with probability 1/2). The previous version of this slide said "7 genotypes" and "this forms a normal distribution"; both are corrected here.
Only 1 in 64 F~2~ plants is as white (or as dark) as a parent. Nilsson-Ehle used this logic in reverse: the fraction of F~2~ like the parents tells you how many loci are involved.
Schematic, analytical: script 01_wheat_f2.R.
""")

new("Title & Photo", {100: "More loci + environment = a bell curve",
                      102: "F~2~ of a cross between pure lines; additive loci. Analytical; environmental effect of illustrative size."},
    pictures={101: F + "02_more_loci.png"},
    notes="""
(1.5 min) Left to right: one locus gives three classes (1 : 2 : 1); three loci give seven. Add a little environmental variation and the classes blur together (Neal's hypothetical continuous kernel-colour figure): the trait now looks continuous, even though only three genes are involved. With 20 loci the distribution is smooth and close to normal.
Why normal? Each trait value is the sum of many small, independent contributions (alleles and environmental effects). Sums of many small independent effects tend towards a normal distribution (the central limit theorem). It is an approximation, not a law: a few genes of large effect, dominance or epistasis can make the distribution skewed or lumpy.
Neal's point: there is really no distinction between "meristic" traits (counted in discrete classes, e.g. seeds per pod) and continuous traits; both are quantitative traits.
Note also that the 20-locus distribution is narrower on this 0–1 scale: with more loci, each of small effect, extreme types become very rare.
Script 02_more_loci.R (binomial genotype frequencies convolved with normal environmental noise, SD = 0.06 on the 0–1 scale; no simulation).
""")

new("Check Question (4 options)", {101: ("Four loci, two additive alleles each. How many phenotype classes in the F~2~?", 30),
                                   105: "5", 109: "8", 113: "9", 117: "16"},
    notes="""
(1.5 min) Answer: 9. There are 2 × 4 = 8 alleles per plant, so 0 to 8 "plus" alleles: 9 classes. 16 is the number of gamete types (2^4^); 8 is the number of alleles.
Follow-up: what fraction of the F~2~ looks like one of the pure parents? 1/4^4^ = 1/256 for each parent.
""")

# ---------------------------------------------------------------- Section 02
new("Section", {102: "02", 103: "Partitioning phenotypic variation"},
    notes="(10 s) Genes are important, but so is the environment. How much of the variation is due to each?")

new("Key Concept", {102: "V~P~ = V~G~ + V~E~",
                    103: "Phenotypic variance (V~P~) is the sum of variance due to genetic differences between individuals (V~G~) and variance due to differences in their environments (V~E~).",
                    105: "This is about **variation among individuals in a population**, not about how much of one individual's trait \"comes from\" its genes. It assumes genetic and environmental effects simply add up (Neal 2004)."},
    notes="""
(1.5 min) Example from the previous version: a newborn animal with "good" genes can grow large, but only in a good environment.
Variance = the average squared deviation from the mean. V~P~ is easy to measure: measure the trait in many individuals. V~G~ and V~E~ are not directly observable, so we need clever experiments.
Common misconception: "height is 80% genetic" does not mean 80% of your height comes from your genes. It means 80% of the variation in height among people in that population is associated with genetic differences.
The additivity assumption fails when there is genotype-by-environment interaction (section 05) or when good genes and good environments go together (e.g. well-fed parents with well-fed offspring).
""")

new("Two Columns", {100: "Two ways to isolate one component",
                    101: "REMOVE GENETIC VARIATION", 103: ["Use genetically identical individuals: clones, inbred lines, or the F~1~ of two inbred lines",
                                                          "V~G~ = 0, so V~P~ = V~E~"],
                    104: "REMOVE ENVIRONMENTAL VARIATION", 106: ["Raise genetically different individuals in identical conditions",
                                                                 "V~E~ ≈ 0, so V~P~ ≈ V~G~",
                                                                 "Never perfect: development itself is noisy"]},
    notes="""
(1 min) Replaces the five build slides of the previous version ("Expressed symbolically" and "The logic of the experiments").
Left: any variation among genetically identical individuals must be environmental (including developmental noise and measurement error).
Right: in practice we cannot make environments truly identical: tiny differences in light, nutrients or position in the pot, and random developmental noise, always remain. So the left-hand approach is the one breeders and geneticists actually use, as East did.
Once we know V~E~, we can get V~G~ from a genetically variable group: V~G~ = V~P~ − V~E~.
""")

new("Figure & Bullets", {100: "East's tobacco flowers",
                         102: ["Long-flowered tobacco, *Nicotiana longiflora*",
                               "Edward East, 1912–1916: the first partitioning of variance",
                               "Trait: corolla (flower tube) length",
                               "Crossed a short-flowered and a long-flowered inbred line",
                               "Measured P~1~, F~1~ and F~2~ plants"]},
    pictures={101: I + "nicotiana_half.jpg"},
    notes="""
(1 min) Context first: organism, question, measurement. Edward M. East (Connecticut Agricultural Experiment Station, later Harvard) began these crosses in 1912 and published in 1916 (Neal, "Partitioning phenotypic variation into different components").
Both parental lines were inbred (homozygous), so each was genetically uniform. Ask: which generations have no genetic variation? (Both parents and the F~1~.) Which does? (The F~2~, because alleles segregate.)
Photo: jacilluch, Flickr, CC BY-SA 2.0 (carried over from the previous version).
""")

new("Title, Bullets & Photo", {100: "Only the F~2~ has genetic variation",
                               103: "TAKEAWAY",
                               105: ("The P~1~ lines and the F~1~ vary only because of the environment. The F~2~ is far more variable because alleles segregate.", 18)},
    pictures={101: I + "east_histograms.jpg"},
    notes="""
(1 min) Read the figure: the two parental lines (top, short and long), the F~1~ (middle, intermediate and as narrow as the parents) and the F~2~ (bottom, intermediate mean but much wider).
The F~1~ are all heterozygous at every locus and genetically identical, so their spread is environmental. The F~2~ mean is similar to the F~1~ mean, but the spread is much greater.
Figure: Neal (2004), figure of East's breeding experiments, data from East (1916), carried over from the previous version of the lecture. Copyright Blackwell/Wiley; teaching use only.
""")

new("Title, Bullets & Photo", {100: "Partitioning East's data",
                               103: "THE NUMBERS",
                               105: (["V~E~ = 5.2 (mean variance of P~1~, P~2~ and F~1~)",
                                      "V~P~ of the F~2~ = 40.5",
                                      "V~G~ = 40.5 − 5.2 = 35.3",
                                      "87% genetic, 13% environmental, **in this F**~2~"], 17)},
    pictures={101: F + "03a_variance_vg_ve.png"},
    notes="""
(1.5 min) Correction from the previous version: V~E~ = 5.2 is the average variance of the three genetically uniform groups (both parental lines and the F~1~), not the F~1~ alone (Neal, "Genetic and environmental components").
Then rearrange V~P~ = V~G~ + V~E~: V~G~ = V~P~ − V~E~ = 40.5 − 5.2 = 35.3.
Key caveat (Neal's first "general point"): these percentages are not fixed properties of corolla length. In the parental lines V~G~ = 0, so the same trait would be "0% genetic". In a more variable greenhouse V~E~ would be larger. The answer to "nature or nurture?" depends on the population and the environment.
Variance units: mm² (corolla length in mm). Script 03_variance_budget.R; empirical values from Neal (2004).
""")

s = new("Title Only", {100: "Genetic variance has parts"},
        notes="""
(1.5 min) Neal's hypothetical two-locus series (section "Partitioning the components of genetic variation"). Red alleles are capital letters.
Row 1, additive: each red allele adds one unit, wherever it is.
Row 2, dominance: one red allele at a locus does as much as two (A is dominant to a), so Aabb = AAbb.
Row 3, dominance plus epistasis: A only has an effect when B is also present (interaction between loci).
So V~G~ = V~A~ + V~D~ + V~I~ (Neal 2004). V~D~ and V~I~ are hard to separate, so they are lumped as non-additive variance.
Why does the additive part matter most? Parents pass on **alleles**, not genotypes. Dominance and epistasis depend on combinations of alleles that are broken up every generation by segregation and recombination, so they do not make offspring resemble their parents in a predictable way. V~A~ is the main cause of resemblance between relatives, and so of the response to selection.
""")
table(s, ["Genotype", "aabb", "Aabb", "AAbb", "AABb", "AABB"],
      [["Additive", "0", "1", "2", "3", "4"],
       ["Dominance (A over a)", "0", "2", "2", "4", "4"],
       ["Dominance + epistasis", "0", "0", "0", "4", "4"]],
      col_w=[3.6, 1.5, 1.5, 1.5, 1.5, 1.5], top=2.25, row_h=0.62, size=18)
textbox(s, ["V~G~ = V~A~ + (V~D~ + V~I~)",
            "additive + (dominance + interaction between loci)"], 0.8, 5.0, 11.0, 1.2, size=20, color=INK)
textbox(s, "Phenotypic scores. After Neal (2004), “Partitioning the components of genetic variation”.", 0.8, 6.25, 9.0, 0.35, size=14)

# ---------------------------------------------------------------- Section 03
new("Section", {102: "03", 103: "Heritability"},
    notes="(10 s) How much do offspring resemble their parents, and what does that tell us?")

new("Key Concept", {102: "Heritability, h² = V~A~ / V~P~",
                    103: "Narrow-sense heritability is the proportion of phenotypic variance in a population that is due to additive genetic effects.",
                    105: "Broad-sense heritability, H² = V~G~ / V~P~, includes dominance and interaction. Like V~P~, both describe a population in an environment, not an individual. h² = 0 means no additive variance, not that genes are unimportant."},
    notes="""
(1.5 min) Neal (2004), "Heritability". In this course, "heritability" means narrow-sense h² unless stated.
Why "h squared"? Historical: it was first calculated as the square of a path coefficient (Neal explains this in a footnote). It is not something squared that you need to take the root of.
Correction from the previous version: it said "if h² = 0, genetics is not at all important". Not so: a trait can have V~G~ > 0 (dominance or epistasis) and still have h² = 0; and a trait can matter enormously for the organism yet have h² ≈ 0 because selection has used up its additive variance (section 05).
East's tobacco: H² = 35.3/40.5 = 0.87.
""")

new("Title, Bullets & Photo", {100: "Parent–offspring slope = h²",
                               103: "TAKEAWAY",
                               105: (["Regress mean offspring value on **mid-parent** value (the mean of the two parents)",
                                      "Slope = h²; East's tobacco: 0.83",
                                      "On a single parent, slope = h²/2"], 17)},
    pictures={101: F + "04_east_regression.png"},
    notes="""
(2 min) East's F~2~ families: each point is one pair of parents and the mean of their offspring (Neal's parent–offspring regression figure).
Reference lines: if h² = 1 (dashed), offspring on average equal the mid-parent; if h² = 0 (dotted, flat), knowing the parents tells you nothing about the offspring. The real slope, 0.83, is in between: offspring "regress" part of the way back towards the population mean.
Single parent: each parent passes on half its alleles, so the slope of offspring on one parent estimates h²/2. The previous version just said "parent"; it is mid-parent that gives h² directly (the practical now uses mid-parent values).
Caveat (link to the practical's assumptions list): parents and offspring often share an environment (territory, diet, maternal care), which also makes them similar and inflates the estimate. Cross-fostering experiments break this link.
Figure: empirical. Points digitised by eye from Neal's parent–offspring regression figure (East 1916); line drawn with the published slope 0.8348 through the mean of the points. A least-squares fit to the digitised points gives 0.838. Script 04_east_regression.R.
""")

new("Title, Bullets & Photo", {100: "Completing the partition",
                               103: "THE NUMBERS",
                               105: (["V~A~ = h² × V~P~ = 0.8348 × 40.5 = 33.8",
                                      "V~D~ + V~I~ = V~G~ − V~A~ = 35.3 − 33.8 = 1.5",
                                      "Additive 83%, non-additive 4%, environment 13%"], 17)},
    pictures={101: F + "03b_variance_va.png"},
    notes="""
(1 min) Rearrange h² = V~A~/V~P~: V~A~ = h² V~P~. Then V~G~ = V~A~ + (V~D~ + V~I~): V~D~ + V~I~ = V~G~ − V~A~.
Correction from the previous version: the formula was written as V~A~ − V~G~ (the numbers, 35.3 − 33.8, were already the right way round).
Almost all of East's genetic variance is additive, which is why the F~2~ responds well to selection. This completes Neal's table summarising the partition of East's data.
Script 03_variance_budget.R.
""")

new("Activity", {101: "3 MIN · PAIRS",
                 102: "Two inbred lines of mice are crossed. Partition the variance in tail length.",
                 103: ["F~1~ variance = 2.5;  F~2~ variance = 12.5;  mid-parent–offspring slope = 0.75",
                       "Find h², V~E~, V~G~, V~A~ and V~D~ + V~I~",
                       "What is H², the broad-sense heritability?"]},
    notes="""
(3 min) Neal's end-of-chapter problem on mouse tails, parts (a) and (b). Part (c) comes after the breeder's equation.
Answers: h² = 0.75. V~E~ = 2.5 (the F~1~ is genetically uniform). V~P~ = 12.5, so V~G~ = 12.5 − 2.5 = 10. V~A~ = 0.75 × 12.5 = 9.375. V~D~ + V~I~ = 10 − 9.375 = 0.625. H² = 10/12.5 = 0.8.
Neal's problem does not say "mid-parent"; we assume it, as in his parent–offspring regression figure.
""")

new("Title, Bullets & Photo", {100: "Heritability in the wild: Darwin's finches",
                               103: "TAKEAWAY",
                               105: (["Medium ground finch, *Geospiza fortis*, Daphne Major",
                                      "Bill depth: h² ≈ 0.8",
                                      "Same slope in both years; the environment shifts the line up or down"], 17)},
    pictures={101: I + "finch_heritability.jpg"},
    notes="""
(1.5 min) No breeding experiment needed: in a marked wild population we can measure parents and their offspring.
Each point is one family: offspring bill depth against mid-parent bill depth, for families raised in 1976 (filled) and 1978 (open). The slopes are similar (h² ≈ 0.8; Boag 1983 reported h² = 0.79 for bill depth), but the 1978 line is lower: offspring in that year were smaller for a given mid-parent value. That is an environmental (year) effect on the mean, while the slope still estimates h².
Bill depth matters: after the 1977 drought, finches with deeper bills survived better because they could crack the large, hard seeds that remained. With h² this high, that selection produced a measurable evolutionary response in the next generation.
If you want a V~A~ example: V~A~ = h² × V~P~. The previous version used h² = 0.793 and V~P~ = 0.447 mm², giving V~A~ = 0.354 mm² (it printed 0.44 in one place; 0.793 × 0.447 = 0.354). These two values could not be traced to a source; confirm before using them.
Figure: Boag & Grant data (Boag 1983, *Evolution* 37: 877–894; Boag & Grant 1978, *Nature* 274: 793–794), carried over from the previous version. Teaching use only.
""")

new("Title, Bullets & Photo", hidden=True, fields={100: "Racehorses: is winning heritable?",
                               103: "ASK FIRST",
                               105: (["Lifetime winnings: **yes, but only just:** h² ≈ 0.1",
                                      "Stud fees: **no** detectable link to genetic merit"], 18)},
    pictures={101: I + "horse_racing.jpg"},
    notes="""
(1.5 min) HIDDEN by default to keep within 45 minutes; a good optional example of low h². Before revealing, ask the two questions and take a show of hands.
Wilson & Rambaut (2008) *Biology Letters* 4: 173–175, "Breeding racehorses: what price good genes?" Thoroughbred pedigree of 4,476 horses; lifetime earnings (log scale) of about 2,500 horses; stud fees of 554 stallions.
h² of lifetime earnings = 0.095 ± 0.034 (V~A~ = 0.299 of V~P~ = 3.13): about 90% of the variation in earnings is not additive genetic (training, jockey, luck, which races were entered).
Stud fees (the price a mare's owner pays to breed from a stallion) showed no relationship with the stallion's estimated breeding value for earnings: roughly 2 cents of expected prize money per dollar of stud fee, not significantly different from zero.
Lesson: a trait can be heritable and still respond only slowly to selection, and a high price is not evidence of good genes.
Photo carried over from the previous version; source and licence to confirm.
""")

# ---------------------------------------------------------------- Section 04
new("Section", {102: "04", 103: "Response to selection"},
    notes="(10 s) Heritability lets us predict how fast a trait will change under selection.")

new("Key Concept", {102: "The breeder's equation: R = h²S",
                    103: "The response to selection (R) equals the heritability (h²) multiplied by the selection differential (S).",
                    105: "S = mean of the selected parents − mean of the whole population.  R = mean of their offspring − mean of the whole population."},
    notes="""
(1 min) Intuition: h² is the slope of offspring on mid-parent. If the selected parents are S units above the population mean, their offspring are expected to be h² × S units above it.
S is sometimes called the selection differential or the strength of selection; Neal calls it the intensity of selection or selection pressure.
Assumptions: h² is constant and the environment is the same in both generations.
""")

new("Title, Bullets & Photo", {100: "Predicting the response: heavier sheep",
                               103: "WORKED EXAMPLE",
                               105: (["Flock mean 100 kg; breed only from sheep averaging 125 kg",
                                      "S = 125 − 100 = 25 kg",
                                      "h² = 0.5, so R = 0.5 × 25 = 12.5 kg",
                                      "Next generation: 112.5 kg"], 17)},
    pictures={101: F + "05_breeders_equation.png"},
    notes="""
(1.5 min) Kept from the previous version, redrawn. Ask: what if h² were 0? (R = 0: the offspring average 100 kg however hard we select.) What if h² were 1? (R = 25: offspring average 125 kg.)
The offspring lie between the selected parents and the old mean: "regression towards the mean" (Neal, "Response to selection").
Repeat each generation and the mean keeps moving (Neal's two-generation selection figure): two generations at S = 25 gives 125 kg, if h² stays at 0.5. Does it? Section 05.
Figure: analytical, illustrative values. Body weight drawn as normal with SD 15 kg (assumed); the heaviest 12% are selected (truncation at 117.6 kg gives a selected mean of exactly 125 kg); the offspring curve keeps the same SD for simplicity. Script 05_breeders_equation.R.
""")

new("Check Question (4 options)", {101: ("Mean tail length 80 mm; h² = 0.75. Parents with 88 mm and 92 mm tails: expected mean of their offspring?", 28),
                                   105: "80 mm", 109: "87.5 mm", 113: "90 mm", 117: "97.5 mm"},
    notes="""
(1.5 min) Answer: 87.5 mm. Mid-parent = 90, S = 90 − 80 = 10, R = 0.75 × 10 = 7.5, so 80 + 7.5 = 87.5 mm.
80 mm assumes h² = 0; 90 mm assumes h² = 1; 97.5 mm adds R to the mid-parent instead of to the population mean.
Follow-up (Neal's mouse-tail problem, part c): parents averaging 70 mm: S = −10, R = −7.5, offspring 72.5 mm.
Neal gives the units as cm; mm are more plausible for mice, and the arithmetic is the same.
""")

# ---------------------------------------------------------------- Section 05
new("Section", {102: "05", 103: "Heritability is not a fixed property"},
    notes="(10 s) h² is a ratio of two variances. Change either variance and h² changes.")

new("Title, Bullets & Photo", {100: "Heritability is slippery",
                               103: "h² = V~A~ / V~P~",
                               105: (["Selection removes alleles: V~A~ falls",
                                      "Inbreeding or a founder event: V~A~ falls",
                                      "A more variable environment: V~E~ rises",
                                      "Either way, h² falls"], 17)},
    pictures={101: I + "icy_road_signs.jpg"},
    notes="""
(1 min) Kept from the previous version (the "slippery" signs). Heritability is a ratio of two variances, measured in one population in one environment (Neal, "Heritability").
Examples from East: either parental line has h² = 0, because V~A~ = 0, even though corolla length in the F~2~ has h² = 0.83. Height of plants grown in uniform conditions will have a higher h² than the same seed stock grown on patchy soil, because V~E~ is smaller.
Image: stock photograph from the previous version (shutterstock_1015220521); check that the licence covers the course.
""")

new("Title, Bullets & Photo", {100: "Selection uses up additive variance",
                               103: "TAKEAWAY",
                               105: (["Illinois maize, selected for high or low seed oil since 1896",
                                      "h² fell to about a third in both lines",
                                      "The lines still diverged: V~A~ not yet exhausted"], 17)},
    pictures={101: F + "06_corn_heritability.png"},
    notes="""
(1.5 min) The Illinois long-term selection experiment on maize (*Zea mays*) began in 1896, before Mendel's work was rediscovered, from 163 ears with 4–6% oil (Neal, "Empirical examples of selection of quantitative characters"). The high and low lines left the original range within a few generations, and were still diverging after 76 generations.
As selection fixes favourable alleles, V~A~ shrinks and so does h² (Neal's table of oil-content heritability): high line 0.32 to 0.12, low line 0.50 to 0.15. With at least 20 loci affecting oil content, V~A~ takes a long time to run out; new mutations also add a little each generation.
Response continues for a long time, but more slowly: the breeder's equation is a short-term prediction.
Figure: empirical, Dudley (1977) via Neal (2004)'s table of oil-content heritability, plotted at the mid-point of each block of generations. Script 06_corn_heritability.R.
""")

s = new("Title Only", {100: "Fitness traits have low heritability"},
        notes="""
(2 min) Neal's table of heritabilities in domestic species (from Hartl & Clark 1989). Traits close to fitness (calving interval, eggs per hen, litter size, maize yield) have low h²; "trivial" traits have high h². Mousseau & Roff (1987, *Heredity*) found the same pattern in wild animals.
Neal's explanation: fitness traits have been under strong, persistent selection, which has used up their additive variance.
But that is not the whole story. Houle (1992, *Genetics* 130: 195–204) compared traits on a common scale (the additive genetic coefficient of variation) and found that fitness traits often have *more* additive variance than morphological traits. Their low h² arises mainly because V~E~ (and non-additive variance) is very large: fitness depends on everything else that happens to the organism.
So: low h² does not prove that V~A~ is small. Correction to the previous version, which said "traits closely related to fitness are always under strong selection" and that the pattern simply "aligns with" selection reducing V~A~.
""")
table(s, ["Species", "Trait less tied to fitness", "h²", "Trait closely tied to fitness", "h²"],
      [["Cattle", "Wither height", "0.60", "Calving interval", "0.25"],
       ["Poultry", "Egg weight", "0.55", "Eggs per hen", "0.10"],
       ["Pigs", "Back-fat thickness", "0.60", "Litter size", "0.15"],
       ["Maize", "Plant height", "0.53", "Yield", "0.13"]],
      col_w=[1.7, 3.3, 1.0, 3.6, 1.0], top=2.3, row_h=0.55, size=17)
textbox(s, ["**Two explanations:** selection has used up V~A~ (Neal), and",
            "fitness integrates many traits, so V~E~ is large (Houle 1992)"], 0.8, 5.15, 11.7, 1.0, size=19, color=INK)
textbox(s, "Approximate h² values from Neal (2004), after Hartl & Clark (1989).", 0.8, 6.25, 9.5, 0.35, size=14)

new("Check Question (2 options)", {101: ("Dog size: h² = 0.8 within white dogs and within black dogs. White dogs are larger. Is the difference genetic?", 28),
                                   105: "Yes: size is highly heritable", 109: "Can't tell from h²"},
    notes="""
(1.5 min) Answer: can't tell. Neal's thought experiment (section "Intelligence, race and societal class"): feed white puppies a rich diet and black puppies a poor one. Within each group, 80% of the variance in size is additive genetic, yet the 15 cm difference between the groups is entirely due to diet.
Heritability is measured within a population. It says nothing about the cause of differences between populations. Neal uses this to dismantle the hereditarian argument about IQ differences between groups (section "Intelligence, race and societal class"): high within-group h² of IQ does not imply that group differences are genetic, and the evidence points to environmental causes.
Also worth saying: even a trait with h² near 1 can be changed by the environment (Neal's phenylketonuria example: a diet abolishes the effect).
""")

new("Key Concept", {102: "Reaction norm",
                    103: "The range of phenotypes that a single genotype produces across a range of environments.",
                    105: "A sloping reaction norm means the trait is plastic (phenotypic plasticity). If genotypes have different slopes, there is a genotype-by-environment interaction (G×E)."},
    notes="""
(0.5 min) A reaction norm is a powerful way to see genes and environment at the same time. Real reaction norms are usually curved and have limits (an organism cannot breed earlier and earlier indefinitely).
""")

new("Title & Photo", {100: "Four patterns of reaction norms",
                      102: "Each line is one genotype measured across environments. Schematic."},
    pictures={101: F + "07_reaction_norms.png"},
    notes="""
(1.5 min) Replaces three slides of the previous version.
1. Plastic, no genetic variation: the environment changes the trait, but all genotypes respond identically (V~G~ = 0).
2. Genetic variation, not plastic: genotypes differ, the environment has no effect.
3. Both, no G×E: genotypes differ by a constant amount everywhere; the lines are parallel.
4. G×E: the effect of the environment depends on the genotype (or, equally, the effect of the genotype depends on the environment). The lines have different slopes and may cross.
Ask: in which panels would a common-garden experiment find genetic differences? (2, 3 and 4; in panel 4 the size, and even the sign, of the difference depends on which garden.)
Script 07_reaction_norms.R.
""")

new("Title, Bullets & Photo", {100: "G×E makes h² environment-specific",
                               103: "TAKEAWAY",
                               105: ("The same two genotypes give high V~G~ in environment C and none in B. h² measured in one environment need not apply in another.", 17)},
    pictures={101: F + "08_gxe_heritability.png"},
    notes="""
(1.5 min) Kept from the previous version, with a correction: where the lines cross (environment B) it is the genetic variance that is zero, not the phenotypic variance (the old slide said "Var(P) = 0"). Environmental variance remains, so V~P~ = V~E~ there and h² = 0.
In environment C the genotypes are far apart: V~G~, and so h², are high. The response to selection therefore also depends on the environment.
Schematic: two genotypes at equal frequency; V~G~ = (difference in genotype means / 2)²; environmental scatter not drawn and assumed equal in all environments. Script 08_gxe_heritability.R.
""")

new("Title & Photo", hidden=True, fields={100: "Plasticity and phenological mismatch",
                      102: "Reaction norms for timing against spring temperature. Schematic; based on oak, winter moth and great tit."},
    pictures={101: F + "09_phenology_mismatch.png"},
    notes="""
(2 min) Kept from the previous version and redrawn. Many organisms time breeding plastically, using temperature as a cue: in warm springs, oak leaves burst earlier, winter moth caterpillars hatch earlier and great tits (*Parus major*, musvit) lay earlier.
If all three have the same plasticity (left), the food chain stays in step as springs warm. If the bird is less plastic than its food (right), chicks hatch after the caterpillar peak in warm years: a phenological mismatch, which creates selection for earlier breeding.
Whether the population can respond depends on heritable variation in laying date, and in plasticity itself (G×E: genotypes with steeper reaction norms). Reaction norms are not straight lines indefinitely; there are limits, and limits differ between species.
HIDDEN by default to keep the lecture within 45 minutes; unhide it if there is time (it links nicely to the great tit example in section 06). Schematic, illustrative slopes. Script 09_phenology_mismatch.R.
""")

# ---------------------------------------------------------------- Section 06
new("Section", {102: "06", 103: "Selection in the wild"},
    notes="(10 s) Two examples where people have changed selection on a polygenic trait, and the trait has responded.")

new("Title, Bullets & Photo", {100: "Bird feeders and great tit bills",
                               103: "TAKEAWAY",
                               105: (["UK great tits have longer bills than Dutch ones (A)",
                                      "UK bills grew longer from the 1980s to 2000s (B)",
                                      "Gene variants for longer bills: more feeder visits",
                                      "Selection on a polygenic trait, caught in the act"], 16)},
    pictures={101: I + "bosse2017.jpg"},
    notes="""
(1.5 min) Bosse et al. (2017) Recent natural selection causes adaptive evolution of an avian polygenic trait. *Science* 358: 365–368.
Great tits (*Parus major*, musvit) from Wytham Woods, Oxford, and from the Netherlands were genotyped at hundreds of thousands of markers. Regions of the genome showing signatures of recent selection in the UK were enriched for genes affecting bill morphology (notably the collagen gene *COL4A5*). Bill length is polygenic and heritable.
(A) UK birds have longer bills than mainland European birds. (B) In Wytham, bill length increased in birds born between the 1980s and the late 2000s. Birds carrying the long-bill variants of *COL4A5* visited feeders more often, and those variants were associated with higher reproductive success.
Interpretation offered by the authors, with caution: British people spend about twice as much on bird food as mainland Europeans, and have done so for a century, which may have selected for longer bills. Quote from co-author Lewis Spurgin (UEA): "Although we can't say definitively that bird feeders are responsible, it seems reasonable to suggest that the longer beaks amongst British great tits may have evolved as a response to this supplementary feeding."
Figure from Bosse et al. (2017), carried over from the previous version; teaching use only.
""")

new("Figure & Bullets", {100: "Trophy hunting of bighorn sheep",
                         102: (["Bighorn sheep, *Ovis canadensis*, Ram Mountain, Alberta",
                               "Hunters may shoot rams with horns above a size limit",
                               "Big-horned rams are often shot before they breed much",
                               "Horn size and body weight are heritable",
                               "**Predict:** what happens to horn size?"], 18)},
    pictures={101: I + "bighorn_ram.jpg"},
    notes="""
(1 min) Context before data. Ask the room to predict the direction of change, and what they would need to know to predict its speed (h² and S: the breeder's equation).
Answer: selection on rams is against large horns, because large-horned rams are removed, often at 4–6 years old, before they achieve high mating success. S is negative, and horn length is heritable, so the response should be smaller horns.
Photo carried over from the previous version (file "licensed-image.jpeg"); source and licence to confirm.
""")

new("Figure & Bullets", {100: "Horns shrank under trophy hunting",
                         102: ["1972–2002: mean weight (a) and horn length (b) of 4-year-old rams both fell",
                               "Selection by hunters, on heritable traits, against large rams",
                               "An unintended response to artificial selection"]},
    pictures={101: I + "coltman2003.jpg"},
    notes="""
(1.5 min) Coltman, D. W. et al. (2003) Undesirable evolutionary consequences of trophy hunting. *Nature* 426: 655–658. Figure 2: (a) mean weight of 4-year-old rams (N = 133), (b) mean horn length (N = 119), (c) population size (ewes aged 2 or more plus yearlings), 1972–2002.
Using the pedigree, Coltman et al. estimated breeding values and found that the breeding values for horn size and weight declined too, which points to an evolutionary (genetic) response, not only an environmental one.
The interpretation has been debated: Traill, Schindler & Coulson (2014, *Proc. R. Soc. B*) argued that demographic and environmental change could explain much of the decline; Pigeon et al. (2016, *Evolutionary Applications*), using more years and genomic data, again found an evolutionary decline in horn size, slower than first thought, which reversed partly once hunting pressure fell.
Point for students: separating genetic and environmental change in the wild is hard, for exactly the reasons we met in sections 02 and 05.
Figure: Coltman et al. (2003) Fig. 2 (caption removed), carried over from the previous version; teaching use only.
""")

# ---------------------------------------------------------------- Summary
new("Summary", {101: "Key points",
                102: (["Many additive loci + environment → continuous traits",
                       "V~P~ = V~G~ + V~E~;  V~G~ = V~A~ + (V~D~ + V~I~)",
                       "h² = V~A~/V~P~ = slope of offspring on mid-parent",
                       "Response to selection: R = h²S",
                       "h² is specific to a population and environment",
                       "Wild traits respond to selection, including ours"], 17),
                106: (["**Read:** Neal (2004), *Polygenic inheritance, quantitative genetics and heritability*",
                       "**Practise:** Neal's end-of-chapter problem",
                       "**Practical:** Heritability from a linear regression"], 15)},
    notes="""
(1 min) One point per section.
Practical URL: """ + URL + """
""")

new("Activity", {101: "PRACTICAL · R",
                 102: ("How heritable is wing length in two bird populations?", 32),
                 103: ["Plot the data, then fit `lm(Offspring_value ~ Midparent_value)`",
                       "Read h² (and its confidence interval) from the slope; V~A~ = h² × V~P~",
                       "Compare the two populations: why might h² differ?"]},
    notes="""
(30 s) Previews the practical "Heritability from a linear regression": """ + URL + """
The practical uses mid-parent values (column Midparent_value), 100 families per population. Population 1 gives h² ≈ 0.73 (95% CI 0.49–0.98), population 2 h² ≈ 0.18 (−0.09 to 0.46); V~A~ is about 2.8 mm² in both, so the difference in h² is environmental. Model answers are in the book's Solutions chapter.
""")

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")

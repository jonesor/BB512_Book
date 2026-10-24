"""Build "BB512 - 07 - Evolution of life histories" from BB512_Theme.pptx.
Run from the lecture folder:  python3 build_deck.py <path/to/BB512_Theme.pptx> <out.pptx>
Inline markup in text: *italic*, **bold**, ~subscript~, ^superscript^, and backticks for code (Menlo).
Helper functions are copied unchanged from polygenic_quantitative_genetics/build_deck.py.
"""
import re, sys, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

THEME, OUT = sys.argv[1], sys.argv[2]
FOREST = RGBColor(0x1F, 0x4D, 0x3F); INK = RGBColor(0x1F, 0x26, 0x23)
PALE = RGBColor(0xEF, 0xEE, 0xE6); BG = RGBColor(0xFB, 0xFA, 0xF6)
GREY = RGBColor(0x6E, 0x6F, 0x68); CLAY = RGBColor(0xA4, 0x50, 0x3A); GOLD = RGBColor(0xC0, 0x8A, 0x2E)
FONT = "Avenir Next"

prs = Presentation(THEME)
# Remove the theme's example slides
sldIdLst = prs.slides._sldIdLst
for sldId in list(sldIdLst):
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
LAYOUT = {l.name: l for l in prs.slide_layouts}

TOKEN = re.compile(r"(`.+?`|\*\*.+?\*\*|\*.+?\*|~.+?~|\^.+?\^)")

def add_runs(par, text, size=None, color=None, bold=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        r = par.add_run()
        fmt = {}
        if part.startswith("`"):
            part = part[1:-1]; fmt["code"] = True
        elif part.startswith("**"):
            part = part[2:-2]; fmt["bold"] = True
        elif part.startswith("*"):
            part = part[1:-1]; fmt["italic"] = True
        elif part.startswith("~"):
            part = part[1:-1]; fmt["baseline"] = "-25000"
        elif part.startswith("^"):
            part = part[1:-1]; fmt["baseline"] = "30000"
        r.text = part
        if fmt.get("bold"): r.font.bold = True
        elif bold is not None: r.font.bold = bold
        if fmt.get("italic"): r.font.italic = True
        if fmt.get("code"): r.font.name = "Menlo"
        if "baseline" in fmt: r._r.get_or_add_rPr().set("baseline", fmt["baseline"])
        if size: r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color

def set_text(shape, lines, size=None, color=None):
    tf = shape.text_frame
    if isinstance(lines, str):
        lines = [lines]
    # clear existing paragraphs but keep the first paragraph's properties
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, line in enumerate(lines):
        par = first if i == 0 else tf.add_paragraph()
        level = 0
        while line.startswith("  "):
            level += 1; line = line[2:]
        par.level = level
        add_runs(par, line, size=size, color=color)

def ph(slide, idx):
    for s in slide.placeholders:
        if s.placeholder_format.idx == idx:
            return s
    raise KeyError(f"placeholder idx {idx} not on layout {slide.slide_layout.name}")

def add_slide_number(slide):
    for lph in slide.slide_layout.placeholders:
        if lph.placeholder_format.type is not None and "SLIDE_NUMBER" in str(lph.placeholder_format.type):
            el = copy.deepcopy(lph._element)
            slide.shapes._spTree.append(el)
            # make sure it contains a slidenum field
            txBody = el.find(qn("p:txBody"))
            if txBody is not None and not txBody.findall(".//" + qn("a:fld")):
                pass
            return

def new(layout, fields=None, notes=None, pictures=None, number=False, hidden=False):
    slide = prs.slides.add_slide(LAYOUT[layout])
    if hidden:
        slide._element.set("show", "0")
    for idx, val in (fields or {}).items():
        if isinstance(val, tuple):
            set_text(ph(slide, idx), val[0], size=val[1])
        else:
            set_text(ph(slide, idx), val)
    for idx, path in (pictures or {}).items():
        ph(slide, idx).insert_picture(path)
    # remove unused empty placeholders (keeps the slide clean in Keynote)
    for s in list(slide.placeholders):
        if s.has_text_frame and not s.text_frame.text.strip() and s.placeholder_format.idx not in (fields or {}):
            s._element.getparent().remove(s._element)
    if number and layout not in ("Title", "Photo"):
        add_slide_number(slide)
    if layout == "Photo - 3 Up":  # two-line captions: move up so they clear the footer
        lay = {l.placeholder_format.idx: l for l in slide.slide_layout.placeholders}
        for s in slide.placeholders:
            idx = s.placeholder_format.idx
            if idx in (101, 103, 105):
                s.left, s.width = lay[idx].left, lay[idx].width
                s.top, s.height = Inches(6.12), Inches(0.5)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes.strip()
    return slide

def table(slide, header, rows, col_w, top=2.15, row_h=0.52, size=15, first_col_bold=True, left=0.8,
          highlight_rows=()):
    nrows, ncols = len(rows) + 1, len(header)
    shape = slide.shapes.add_table(nrows, ncols, Inches(left), Inches(top), Inches(sum(col_w)), Inches(row_h * nrows))
    tbl = shape.table
    # strip the default table style banding
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    for i in range(nrows):
        tbl.rows[i].height = Inches(row_h)
        for j in range(ncols):
            cell = tbl.cell(i, j)
            text = header[j] if i == 0 else rows[i - 1][j]
            cell.text = ""
            par = cell.text_frame.paragraphs[0]
            par.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            if i == 0:
                add_runs(par, text, size=size, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
                cell.fill.solid(); cell.fill.fore_color.rgb = FOREST
            else:
                hl = (i - 1) in highlight_rows
                add_runs(par, text, size=size, color=CLAY if hl and j > 0 else INK,
                         bold=(j == 0 and first_col_bold) or hl)
                cell.fill.solid(); cell.fill.fore_color.rgb = PALE if i % 2 == 1 else BG
            for r in par.runs:
                r.font.name = FONT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
    return tbl

def textbox(slide, text, left, top, width, height, size=16, color=GREY, italic=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        add_runs(par, line, size=size, color=color)
        for r in par.runs:
            r.font.name = FONT
            if italic: r.font.italic = True
    return tb


F = "figures/"; I = "images/"
BOOK = "https://jonesor.github.io/BB512_Book/"
URL_EGGS = BOOK + "how-many-eggs-should-a-bird-lay.html"
URL_FOS = BOOK + "trade-offs-and-the-declining-force-of-selection.html"

# ---------------------------------------------------------------- 1 Title
new("Title", {103: ("Evolution of life histories", 36),
              104: "Why doesn't every organism mature at birth, breed constantly and live forever?"},
    pictures={100: I + "title_blue_whale.jpg"},
    notes="""
(30 s) Guiding question for today. My aim is to get the gist across; the textbook has the equations.
Photo: blue whale, *Balaenoptera musculus*, Neal's first example of a life history: about 150 t, a single calf every few years. Carried over from the previous version of the lecture; source and licence to confirm.
""")

# ---------------------------------------------------------------- 2 Objectives
new("Agenda", {101: "Life history evolution",
               103: (["Explain why life histories differ, using trade-offs in allocating limited resources",
                      "Use a cost–benefit argument to explain why intermediate trait values are often favoured",
                      "Evaluate evidence for Lack's hypothesis, including the cost of reproduction to parents",
                      "Explain why the force of selection declines with age, and how this lets senescence evolve",
                      "Predict how extrinsic mortality shapes the pace of life and ageing"], 18)},
    notes="""
(1 min) Objective 1 goes with section 01, 2 with section 02 (practised in the guppy check question), 3 with section 03, and 4 and 5 with section 04. They are general principles; Lack's starlings, kestrels and red deer are the evidence we use.
Reading: Neal (2004), "Evolution of life histories" (all sections). Neal is thinner on senescence than on clutch size; the lecture covers more of that.
After the lecture: practicals "How many eggs should a bird lay?" and "Trade-offs and the declining force of selection".
""")

# ---------------------------------------------------------------- 3 Hook
new("Statement", {101: ("A “Darwinian demon” would mature at birth, have endless offspring and never die. Why doesn't it exist?", 34),
                  103: "Law (1979); Neal (2004), “Evolution of life histories”"},
    notes="""
(1 min) Neal opens the chapter with this "ideal life history": live as long as possible, start breeding early, breed often, and produce vast numbers of offspring that all survive. Richard Law (1979) called this imaginary organism a Darwinian demon.
Ask the room: why doesn't it exist? Collect answers; the answer we build on today is that resources are limited, so improving one trait costs another (trade-offs), and the environment decides which costs are worth paying.
""")

# ---------------------------------------------------------------- 4-5 Variation
new("Photo - 3 Up", {101: (["**Bumblebee bat:** 2 g", "*Craseonycteris thonglongyai*"], 14),
                     103: (["**Blue whale:** about 150 t", "*Balaenoptera musculus*"], 14),
                     105: (["**Bristlecone pine:** ~4,800 years", "*Pinus longaeva*"], 14)},
    pictures={100: I + "third_bumblebee_bat.jpg", 102: I + "third_blue_whale.jpg", 104: I + "third_bristlecone.jpg"},
    notes="""
(1 min) Huge variation in life histories: size and lifespan.
Bumblebee bat (Thailand and Myanmar) and Etruscan shrew (*Suncus etruscus*) are the smallest mammals, about 2 g. Blue whale: the largest animal; adults about 150 t (Neal), up to about 170 t. Calves are 2–3 t and 6–8 m long at birth and about 20 t at weaning at around 7 months (Guinness World Records). (Corrected from the previous notes, which said a 12 t calf reaching 60 t in 6 months.)
Bats have the largest offspring relative to their mother's size of any mammals: a newborn is typically 20–30% of the mother's post-partum mass, against about 5% in humans.
Great Basin bristlecone pine: the oldest known individual ("Methuselah", White Mountains, California) is about 4,800 years old. Harriet the Galápagos tortoise (*Chelonoidis niger*) lived about 175 years (died 2006).
Photos carried over from the previous version; sources and licences to confirm (the bumblebee bat file is from Wikimedia Commons).
""")

new("Photo - 3 Up", {101: (["**Kiwi:** one huge egg", "*Apteryx australis*"], 14),
                     103: (["**Atlantic cod:** millions of eggs", "*Gadus morhua*"], 14),
                     105: (["**Orchids:** ~4 million seeds per pod", "Orchidaceae"], 14)},
    pictures={100: I + "third_kiwi.jpg", 102: I + "third_cod.jpg", 104: I + "third_orchid.jpg"},
    notes="""
(1 min) ...and in offspring number and size.
Kiwi: usually a single egg, up to about 20% of the female's body mass (the X-ray shows how much of her it fills).
Atlantic cod: about 500,000 eggs per kg of female; a record of about 9 million eggs has been reported for a 34 kg female. Almost all die.
Orchids: dust-like seeds. A single capsule of *Cycnoches chlorochilon* holds almost 4 million seeds (Kew); a plant can carry several capsules. Neal says "as many as 1 billion (10^9^)" for some species; I could not find support for that per-plant figure, so the slide uses the per-capsule one. Each seed is about the size of a fungal spore and depends on mycorrhizal fungi to germinate. Contrast: the coconut palm makes a few nuts of several kg each (Neal's number–size trade-off example).
Others from the previous version, for discussion: a Mexican dung beetle (*Copris*) produces only a handful of offspring in its life, one egg per dung ball; ghost moths (Hepialidae) can lay tens of thousands of eggs; annual and monocarpic plants breed once and die.
Photos carried over from the previous version; sources and licences to confirm.
""")

# ---------------------------------------------------------------- Section 01
new("Section", {102: "01", 103: "Life histories and trade-offs"},
    notes="(10 s) First: what a life history is, and why organisms can't have everything.")

new("Key Concept", {102: "Life history",
                    103: "The schedule of an organism's growth, reproduction and survival: size at birth, growth rate, age and size at maturity, number and size of offspring, and lifespan.",
                    105: "Fitness depends on how many offspring are produced and when. No strategy is “best”: each is a different solution to the same problem of leaving descendants."},
    notes="""
(1 min) For fitness, the most important events in life are reproduction and death. Life history theory explains the "design" of these traits as solutions to the problem of maximising fitness.
Link back to the life tables lectures: a life history is summarised by its survivorship (l~x~) and fecundity (m~x~) schedules, and selection acts on both (Neal, "Evolution of life histories", introduction).
""")

new("Key Concept", {102: "Trade-off",
                    103: "A fitness benefit through one trait comes at a cost to another, because energy, time and nutrients are limited.",
                    105: "Example: resources put into eggs this year cannot be spent on repairing the body to breed again next year. (A night at the pub is fun; the hangover is not.)"},
    notes="""
(1.5 min) Everyday analogies from the previous version: "don't put all your eggs in one basket" (spreading risk); big cars carry more people but use more fuel; the pub and the hangover.
The biological version: Gadgil & Bossert (1970): "An organism's life history may be looked upon as resultant of three biological processes, namely, maintenance, growth, and reproduction. Any organism has limited resources of time and energy at its disposal. The three component processes of the life history compete for these limited resources."
Neal's examples: number versus size of offspring (orchid seeds versus coconuts).
""")

new("Title, Bullets & Photo", {100: "One budget, three uses",
                               103: "TAKEAWAY",
                               105: (["Maintenance, growth and reproduction compete for the same budget. Different splits give slow and fast life histories.", "", "*Schematic: illustrative values, not data.*"], 18)},
    pictures={101: F + "01_allocation.png"},
    notes="""
(1.5 min) Redrawn from the elephant and mouse bars of the previous version. A slow life history (elephant) spends more on maintenance and growth, and breeds late and slowly; a fast one (mouse) puts much more into reproduction and little into maintenance.
This has consequences for population dynamics: fast species grow quickly but crash easily; slow species change slowly.
Keep this picture in mind: it comes back at the end as the "disposable soma" theory of ageing.
Schematic: shares are illustrative, not measured. Script 01_allocation.R.
""")

new("Two Columns", hidden=True, fields={100: "What decides the split?",
                    101: "EXTRINSIC: THE ENVIRONMENT", 103: ["Predation, disease, climate, food, competition",
                                                         "Set the chances of surviving, and the rewards for breeding now or later"],
                    104: "INTRINSIC: THE ORGANISM", 106: ["Trade-offs between traits",
                                                          "Physiological, developmental and genetic constraints",
                                                          "Limit which combinations of traits are possible"]},
    notes="""
(1.5 min) HIDDEN by default to keep within 45 minutes; the extrinsic side returns in section 04. Life histories evolve through the interaction of the environment and biological constraints. The two interact: for example, high predation (extrinsic) changes how the reproduction–survival trade-off (intrinsic) is resolved. We return to extrinsic mortality in section 04.
""")

new("Figure & Bullets", {100: "An early idea: r- and K-selection",
                         102: (["MacArthur & Wilson (1967); Pianka (1970)",
                                "Far below K (disturbed habitats): selection for high r: breed early, many small offspring",
                                "Near K (stable, crowded habitats): selection for competitive ability: breed late, few large offspring",
                                "Useful shorthand, but too simple (Neal)"], 17)},
    pictures={101: F + "02_logistic_rK.png"},
    notes="""
(2 min) Terms from the logistic growth model: dN/dt = rN(1 − N/K), with r the intrinsic rate of increase and K the carrying capacity.
Unstable or frequently disturbed environments (fire, frost, drought) were thought to keep populations well below K most of the time. There, growth is not limited by resources, so selection should favour traits that raise r: early maturity, high fecundity, short generation time. Stable environments keep populations near K, where offspring compete for resources, so selection should favour competitive ability, survival and fewer, better-provisioned offspring.
The theory bundled three things on one axis: a demographic context (below K versus near K), an ecological context (unstable versus stable) and a suite of traits. Neal's table of r- and K-selected correlates (after Pianka 1970) lists them.
Problems: many species don't fit (next slide), and laboratory selection experiments on protozoa and fruit flies only partly confirmed it (Neal). Neal's examples: the multimammate rat *Mastomys natalensis* (r-like) and red deer (K-like).
Figure: analytical logistic curve, illustrative parameters. Script 02_logistic_rK.R.
""")

new("Check Question (4 options)", {101: ("Elms are huge, live for centuries and breed for decades, yet make millions of tiny seeds. r or K?", 28),
                                   105: "r-selected", 109: "K-selected",
                                   113: "Both: adults and seeds face different selection", 117: "Can't tell without knowing the population's density"},
    notes="""
(1.5 min) Answer: C. Distractor D is tempting because r/K is defined by density relative to K, but the point here is that one life cycle can contain both r-like and K-like stages. Neal's American elm example. Canopy trees compete for light and space (K-like traits: large, long-lived, iteroparous), but a seedling can only establish when a canopy tree falls, which is rare and unpredictable. So it pays to make huge numbers of small, well-dispersed seeds (r-like). Different stages of one life cycle face different selection.
Neal's other "misfits": periodical cicadas and bamboos (long-lived but semelparous and synchronised) and Pacific sockeye salmon (large, long generation time, breed once).
""")

new("Key Concept", {102: "The fast–slow continuum",
                    103: "Across species, life-history traits tend to vary together: fast species mature early, breed a lot and die young; slow species do the opposite.",
                    105: "A description of a pattern, not a mechanism. Body size explains much of it, but not all: bats live far longer than mice of the same size (section 04)."},
    notes="""
(1.5 min) The modern replacement for r/K. The fast–slow continuum is more modest: it is mainly an empirical description of how traits covary across species, without assuming that population density relative to K is the cause.
Life histories also vary along other axes (for example, how reproduction is spread over the lifespan), and they are shaped by trade-offs, age-specific mortality and ecology, not by one selective regime.
""")

# ---------------------------------------------------------------- Section 02
new("Section", {102: "02", 103: "When to start breeding"},
    notes="(10 s) At what age and size should an organism mature?")

new("Two Columns", {100: "Breed now, or wait?",
                    101: "BENEFITS OF WAITING", 103: ["Grow larger first",
                                                      "Larger parents can make more, or bigger, offspring",
                                                      "Bigger offspring may survive better"],
                    104: "COSTS OF WAITING", 106: ["Risk of dying before breeding at all",
                                                   "Longer generation time: the lineage grows more slowly"]},
    notes="""
(2 min) Neal ("Generation times: when to start breeding"): r = ln(R~0~)/T, so for the same lifetime output, breeding earlier (smaller T) gives faster growth and higher fitness. Generation time usually has a much bigger effect on r than R~0~ does.
So why wait? Larger size: more offspring, fewer predators, more tolerance of harsh conditions, better competitors (Neal). And when R~0~ is close to 1, small improvements in fertility or survival can be worth more than breeding earlier.
Ask: what would tip the balance towards breeding early? (High mortality; no gain from size.)
""")

new("Title, Bullets & Photo", {100: "Costs and benefits give an intermediate optimum",
                               103: "TAKEAWAY",
                               105: ("Fitness ≈ benefit − cost. The optimum is where the gap between the curves is largest, not where they cross.", 18)},
    pictures={101: F + "03_optimality.png"},
    notes="""
(2 min) The logic of an optimality model. Benefits usually saturate (diminishing returns), costs usually accelerate. Net benefit (gold) is the vertical gap between them; it is greatest where the two curves have the same slope, so the optimum is an intermediate trait value.
Correction to the previous diagram, which drew two U-shaped curves and marked the optimum where they came closest together.
The same logic applies to age at maturity (this section) and clutch size (next section).
Schematic: B(x) = 10(1 − e^−x/3^), C(x) = 0.12x²; optimum at x = 3.85. Script 03_optimality.R.
""")

new("Check Question (4 options)", {101: ("Predators start killing many adults, but not juveniles. How should age at first breeding evolve?", 28),
                                   105: "Earlier", 109: "Later",
                                   113: "No change: only juveniles matter", 117: "No change: predation isn't heritable"},
    notes="""
(1.5 min) Answer: A. If adults are likely to die soon, the cost of waiting (dying before breeding) rises, so the optimum on the previous slide shifts to an earlier age. Draw it: the cost curve gets steeper and the optimum moves left.
Evidence: Trinidadian guppies (*Poecilia reticulata*). Where the cichlid *Crenicichla alta* eats mainly large adults, guppies mature earlier and smaller, and put more into reproduction, with more, smaller offspring; where the killifish *Rivulus hartii* eats mainly juveniles, the reverse. The differences are heritable. Moving guppies from a high-predation site to a stream with only *Rivulus* produced evolution in the predicted direction within 11 years, roughly 30–60 generations (Reznick, Bryga & Endler 1990, *Nature* 346: 357–359).
Follow-up: what if predators killed mainly juveniles? (Later maturity and fewer, larger offspring: the *Rivulus* streams.)
D is wrong because the trait that evolves is the guppies' age at maturity, not the predation.
""")

new("Figure & Bullets", hidden=True, fields={100: "Example: maternal age and infant mortality",
                         102: (["USA, 2006: infant deaths per 1,000 live births, by mother's age",
                                "Highest for the youngest and oldest mothers; lowest at 30–34",
                                "Young mothers may still be growing; with age, egg quality falls and complications rise",
                                "Caution: education and income also differ with age"], 17)},
    pictures={101: I + "infant_mortality.jpg"},
    notes="""
(1.5 min) HIDDEN by default to keep within 45 minutes. A familiar U-shaped cost curve in humans; the lowest infant mortality is at an intermediate maternal age.
Be careful with the interpretation: this is maternal age at a birth, not age at maturity, and in modern humans much of the pattern at young ages reflects social and economic factors rather than biology.
Chart carried over from the previous version ("Human infant mortality, USA 2006"; first bar is the total). Data presumably CDC/NCHS linked birth and infant death data; source to confirm.
""")

# ---------------------------------------------------------------- Section 03
new("Section", {102: "03", 103: "How many offspring?"},
    notes="(10 s) Once mature, how many offspring should an organism have at each breeding attempt?")

new("Figure & Bullets", {100: "Lack's hypothesis",
                         102: (["David Lack (1910–1973), Oxford ornithologist",
                                "Single-female clutches range from 1 egg (kiwi, albatrosses) to about 20 (partridges)",
                                "**Lack's hypothesis:** clutch size evolves to the number from which, on average, the most young are raised",
                                "The limit: how much food the parents can bring"], 17)},
    pictures={101: I + "nestbox.jpg"},
    notes="""
(1.5 min) Quote: "The number of eggs in the clutch has been evolved to correspond with that from which, on the average, the most young are raised" (Lack 1968, p. 5, quoted in Neal).
Too many chicks cannot all be fed; too few wastes the opportunity. Birds are good for testing this: eggs are easy to count, chicks can be ringed in the nest and recaptured later.
Correction to the previous version: it said clutch size varies "from 1 to about 50" and that an ostrich "can lay 50 eggs". A single female ostrich lays up to about a dozen eggs; nests of 20–40 (occasionally 50) contain eggs from several females.
The practical "How many eggs should a bird lay?" works through the logic with the fictional big bird; this slide and the next deliberately do not show its answer.
Photo: nest box clutch, carried over from the previous version; source to confirm (replace with a higher-resolution photo if available).
""")

new("Title, Bullets & Photo", {100: "Lack's Swiss starlings",
                               103: "THE DATA",
                               105: (["Most common brood: 5 eggs",
                                      "Most young recovered per brood: 7, then 5 and 6",
                                      "Fewer than 2% of 15,757 chicks were recovered: a weak test"], 17)},
    pictures={101: F + "04_lack_starlings.png"},
    notes="""
(2 min) Starlings (*Sturnus vulgaris*) in Switzerland, Lack (1948), as tabulated in Neal ("Number of offspring"). 3,516 broods; mean 4.5 eggs. Young were ringed in the nest and counted if recovered more than 3 months after fledging.
Big broods did not produce proportionally more survivors: nestlings in big broods were lighter (15-day-old chicks: 88 g in broods of 2, 78 g in broods of 5, 71 g in broods of 7), and lighter fledglings survive worse.
But the peak (0.108 young per brood at 7 eggs) is slightly above the most common clutch (5 eggs, 0.104). Neal notes that Lack "ensured the support of his hypothesis by combining the data for clutch sizes of 7 and 8". With only 306 recoveries, the data cannot really distinguish 5 from 7.
Neal's other explanation: clutch size is partly environmental, so a genotype "for 5" lays 3–7 eggs (Mountford 1968). Next: the cost to the parents.
Figure: empirical, published summary values. Script 04_lack_starlings.R.
""")

new("Activity", {101: "2 MIN · PAIRS",
                 102: "Broods of 7 raised slightly more young. Why might laying 5 still be best?",
                 103: ["Think about the parents, not just this year's chicks",
                       "Female A raises 6 young a year and survives 4 breeding seasons",
                       "Female B raises 5 young a year and survives 5 breeding seasons",
                       "Which leaves more descendants?"]},
    notes="""
(2 min) Neal's example. A: 6 × 4 = 24 young; B: 5 × 5 = 25 young. B is fitter, despite raising fewer young each year.
General point: what matters is lifetime reproductive success. If raising a big brood reduces the parent's own survival (a cost of reproduction), the best clutch is smaller than the one that fledges most young in a single season. Neal shows this with adult mortality and relative fitness plotted against brood size; the blue tit (Nur 1984) is his empirical example.
""")

new("Figure & Bullets", {100: "Testing it: kestrel brood manipulations",
                         102: (["Common kestrel, *Falco tinnunculus*, the Netherlands",
                                "Chicks moved between nests: broods reduced, left as they were, or enlarged",
                                "Parents of enlarged broods worked harder, and fewer of them survived to breed again",
                                "A cost of reproduction: today's extra chick is paid for later"], 17)},
    pictures={101: I + "kestrel.jpg"},
    notes="""
(1.5 min) Correction to the previous version, which named the species as the sparrowhawk and labelled a simulated clutch curve as Daan et al. (1990). The studies are on kestrels:
Dijkstra, C., Bult, A., Bijlsma, S., Daan, S., Meijer, T. & Zijlstra, M. (1990) Brood manipulations in the kestrel (*Falco tinnunculus*): effects on fitness of parents and offspring. *Journal of Animal Ecology* 59: 269–286.
Daan, S., Dijkstra, C. & Tinbergen, J. M. (1990) Family planning in the kestrel (*Falco tinnunculus*): the ultimate control of covariation of laying date and clutch size. *Behaviour* 114: 83–116.
Daily energy expenditure of parents (measured with doubly labelled water) rose with brood size (Deerenberg et al. 1995). A widely quoted summary (Gill, *Ornithology*): parents that raised two extra chicks had much higher winter mortality (about 60%) than parents with normal or reduced broods (about 29%). Check these figures against Dijkstra et al. (1990) before quoting them.
The previous version said that "increased clutch size increased reproductive success" in the first analysis, and that the conclusion changed once adult survival was included: the same lifetime argument as the previous slide.
Caveat for keen students: across bird species, experimental brood enlargements cost parental survival less consistently than this classic study suggests (recent meta-analyses), partly because high-quality individuals lay larger clutches and can afford them ("individual optimisation").
Remaining reproductive value (from the previous notes): the expected future reproduction from age x on; it declines with age after maturity, which links to section 04.
Photo: female kestrel, carried over from the previous version; source to confirm.
""")

new("Title, Bullets & Photo", {100: "Count over the whole life, not one season",
                               103: "TAKEAWAY",
                               105: (["If big broods cost parents survival, the lifetime optimum is smaller than Lack's one-season optimum.", "", "*Schematic: illustrative values, not data.*"], 18)},
    pictures={101: F + "05_lifetime_clutch.png"},
    notes="""
(1.5 min) Pale curve: young raised in one season (Lack's measure); it peaks at 6 here. Dark curve: young raised over the whole lifetime, when bigger broods lower the parent's chance of surviving to breed again; it peaks at 5.
Lack was partly right: the food limit is real. But there is also a trade-off between current reproduction and the parent's future reproduction.
Schematic, illustrative functions (chick survival falls with brood size; adult annual survival falls from 0.77 with 1 egg to 0.37 with 6). Both curves scaled to a maximum of 1. Script 05_lifetime_clutch.R.
""")

# ---------------------------------------------------------------- Section 04
new("Section", {102: "04", 103: "Why do organisms age?"},
    notes="(10 s) How long should an organism expect to live, and why does it deteriorate with age?")

new("Key Concept", {102: "Senescence",
                    103: "A decline in survival and/or reproduction with age after maturity: the death rate rises and fertility falls as an individual gets older.",
                    105: "Not inevitable wear and tear: ageing varies enormously across species, and some (e.g. some tortoises and long-lived plants) show little or none. So why does it evolve?"},
    notes="""
(1 min) Senescence = an age-related increase in mortality and/or decrease in fertility.
If ageing were simply wear and tear, all species would age in a similar way. They don't: Jones et al. (2014, *Nature* 505: 169–173, "Diversity of ageing across the tree of life") compared 46 species; some show steep increases in mortality with age (humans, killer whales), others almost flat or even falling mortality (some tortoises, hydra, some plants).
Lifespan is a balance between selection to live longer (more breeding attempts) and the forces that let mortality rise with age. The next slides explain why selection is weak late in life.
""")

new("Title & Photo", {100: "Thought experiment: when are offspring made?",
                      102: "Fictional species: from age 1, survival 0.8 a year and 10 offspring a year. Bars: share alive × 10."},
    pictures={101: F + "06a_thought_experiment_normal.png"},
    notes="""
(1.5 min) The fictional animal from the previous version (and, with 25 ages, in the practical). Each bar is the expected number of offspring at that age, l~x~ × m~x~: everyone who reaches age x has 10 offspring, but only 0.8^x−1^ of the cohort is still alive to do it.
Their sum is R~0~ = 49.54, our measure of fitness. Two-thirds of it (68%) comes from ages 1–5; hardly anything from ages 15+, simply because few individuals are still alive, even though this animal does not age at all.
The arithmetic is on the self-study slide at the end.
Script 06_thought_experiment.R.
""")

new("Title & Photo", {100: "A mutant: an early benefit with a late cost",
                      102: "+1 offspring at age 1, death before age 21: fitness rises from 49.54 to 50.42, so the mutant spreads."},
    pictures={101: F + "06b_thought_experiment_mutant.png"},
    notes="""
(1.5 min) The gain is 1 × l~1~ = 1.00 offspring. The cost is the reproduction lost at age 21: l~21~ × m~21~ = 0.0115 × 10 = 0.12. Net: 49.54 + 1.00 − 0.12 = 50.42.
Correction to the previous version, which gave 50.54 (49.54 + 1, ignoring the cost), and 49.42 for the mutant without the benefit (that one was right).
So a shorter-lived mutant is fitter. Selection on events late in life is weak because few individuals live long enough to experience them. In the practical, students find out how many late years can be sacrificed for one early offspring.
Script 06_thought_experiment.R.
""")

new("Title, Bullets & Photo", {100: "The force of selection fades after maturity",
                               103: "TAKEAWAY",
                               105: (["Strength of selection = share of lifetime reproduction still to come",
                                      "Full strength until breeding starts at 3, then falls",
                                      "Neal's **reproductive value** tells the same story: it peaks at first breeding"], 16)},
    pictures={101: F + "07_red_deer_force.png"},
    notes="""
(2 min) Genes that harm only old deer are barely "seen" by selection. Female red deer (*Cervus elaphus*) on Rhum, 1957, from the life table in Neal (from Lowe 1969). For each age, the strength of selection on surviving the next year = the share of lifetime reproduction that still lies ahead (Hamilton's 1966 force of selection; the population is treated as stationary, as Neal does).
Before age 3 nothing has been produced yet, so dying costs everything: strength 1. After that, each year of breeding "uses up" part of the total. By age 9 only 5% of lifetime reproduction is still to come.
Link to Neal: he explains the same idea with Fisher's reproductive value, v~x~ (expected future offspring of an individual of age x), which peaks at first breeding (1.69 at age 3) and then declines. Neal notes that Hamilton (1966) and Charlesworth (1980) refined Fisher's argument; Hamilton's force of selection is the quantity plotted here. The kink at ages 8–9 is in Lowe's data (Neal discusses whether it is real).
Previous version's human version of this curve: "Evolution doesn't care (so much) what happens to you after 45!" Same shape: flat until maturity, then falling to near zero at the end of reproduction.
Figure: empirical life table, analytical measure. Script 07_red_deer_force.R.
""")

new("Check Question (4 options)", {101: ("In red deer, which harmful mutation will selection remove most effectively?", 30),
                                   105: "Lethal at age 2", 109: "Lethal at age 6",
                                   113: "Lethal at age 10", 117: "All equally: death is death"},
    notes="""
(1.5 min) Answer: A. A deer that dies at age 2 loses all its lifetime reproduction (strength 1.0). Dying at age 6 loses about half (0.53); dying at age 10 loses about 5% (0.05). Selection against the late-acting mutation is about twenty times weaker.
Follow-up: what about a mutation lethal at age 1 versus age 2? Equally strong: before maturity, every age is worth the same.
""")

new("Title & Bullets", {100: "Three ideas about why ageing evolves",
                        101: (["**Mutation accumulation** (Medawar 1952): harmful mutations that act late in life are only weakly selected against, so they build up",
                               "**Antagonistic pleiotropy** (Williams 1957): genes with early benefits and late costs are favoured, like our mutant",
                               "**Disposable soma** (Kirkwood 1977): energy spent on reproduction is not spent on repair: the allocation trade-off again",
                               "All three depend on the declining force of selection"], 19)},
    notes="""
(2 min) Mutation accumulation: drift and recurrent mutation, not selection, decide the frequency of alleles whose effects appear only after most reproduction is over (e.g. late-onset genetic diseases in humans).
Antagonistic pleiotropy: one gene, two effects. Selection favours the early benefit because the late cost falls where selection is weak. Possible example: genes that raise fertility early but increase cancer risk later.
Disposable soma: the physiological version of antagonistic pleiotropy. The body (soma) only needs to last as long as it is likely to be useful; investing in perfect repair is wasted if you are likely to be eaten anyway. This links back to the maintenance–growth–reproduction budget in section 01.
These are not mutually exclusive; all three probably contribute. The practical "Trade-offs and the declining force of selection" explores antagonistic pleiotropy with a life table.
""")

new("Two Columns", {100: "Dangerous lives favour fast living",
                    101: "HIGH EXTRINSIC MORTALITY", 103: ["Future breeding is uncertain",
                                                            "Selection favours breeding early and investing heavily now",
                                                            "Less spent on growth and repair: faster ageing"],
                    104: "LOW EXTRINSIC MORTALITY", 106: ["Future breeding is likely",
                                                           "Selection favours growth, repair and survival",
                                                           "Delayed breeding can pay off: slower ageing"]},
    notes="""
(1.5 min) Extrinsic mortality = deaths from outside causes (predators, accidents, disease, harsh weather) that the organism can do little about by investing in its body. Maintenance means things like disease resistance, repairing DNA damage, and being able to run from predators.
This is Williams' (1957) prediction: when extrinsic mortality is high, the force of selection falls faster with age, so senescence should evolve to be faster.
Caveat for keen students: the theory is subtler than this. If extra mortality hits all ages equally and does not change density-dependent processes, it need not change the evolution of ageing (Abrams 1993; reviewed by Moorad, Promislow & Silvertown 2019, *TREE*). The prediction is still well supported by comparisons such as the next slide.
""")

new("Check Question (4 options)", {101: ("Island opossums have lived without predators for thousands of years. How should they differ from mainland opossums?", 28),
                                   105: "They age more slowly and live longer", 109: "They age faster and die younger",
                                   113: "No difference: ageing is wear and tear", 117: "They breed earlier and more intensively"},
    notes="""
(1.5 min) The island is Sapelo Island, Georgia, USA. Answer: A. Austad (1993, *Journal of Zoology* 229: 695–708), "Retarded senescence in an insular population of Virginia opossums (*Didelphis virginiana*)". Island opossums aged more slowly and lived longer than mainland opossums, which face predators such as dogs and foxes; reproduction also declined less with age (island two-year-olds kept larger litters, which mainland two-year-olds did not). Opossums normally live only 2–3 years.
D is the prediction for the opposite case (higher extrinsic mortality).
Popular accounts give the isolation time as anything from about 4,000 to 9,000 years; check Austad (1993) before quoting a figure.
""")

new("Figure & Bullets", {100: "Safer lives, slower ageing",
                         102: (["Flight lowers extrinsic mortality",
                                "Bats live several times longer than ground-dwelling mammals of the same size",
                                "Brandt's bat (*Myotis brandtii*, about 7 g): over 41 years in the wild",
                                "A house mouse (*Mus musculus*) rarely lives beyond 3–4 years, even in the lab"], 17)},
    pictures={101: I + "bat.jpg"},
    notes="""
(1 min) The comparative evidence, after the opossum prediction. Wilkinson & South (2002, *Ageing Research Reviews*): bats live on average about 3.5 times longer than non-flying placental mammals of similar size. Birds, also flying, live longer than mammals of the same size. Brandt's bat record: a male recaptured 41 years after ringing in Siberia (Podlutsky et al. 2005, *Journal of Gerontology*).
The previous version showed a rat and a bat to make this contrast.
Photo: noctule bat, carried over from the previous version; source and licence to confirm.
""")

new("Figure & Bullets", {100: "Life after reproduction: menopause",
                         102: (["Females live long after their last offspring: humans and four toothed whales (killer whale, pilot whale, beluga, narwhal)",
                                "Selection on survival should be almost zero after the last birth. So why live on?",
                                "**Grandmother hypothesis:** helping relatives raises **inclusive fitness** (own offspring plus relatives' offspring, weighted by relatedness)",
                                "Old killer whale females lead their groups to salmon when food is scarce"], 16)},
    pictures={101: I + "killer_whale.jpg"},
    notes="""
(1.5 min) Our model was a simplification: it counted only an individual's own offspring. Fitness can also include effects on the reproductive success of relatives, weighted by relatedness (kin selection, inclusive fitness: later in the course, in the animal behaviour part).
"Pilot whale" on the slide = short-finned pilot whale. Menopause (a long post-reproductive lifespan) is known in humans and four toothed whales (Ellis et al. 2018, *Scientific Reports*, adding belugas and narwhals). Killer whales: post-reproductive females lead collective movements, especially when salmon are scarce (Brent et al. 2015, *Current Biology*); breeding late in life would also bring mothers into competition with their daughters (Croft et al. 2017, *Current Biology*). Elephants, also long-lived and with matriarchs, keep breeding into old age: no menopause. Recent work suggests some wild chimpanzees (Ngogo, Uganda) also have a post-reproductive lifespan (Wood et al. 2023, *Science*); still debated.
Previous version's link: Grandmother hypothesis in killer whales and elephants, http://phenomena.nationalgeographic.com/2015/03/05/why-killer-whales-go-through-menopause-but-elephants-dont/
The previous version's kin selection slide (meerkats, relatedness diagram, black widow spider, praying mantis) is folded into this slide; the spider and mantis illustrate sexual cannibalism rather than kin selection.
Photo: killer whale (*Orcinus orca*), carried over from the previous version; source and licence to confirm.
""")

# ---------------------------------------------------------------- Summary
new("Summary (Key Points Only)", {101: "Key points",
                102: (["A life history is a schedule of growth, reproduction and survival",
                       "Limited resources force trade-offs; costs and benefits often give intermediate optima",
                       "Clutch size: count young over the parent's whole life, not one season",
                       "The force of selection fades after maturity, so senescence can evolve",
                       "High extrinsic mortality favours fast lives and faster ageing"], 18)},
    notes="""
(1 min) One point per section (section 04 has two).
Reading: Neal (2004), "Evolution of life histories". Self-study: the last two slides (in the PDF).
Practicals: """ + URL_EGGS + """ and """ + URL_FOS + """
""")

new("Activity", {101: "PRACTICAL · EXCEL",
                 102: ("How many eggs should a bird lay, and what is an early baby worth?", 32),
                 103: ["Big bird: combine clutch size and chick survival to find the best clutch",
                       "Life table: add one baby at age 1, remove late survival, and recalculate R~0~",
                       "How much late reproduction matches one extra early baby?"]},
    notes="""
(30 s) Previews the two practicals.
"How many eggs should a bird lay?": """ + URL_EGGS + """
"Trade-offs and the declining force of selection": """ + URL_FOS + """ (life table with 25 ages, survival 0.8, 10 offspring per year; R~0~ = 49.81, or 50.81 with one extra offspring at age 1).
""")

new("Title & Bullets", {100: "Questions to consider",
                        101: (["Why can some baby birds leave the nest almost at once, while others take weeks to fledge?",
                               "Why can foals run shortly after birth, while human babies take a year to walk?",
                               "Why do some sharks lay eggs while others give birth to live young?",
                               "Why do some animals help to raise young that are not their own?"], 19)},
    notes="""
(0 min: for review; students have the PDF.) Hints: altricial versus precocial young (offspring number versus size and parental investment); egg-laying versus live birth (investment per offspring); helping others' young (kin selection, inclusive fitness).
""")

# ---------------------------------------------------------------- Self-study
new("Title & Bullets", {100: "Self-study: the mutant in numbers",
                        101: (["Fitness: R~0~ = Σ l~x~m~x~ (expected lifetime offspring)",
                               "Normal: m~x~ = 10, l~x~ = 1, 0.8, 0.64, …, 0.0115 at ages 1–21, so R~0~ = 49.54",
                               "Mutant gains 1 × l~1~ = 1.00 offspring at age 1",
                               "… and loses l~21~ × m~21~ = 0.0115 × 10 = 0.12 at age 21",
                               "R~0~ = 49.54 + 1.00 − 0.12 = 50.42: the mutant is fitter",
                               "Practical: how many late years could it lose before it is no longer fitter?"], 18)},
    notes="""
Self-study slide for the PDF; not presented. Numbers checked in scripts/00_check_numbers.R. l~x~ = 0.8^x−1^ (l~1~ = 1).
""")

new("Title & Bullets", {100: "Self-study: Cole's paradox",
                        101: (["Annual plant: λ = B~A~ (seeds per plant)",
                               "Perennial: λ = B~P~ + s~A~ (seeds, plus the adult if it survives)",
                               "Equal fitness when B~A~ = B~P~ + s~A~: one extra seed would do. So why be perennial? (Cole 1954)",
                               "Solution (Charnov & Schaffer 1973): seeds rarely survive. With seedling survival s~j~: B~A~ = B~P~ + s~A~/s~j~",
                               "Example: s~A~ = 0.9, s~j~ = 0.05, so the annual needs 18 extra seeds",
                               "Perennials win when adults survive much better than juveniles"], 18)},
    notes="""
Self-study slide for the PDF; not presented. Neal (2004), "Frequency of reproduction: annual and perennial strategies". Should an organism breed once (semelparity) or repeatedly (iteroparity)? The textbook covers this; there is no time for it in the lecture.
""")

prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
